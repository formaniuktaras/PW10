"""InventoryLite GUI with cash-basis accounting and moving-average inventory.

The app focuses on a lightweight workflow for a trading business:
- Cash basis only: income/expense are registered when money changes hands.
- Inventory cost uses moving-average per product and warehouse.
- Direct-costing: only variable costs (purchase price) are included into COGS; fixed
  expenses are tracked separately via cash transactions.
"""
from __future__ import annotations

import csv
import io
import logging
import math
import traceback
import webbrowser
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import sys
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, messagebox, filedialog
import sqlite3
import re
from typing import Optional
from openpyxl import load_workbook

import db
import labels
from label_templates_ui import TemplateManagerDialog
from utils import (
    APP_NAME,
    VERSION,
    SingleInstance,
    backup_all_data,
    configure_logging,
    get_data_dir,
    get_db_path,
    get_lock_path,
    get_log_path,
    Settings,
    apply_base_currency_settings,
    get_base_currency_code,
    get_base_currency_name,
    get_base_currency_decimals,
    open_data_folder,
    open_file,
    restore_all_data,
    show_error,
    backup_database,
    bind_common_shortcuts,
)
from ui_components import DatePicker, TableFrame, simple_prompt


def ensure_rate_for_date(currency_code: str, rate_date: str) -> float:
    currency_code = currency_code.strip().upper()
    base_currency = get_base_currency_code()
    if not currency_code or currency_code == base_currency:
        return 1.0
    existing = db.rate_on_date(currency_code, rate_date)
    if existing is not None:
        return existing
    suggestion: float | None = None
    try:
        suggestion = db.rate_on_or_before(currency_code, rate_date)
    except Exception:
        suggestion = None
    while True:
        defaults = [f"{suggestion:.4f}" if suggestion else ""]
        values = simple_prompt(
            "Курс валюти",
            [f"Курс {currency_code} -> {base_currency} на {rate_date}"],
            defaults,
        )
        if not values:
            raise ValueError("Курс не вказано")
        try:
            rate = float(values[0])
        except ValueError:
            messagebox.showerror("Курс", "Введіть числовий курс")
            continue
        if rate <= 0:
            messagebox.showerror("Курс", "Курс має бути більшим за 0")
            continue
        db.add_currency_rate(currency_code, rate_date, rate)
        return rate


class InventoryApp(tk.Tk):
    def __init__(self, settings: Settings | None = None) -> None:
        super().__init__()
        self.title(APP_NAME)
        self.geometry("1180x720")
        self.iconbitmap(default="icons/app.ico") if Path("icons/app.ico").exists() else None
        self.settings = settings or Settings()
        apply_base_currency_settings(self.settings)
        self.status_var = tk.StringVar(value="Готово")
        self.status_bar: ttk.Label | None = None
        self.chart_palette = ["#2563eb", "#16a34a", "#f97316", "#a855f7", "#0ea5e9", "#ef4444", "#6366f1"]
        bind_common_shortcuts(self)
        self.create_menu()
        self.apply_settings()

        notebook = ttk.Notebook(self)
        notebook.pack(fill=tk.BOTH, expand=True)

        self.brands_frame = ttk.Frame(notebook)
        self.categories_frame = ttk.Frame(notebook)
        self.products_frame = ttk.Frame(notebook)
        self.counterparties_frame = ttk.Frame(notebook)
        self.warehouses_frame = ttk.Frame(notebook)
        self.channels_frame = ttk.Frame(notebook)
        self.currencies_frame = ttk.Frame(notebook)
        self.purchases_frame = ttk.Frame(notebook)
        self.extra_costs_frame = ttk.Frame(notebook)
        self.sales_frame = ttk.Frame(notebook)
        self.inventory_frame = ttk.Frame(notebook)
        self.cash_frame = ttk.Frame(notebook)
        self.stock_frame = ttk.Frame(notebook)
        self.reports_frame = ttk.Frame(notebook)
        self.export_frame = ttk.Frame(notebook)

        notebook.add(self.brands_frame, text="Бренди")
        notebook.add(self.categories_frame, text="Категорії")
        notebook.add(self.products_frame, text="Товари")
        notebook.add(self.counterparties_frame, text="Контрагенти")
        notebook.add(self.warehouses_frame, text="Склади")
        notebook.add(self.channels_frame, text="Канали продажу")
        notebook.add(self.currencies_frame, text="Валюти")
        notebook.add(self.purchases_frame, text="Закупівлі")
        notebook.add(self.extra_costs_frame, text="Супутні витрати")
        notebook.add(self.sales_frame, text="Продажі")
        notebook.add(self.inventory_frame, text="Інвентаризація")
        notebook.add(self.cash_frame, text="Каса")
        notebook.add(self.stock_frame, text="Залишки")
        notebook.add(self.reports_frame, text="Звіти")
        notebook.add(self.export_frame, text="Експорт")

        self.create_brands_tab()
        self.create_categories_tab()
        self.create_products_tab()
        self.create_counterparties_tab()
        self.create_warehouses_tab()
        self.create_channels_tab()
        self.create_currencies_tab()
        self.create_purchases_tab()
        self.create_extra_costs_tab()
        self.create_sales_tab()
        self.create_inventory_tab()
        self.create_cash_tab()
        self.create_stock_tab()
        self.create_reports_tab()
        self.create_export_tab()
        # "Про програму" is opened from the File menu

        self.refresh_all()

    # Menu
    def create_menu(self) -> None:
        menubar = tk.Menu(self)
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="Резервна копія всіх даних", command=self.on_backup_all)
        file_menu.add_command(label="Відновлення з резервної копії", command=self.on_restore_all)
        file_menu.add_separator()
        file_menu.add_command(label="Резервна копія БД", command=self.on_backup)
        file_menu.add_separator()
        file_menu.add_command(label="Налаштування", command=self.open_settings_dialog)
        file_menu.add_command(label="Скинути профіль гарячих клавіш", command=self.reset_hotkeys_profile)
        file_menu.add_separator()
        file_menu.add_command(label="Документація/FAQ", command=self.open_docs)
        file_menu.add_separator()
        file_menu.add_command(label="Про програму", command=self.show_about)
        file_menu.add_separator()
        file_menu.add_command(label="Вихід", command=self.destroy)
        menubar.add_cascade(label="Файл", menu=file_menu)
        self.config(menu=menubar)

    def default_workdir(self) -> Path:
        path = self.settings.get("files", "working_dir") or str(get_data_dir())
        try:
            return Path(path)
        except Exception:
            return get_data_dir()

    def open_settings_dialog(self, section: str = "general") -> None:
        SettingsDialog(self, section)

    def reset_hotkeys_profile(self) -> None:
        self.settings.set("Типовий", "hotkeys", "profile")
        self.settings.save()
        messagebox.showinfo("Гарячі клавіші", "Профіль клавіш скинуто до типового.")

    def open_docs(self) -> None:
        url = self.settings.get("support", "docs_url") or "https://example.com/docs"
        webbrowser.open(url)

    def apply_settings(self) -> None:
        self.apply_theme()
        self.apply_status_bar()
        self.apply_window_modes()
        self.apply_editor_font()

    def apply_theme(self) -> None:
        theme = (self.settings.get("general", "theme") or "system").lower()
        style = ttk.Style()
        try:
            if theme == "dark":
                palette = self._apply_dark_theme(style)
            elif theme == "light":
                palette = self._apply_light_theme(style)
            else:
                palette = self._apply_system_theme(style)
            self._apply_palette(palette)
        except tk.TclError:
            logging.warning("Не вдалося застосувати тему %s", theme)

    def _preferred_base_theme(self, style: ttk.Style) -> str:
        for candidate in ("vista", "xpnative", "clam", "default"):
            if candidate in style.theme_names():
                return candidate
        return style.theme_use()

    def _apply_system_theme(self, style: ttk.Style) -> dict[str, str]:
        base_theme = self._preferred_base_theme(style)
        style.theme_use(base_theme)
        return {
            "bg": style.lookup("TFrame", "background") or style.lookup(".", "background") or "#f0f0f0",
            "surface": style.lookup("TNotebook", "background") or "#e6e6e6",
            "surface_alt": style.lookup("TButton", "background") or "#d9d9d9",
            "text": style.lookup("TLabel", "foreground") or "#000000",
            "muted": "#555555",
            "accent": "#4a90e2",
        }

    def _apply_light_theme(self, style: ttk.Style) -> dict[str, str]:
        base_theme = self._preferred_base_theme(style)
        palette = {
            "bg": "#f7f7f7",
            "surface": "#ffffff",
            "surface_alt": "#ededed",
            "text": "#202020",
            "muted": "#5b5b5b",
            "accent": "#4a90e2",
        }

        if "inventorylite-light" not in style.theme_names():
            style.theme_create(
                "inventorylite-light",
                parent=base_theme,
                settings={
                    ".": {
                        "configure": {
                            "background": palette["bg"],
                            "foreground": palette["text"],
                            "fieldbackground": palette["surface"],
                            "troughcolor": palette["surface_alt"],
                            "bordercolor": palette["surface_alt"],
                            "focuscolor": palette["accent"],
                        }
                    },
                    "TFrame": {"configure": {"background": palette["bg"]}},
                    "TLabel": {
                        "configure": {
                            "background": palette["bg"],
                            "foreground": palette["text"],
                        }
                    },
                    "TButton": {
                        "configure": {
                            "background": palette["surface_alt"],
                            "foreground": palette["text"],
                            "padding": (10, 6),
                        },
                        "map": {
                            "background": [
                                ("pressed", palette["accent"]),
                                ("active", palette["surface"]),
                            ]
                        },
                    },
                    "TEntry": {
                        "configure": {
                            "fieldbackground": palette["surface"],
                            "foreground": palette["text"],
                            "insertcolor": palette["text"],
                        }
                    },
                    "TCombobox": {
                        "configure": {
                            "fieldbackground": palette["surface"],
                            "foreground": palette["text"],
                            "background": palette["surface_alt"],
                            "arrowsize": 14,
                        },
                        "map": {
                            "fieldbackground": [("readonly", palette["surface"])],
                            "background": [
                                ("active", palette["surface_alt"]),
                                ("readonly", palette["surface_alt"]),
                            ],
                        },
                    },
                    "TNotebook": {
                        "configure": {
                            "background": palette["bg"],
                            "tabmargins": (6, 3, 6, 0),
                        }
                    },
                    "TNotebook.Tab": {
                        "configure": {
                            "background": palette["surface_alt"],
                            "foreground": palette["muted"],
                            "padding": (12, 6),
                        },
                        "map": {
                            "background": [("selected", palette["surface"])],
                            "foreground": [("selected", palette["text"])],
                        },
                    },
                    "Treeview": {
                        "configure": {
                            "background": palette["surface"],
                            "fieldbackground": palette["surface"],
                            "foreground": palette["text"],
                            "bordercolor": palette["surface_alt"],
                            "lightcolor": palette["surface"],
                            "darkcolor": palette["surface_alt"],
                        },
                        "map": {
                            "background": [("selected", palette["accent"])],
                            "foreground": [("selected", palette["surface"]), ("!selected", palette["text"])],
                        },
                    },
                    "Treeview.Heading": {
                        "configure": {
                            "background": palette["surface_alt"],
                            "foreground": palette["text"],
                            "relief": "flat",
                        },
                        "map": {"background": [("active", palette["surface_alt"])]},
                    },
                },
            )

        style.theme_use("inventorylite-light")
        return palette

    def _apply_dark_theme(self, style: ttk.Style) -> dict[str, str]:
        base_theme = "clam" if "clam" in style.theme_names() else style.theme_use()
        palette = {
            "bg": "#2b2b2b",
            "surface": "#333333",
            "surface_alt": "#3a3a3a",
            "text": "#e6e6e6",
            "muted": "#c0c0c0",
            "accent": "#5a5a5a",
        }

        if "inventorylite-dark" not in style.theme_names():
            style.theme_create(
                "inventorylite-dark",
                parent=base_theme,
                settings={
                    ".": {
                        "configure": {
                            "background": palette["bg"],
                            "foreground": palette["text"],
                            "fieldbackground": palette["surface_alt"],
                            "troughcolor": palette["surface_alt"],
                            "bordercolor": palette["surface"],
                            "focuscolor": palette["accent"],
                        }
                    },
                    "TFrame": {"configure": {"background": palette["bg"]}},
                    "TLabel": {
                        "configure": {
                            "background": palette["bg"],
                            "foreground": palette["text"],
                        }
                    },
                    "TButton": {
                        "configure": {
                            "background": palette["surface_alt"],
                            "foreground": palette["text"],
                            "padding": (10, 6),
                        },
                        "map": {
                            "background": [
                                ("pressed", palette["accent"]),
                                ("active", palette["surface"]),
                            ]
                        },
                    },
                    "TEntry": {
                        "configure": {
                            "fieldbackground": palette["surface_alt"],
                            "foreground": palette["text"],
                            "insertcolor": palette["text"],
                        }
                    },
                    "TCombobox": {
                        "configure": {
                            "fieldbackground": palette["surface_alt"],
                            "foreground": palette["text"],
                            "background": palette["surface_alt"],
                            "arrowsize": 14,
                        },
                        "map": {
                            "fieldbackground": [("readonly", palette["surface_alt"])],
                            "background": [
                                ("active", palette["surface"]),
                                ("readonly", palette["surface_alt"]),
                            ],
                        },
                    },
                    "TNotebook": {
                        "configure": {
                            "background": palette["bg"],
                            "tabmargins": (6, 3, 6, 0),
                        }
                    },
                    "TNotebook.Tab": {
                        "configure": {
                            "background": palette["surface"],
                            "foreground": palette["muted"],
                            "padding": (12, 6),
                        },
                        "map": {
                            "background": [("selected", palette["surface_alt"])],
                            "foreground": [("selected", palette["text"])],
                        },
                    },
                    "Treeview": {
                        "configure": {
                            "background": palette["surface"],
                            "fieldbackground": palette["surface"],
                            "foreground": palette["text"],
                            "bordercolor": palette["surface_alt"],
                            "lightcolor": palette["surface"],
                            "darkcolor": palette["surface_alt"],
                        },
                        "map": {
                            "background": [("selected", palette["accent"])],
                            "foreground": [("selected", palette["text"])],
                        },
                    },
                    "Treeview.Heading": {
                        "configure": {
                            "background": palette["surface_alt"],
                            "foreground": palette["text"],
                            "relief": "flat",
                        },
                        "map": {"background": [("active", palette["surface_alt"])]},
                    },
                },
            )

        style.theme_use("inventorylite-dark")
        return palette

    def _apply_palette(self, palette: dict[str, str]) -> None:
        try:
            self.configure(background=palette.get("bg"))
        except tk.TclError:
            logging.debug("Tk configure not available for background")
        try:
            self.tk_setPalette(
                background=palette.get("bg"),
                foreground=palette.get("text"),
                activeBackground=palette.get("surface"),
                activeForeground=palette.get("text"),
                highlightColor=palette.get("accent"),
                highlightBackground=palette.get("surface"),
                insertBackground=palette.get("text"),
                troughColor=palette.get("surface"),
            )
        except tk.TclError:
            logging.debug("Tk palette is not available")

    def apply_status_bar(self) -> None:
        show_status = bool(self.settings.get("ui", "status_bar"))
        if show_status and not self.status_bar:
            self.status_bar = ttk.Label(self, textvariable=self.status_var, relief=tk.SUNKEN, anchor="w")
            self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        elif not show_status and self.status_bar:
            self.status_bar.destroy()
            self.status_bar = None

    def apply_window_modes(self) -> None:
        self.attributes("-fullscreen", bool(self.settings.get("ui", "fullscreen")))
        compact = bool(self.settings.get("ui", "compact_mode"))
        try:
            self.tk.call("tk", "scaling", 0.9 if compact else 1.0)
        except tk.TclError:
            logging.debug("Tk scaling is not available")

    def apply_editor_font(self) -> None:
        family = self.settings.get("editor", "font_family") or "TkDefaultFont"
        size = int(self.settings.get("editor", "font_size") or 10)
        try:
            default_font = tkfont.nametofont("TkDefaultFont")
            text_font = tkfont.nametofont("TkTextFont")
            default_font.configure(family=family, size=size)
            text_font.configure(family=family, size=size)
        except tk.TclError:
            logging.warning("Не вдалося застосувати шрифт %s", family)

    def on_backup_all(self) -> None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"{APP_NAME}_backup_{timestamp}.zip"
        initialdir = str(self.default_workdir())
        target_path = filedialog.asksaveasfilename(
            title="Зберегти резервну копію",
            defaultextension=".zip",
            initialfile=default_name,
            initialdir=initialdir,
            filetypes=(("ZIP", "*.zip"), ("Усі файли", "*.*")),
        )
        if not target_path:
            return
        try:
            target = backup_all_data(Path(target_path))
            messagebox.showinfo("Резервна копія", f"Створено: {target}")
        except Exception:
            logging.exception("Full backup failed")
            show_error("Резервна копія", "Не вдалося створити копію даних.")

    def on_restore_all(self) -> None:
        archive_path = filedialog.askopenfilename(
            title="Відновити з резервної копії",
            initialdir=str(self.default_workdir()),
            filetypes=(("ZIP", "*.zip"), ("Усі файли", "*.*")),
        )
        if not archive_path:
            return
        if not messagebox.askyesno(
            "Відновлення даних",
            "Відновити всі дані з вибраної копії? Поточні дані буде перезаписано.",
        ):
            return
        try:
            restore_all_data(Path(archive_path))
            messagebox.showinfo(
                "Відновлення даних",
                "Дані відновлено. Перезапустіть додаток, щоб застосувати зміни.",
            )
            self.refresh_all()
        except Exception as exc:
            logging.exception("Restore failed")
            show_error("Відновлення даних", str(exc))

    def on_backup(self) -> None:
        try:
            target = backup_database(get_db_path())
            messagebox.showinfo("Резервна копія", f"Створено: {target}")
        except Exception as exc:
            logging.exception("Backup failed")
            show_error("Резервна копія", str(exc))

    # Brands
    def create_brands_tab(self) -> None:
        columns = [("name", "Назва", 300)]
        self.brand_table = TableFrame(self.brands_frame, columns)
        self.brand_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.brand_table.on_double_click(self.edit_brand)
        self.brand_table.register_context_menu(self.edit_brand, self.delete_brand)

        btns = ttk.Frame(self.brands_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Додати", command=self.add_brand).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_brand).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_brand).pack(side=tk.LEFT, padx=4)

    def add_brand(self) -> None:
        values = simple_prompt("Новий бренд", ["Назва бренду"])
        if not values:
            return
        try:
            db.add_brand(values[0])
            self.refresh_brands()
        except sqlite3.IntegrityError:
            show_error("Бренди", "Бренд з такою назвою вже існує.")
        except Exception:
            logging.exception("Add brand error")
            show_error("Бренди", "Не вдалося додати бренд.")

    def edit_brand(self) -> None:
        brand_id = self.brand_table.selected_id()
        if not brand_id:
            show_error("Бренди", "Оберіть бренд для редагування.")
            return
        rows = [b for b in db.list_brands() if b["id"] == brand_id]
        values = simple_prompt("Редагувати бренд", ["Назва бренду"], [rows[0]["name"]] if rows else None)
        if not values:
            return
        try:
            db.update_brand(brand_id, values[0])
            self.refresh_brands()
            self.refresh_products()
        except sqlite3.IntegrityError:
            show_error("Бренди", "Бренд з такою назвою вже існує.")
        except Exception:
            logging.exception("Edit brand error")
            show_error("Бренди", "Не вдалося змінити бренд.")

    def delete_brand(self) -> None:
        brand_id = self.brand_table.selected_id()
        if not brand_id:
            show_error("Бренди", "Оберіть бренд для видалення.")
            return
        if not messagebox.askyesno("Підтвердження", "Видалити бренд та пов'язані товари?"):
            return
        try:
            db.delete_brand(brand_id)
            self.refresh_brands()
            self.refresh_products()
        except Exception:
            logging.exception("Delete brand error")
            show_error("Бренди", "Не вдалося видалити бренд.")

    # Categories
    def create_categories_tab(self) -> None:
        top = ttk.Frame(self.categories_frame)
        top.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(top, text="Пошук").pack(side=tk.LEFT)
        self.category_search_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.category_search_var, width=30).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="Знайти", command=self.refresh_categories).pack(side=tk.LEFT)

        btns = ttk.Frame(self.categories_frame)
        btns.pack(fill=tk.X, padx=8)
        ttk.Button(btns, text="Коренева", command=lambda: self.add_category(parent_id=None)).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Підкатегорія", command=lambda: self.add_category(parent_id=self.selected_category_id())).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Перейменувати", command=self.edit_category).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Перемістити", command=self.move_category_ui).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Вище", command=lambda: self.bump_category(-1)).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Нижче", command=lambda: self.bump_category(1)).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Службова/Прихована", command=self.toggle_category_flags).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити/Злити", command=self.delete_category).pack(side=tk.LEFT, padx=4)

        columns = ("products", "quantity", "flags")
        self.category_tree = ttk.Treeview(
            self.categories_frame,
            columns=columns,
            show="tree headings",
            selectmode="browse",
        )
        self.category_tree.heading("#0", text="Категорія")
        self.category_tree.heading("products", text="Найменувань")
        self.category_tree.column("products", width=110, anchor="center")
        self.category_tree.heading("quantity", text="Одиниць на складі")
        self.category_tree.column("quantity", width=140, anchor="center")
        self.category_tree.heading("flags", text="Статус")
        self.category_tree.column("flags", width=160, anchor="w")
        self.category_tree.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.category_tree.bind("<Double-1>", lambda e: self.edit_category())
        self.category_tree.bind(
            "<Button-3>",
            lambda e: self.category_tree.event_generate("<<TreeviewSelect>>") or self.category_tree.tk.call(
                self.category_tree, "identify", "row", e.y
            ),
        )

    def selected_category_id(self) -> Optional[int]:
        sel = self.category_tree.selection()
        if not sel:
            return None
        try:
            return int(sel[0])
        except ValueError:
            return None

    def add_category(self, parent_id: Optional[int]) -> None:
        values = category_prompt("Нова категорія", None)
        if not values:
            return
        try:
            name, color, icon, attrs, is_service, is_hidden = values
            db.add_category(name, parent_id, color, icon, attrs, is_service, is_hidden)
            self.refresh_categories()
        except sqlite3.IntegrityError:
            show_error("Категорії", "Категорія з такою назвою вже існує.")
        except Exception:
            logging.exception("Add category error")
            show_error("Категорії", "Не вдалося додати категорію.")

    def edit_category(self) -> None:
        category_id = self.selected_category_id()
        if not category_id:
            show_error("Категорії", "Оберіть категорію для редагування.")
            return
        current = self.categories_index.get(category_id)
        values = category_prompt(
            "Редагувати категорію",
            (
                current.get("name", "") if current else "",
                current.get("color") if current else "",
                current.get("icon") if current else "",
                current.get("typical_attributes") if current else "",
                bool(current.get("is_service")) if current else False,
                bool(current.get("is_hidden")) if current else False,
            ),
        )
        if not values:
            return
        try:
            name, color, icon, attrs, is_service, is_hidden = values
            db.update_category(
                category_id,
                name,
                parent_id=current.get("parent_id") if current else None,
                color=color,
                icon=icon,
                typical_attributes=attrs,
                is_service=is_service,
                is_hidden=is_hidden,
                sort_order=current.get("sort_order") if current else None,
            )
            self.refresh_categories()
            self.refresh_products()
        except sqlite3.IntegrityError:
            show_error("Категорії", "Категорія з такою назвою вже існує.")
        except Exception as exc:
            logging.exception("Edit category error")
            show_error("Категорії", f"Не вдалося змінити категорію: {exc}")

    def bump_category(self, direction: str) -> None:
        category_id = self.selected_category_id()
        if not category_id:
            show_error("Категорії", "Оберіть категорію для переміщення.")
            return
        db.bump_category_order(category_id, direction)
        self.refresh_categories()

    def move_category_ui(self) -> None:
        category_id = self.selected_category_id()
        if not category_id:
            show_error("Категорії", "Оберіть категорію для переміщення.")
            return
        exclude = {category_id}
        exclude.update(db.get_category_descendants(category_id))
        options = [(None, "(Корінь)")]
        for cat in self.flatten_categories():
            if cat["id"] in exclude:
                continue
            options.append((cat["id"], cat["label"]))
        choice = select_category_dialog("Новий батько", options)
        if choice is None:
            return
        try:
            db.move_category(category_id, choice)
            self.refresh_categories()
            self.refresh_products()
        except Exception:
            logging.exception("Move category error")
            show_error("Категорії", "Не вдалося перемістити категорію.")

    def bump_category(self, delta: int) -> None:
        category_id = self.selected_category_id()
        if not category_id:
            show_error("Категорії", "Оберіть категорію для зміни порядку.")
            return
        db.bump_category_order(category_id, delta)
        self.refresh_categories()

    def toggle_category_flags(self) -> None:
        category_id = self.selected_category_id()
        if not category_id:
            show_error("Категорії", "Оберіть категорію.")
            return
        current = self.categories_index.get(category_id)
        if not current:
            return
        try:
            db.update_category(
                category_id,
                current["name"],
                parent_id=current.get("parent_id"),
                color=current.get("color"),
                icon=current.get("icon"),
                typical_attributes=current.get("typical_attributes"),
                is_service=not bool(current.get("is_service")) if messagebox.askyesno(
                    "Статус", "Позначити/зняти позначку 'службова'?"
                )
                else bool(current.get("is_service")),
                is_hidden=not bool(current.get("is_hidden")) if messagebox.askyesno(
                    "Статус", "Приховати/показати категорію?"
                )
                else bool(current.get("is_hidden")),
                sort_order=current.get("sort_order"),
            )
            self.refresh_categories()
        except Exception:
            logging.exception("Toggle flags error")
            show_error("Категорії", "Не вдалося змінити статус.")

    def delete_category(self) -> None:
        category_id = self.selected_category_id()
        if not category_id:
            show_error("Категорії", "Оберіть категорію для видалення.")
            return
        target = None
        if messagebox.askyesno("Злиття", "Злити категорію з іншою (рекомендовано для категорій з товарами)?"):
            options = [(None, "(видалити, якщо порожня)")]
            for cat in self.flatten_categories():
                if cat["id"] == category_id:
                    continue
                options.append((cat["id"], cat["label"]))
            target = select_category_dialog("Цільова категорія", options)
        try:
            db.delete_category(category_id, target)
            self.refresh_categories()
            self.refresh_products()
        except ValueError as exc:
            show_error("Категорії", str(exc))
        except Exception:
            logging.exception("Delete category error")
            show_error("Категорії", f"Не вдалося видалити категорію: {exc}")

    # Products
    def create_products_tab(self) -> None:
        top = ttk.Frame(self.products_frame)
        top.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(top, text="Пошук:").pack(side=tk.LEFT)
        self.product_search_var = tk.StringVar()
        search_entry = ttk.Entry(top, textvariable=self.product_search_var, width=30)
        search_entry.pack(side=tk.LEFT, padx=4)
        search_entry.bind("<Return>", self.on_product_search_enter)
        ttk.Label(top, text="Категорія:").pack(side=tk.LEFT, padx=6)
        self.product_category_filter_var = tk.StringVar()
        self.product_category_filter_combo = ttk.Combobox(
            top, textvariable=self.product_category_filter_var, state="readonly", width=30
        )
        self.product_category_filter_combo.pack(side=tk.LEFT)
        self.include_subcategories_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(top, text="Включно з підкатегоріями", variable=self.include_subcategories_var).pack(
            side=tk.LEFT, padx=6
        )
        ttk.Button(top, text="Оновити", command=self.on_search_products).pack(side=tk.LEFT, padx=6)

        columns = [
            ("sku", "SKU", 120),
            ("barcode", "Штрихкод", 160),
            ("supplier_sku", "Артикул постачальника", 170),
            ("name", "Назва", 230),
            ("brand", "Бренд", 140),
            ("category", "Категорія", 140),
            ("extra_categories", "Додаткові категорії", 200),
            ("unit", "Одиниця", 90),
            ("is_active", "Активний", 90),
        ]
        self.product_table = TableFrame(self.products_frame, columns, selectmode="extended")
        self.product_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.product_table.on_double_click(self.edit_product)
        self.product_table.register_context_menu(self.edit_product, self.delete_product)

        btns = ttk.Frame(self.products_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Додати", command=self.add_product).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_product).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_product).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Імпорт із файлу", command=self.import_products_from_file).pack(side=tk.LEFT, padx=4)
        ttk.Button(
            btns,
            text="Масові дії...",
            command=lambda: open_products_bulk_actions_dialog(self, db.get_connection(), self.product_table),
        ).pack(side=tk.LEFT, padx=4)

    def on_search_products(self) -> None:
        category_id = None
        try:
            selection = self.product_category_combo.current()
            if selection is not None and selection > 0:
                category_id = self.category_choices[selection]["id"]
        except (AttributeError, IndexError):
            category_id = None
        self.refresh_products(
            self.product_search_var.get(),
            category_id=category_id,
            include_subtree=bool(self.include_subcategories_var.get()),
        )

    def on_product_search_enter(self, _event=None) -> None:
        value = self.product_search_var.get().strip()
        prefix = _sanitize_barcode_prefix(self.settings.get("defaults", "product", "barcode_prefix") or "")
        product = db.find_product_by_scan_code(value, prefix)
        if product:
            self.product_search_var.set("")
            self.product_category_filter_var.set("Усі категорії")
            self.refresh_products()
            pid = str(product["id"])
            self.product_table.tree.selection_set(pid)
            self.product_table.tree.focus(pid)
            self.product_table.tree.see(pid)
            return
        self.on_search_products()

    def add_product(self) -> None:
        brands = db.list_brands()
        categories = self.flatten_categories()
        values = product_prompt(brands, categories, "Новий товар", settings=self.settings)
        if not values:
            return
        sku, supplier_sku_legacy, name, brand_id, category_id, unit, is_active, extras, supplier_codes, barcodes = values
        try:
            effective_supplier_sku = supplier_sku_legacy or (
                supplier_codes[0]["supplier_sku"] if len(supplier_codes) == 1 else None
            )
            product_id = db.add_product(sku, name, brand_id, category_id, unit, is_active, effective_supplier_sku)
            db.set_product_categories(product_id, category_id, extras)
            if supplier_codes:
                db.replace_product_supplier_codes(product_id, supplier_codes)
            if barcodes:
                db.replace_product_barcodes(product_id, barcodes)
            self.refresh_products()
        except sqlite3.IntegrityError as exc:
            if "ProductSupplierCodes" in str(exc):
                show_error(
                    "Товари",
                    "Артикул постачальника вже прив’язаний до іншого товару для цього постачальника.",
                )
            elif "ProductBarcodes" in str(exc) or "UNIQUE" in str(exc):
                show_error("Товари", "Цей штрихкод вже прив’язаний до іншого товару.")
            else:
                show_error("Товари", "SKU або назва вже існує.")
        except Exception:
            logging.exception("Add product error")
            show_error("Товари", "Не вдалося додати товар.")

    def edit_product(self) -> None:
        product_id = self.product_table.selected_id()
        if not product_id:
            show_error("Товари", "Оберіть товар для редагування.")
            return
        product = db.get_product(product_id)
        if not product:
            return
        brands = db.list_brands()
        categories = self.flatten_categories()
        supplier_codes = [
            {
                "supplier_id": row["supplier_id"],
                "supplier_sku": row["supplier_sku"],
                "is_primary": bool(row["is_primary"]),
            }
            for row in db.list_product_supplier_codes(product_id)
        ]
        barcodes = [
            {"code": row["code"], "note": row["note"] or ""}
            for row in db.list_product_barcodes(product_id)
        ]
        values = product_prompt(
            brands,
            categories,
            "Редагувати товар",
            (
                product["sku"],
                product.get("supplier_sku"),
                product["name"],
                product["brand_id"],
                product["category_id"],
                product["unit"],
                bool(product["is_active"]),
                db.get_product_additional_categories(product_id),
                supplier_codes,
                barcodes,
            ),
            settings=self.settings,
        )
        if not values:
            return
        sku, supplier_sku_legacy, name, brand_id, category_id, unit, is_active, extras, supplier_codes, barcodes = values
        try:
            db.update_product(product_id, sku, name, brand_id, category_id, unit, is_active, supplier_sku_legacy)
            db.set_product_categories(product_id, category_id, extras)
            db.replace_product_supplier_codes(product_id, supplier_codes)
            db.replace_product_barcodes(product_id, barcodes)
            self.refresh_products()
        except sqlite3.IntegrityError as exc:
            if "ProductSupplierCodes" in str(exc):
                show_error(
                    "Товари",
                    "Артикул постачальника вже прив’язаний до іншого товару для цього постачальника.",
                )
            elif "ProductBarcodes" in str(exc) or "UNIQUE" in str(exc):
                show_error("Товари", "Цей штрихкод вже прив’язаний до іншого товару.")
            else:
                show_error("Товари", "SKU або назва вже існує.")
        except Exception:
            logging.exception("Edit product error")
            show_error("Товари", "Не вдалося змінити товар.")

    def delete_product(self) -> None:
        product_id = self.product_table.selected_id()
        if not product_id:
            show_error("Товари", "Оберіть товар для видалення.")
            return
        if not messagebox.askyesno("Підтвердження", "Видалити товар?"):
            return
        try:
            db.delete_product(product_id)
            self.refresh_products()
        except Exception:
            logging.exception("Delete product error")
            show_error("Товари", "Не вдалося видалити товар.")

    def import_products_from_file(self) -> None:
        file_path = filedialog.askopenfilename(
            title="Файл товарів",
            filetypes=[("CSV", "*.csv"), ("Excel", "*.xlsx *.xls"), ("Усі файли", "*.*")],
            initialdir=str(self.default_workdir()),
        )
        if not file_path:
            return

        try:
            raw_rows, headers = parse_import_file(Path(file_path), encoding=self.settings.get("files", "encoding") or "utf-8")
        except Exception as exc:
            logging.exception("Не вдалося прочитати файл імпорту товарів")
            show_error(
                "Імпорт товарів",
                "Не вдалося прочитати файл. Перевірте формат, кодування та структуру даних.\n" + str(exc),
            )
            return

        if not raw_rows:
            messagebox.showinfo("Імпорт товарів", "У файлі не знайдено рядків із товарами.")
            return

        dialog = ProductsImportDialog(self, raw_rows, headers, settings=self.settings)
        result = dialog.result
        if not result:
            return

        try:
            backup_path = backup_database(get_db_path())
        except Exception:
            logging.exception("Не вдалося створити резервну копію БД перед імпортом товарів")
            show_error("Імпорт товарів", "Не вдалося створити резервну копію БД.")
            return

        try:
            summary = self._process_product_import(result["rows"], result["options"])
            summary = f"Резервну копію створено за шляхом:\n{backup_path}\n\n" + summary
        except Exception:
            logging.exception("Помилка під час імпорту товарів")
            show_error("Імпорт товарів", "Імпорт перервано помилкою. Деталі у логах.")
            return

        messagebox.showinfo("Імпорт товарів", summary)
        self.refresh_products()

    def _process_product_import(self, rows: list[dict], options: dict) -> str:
        mode = options.get("mode", "create")
        create_missing = bool(options.get("create_missing", True))
        extra_mode = options.get("extra_categories_mode", "none")
        update_name = bool(options.get("update_name", True))

        default_brand = self.settings.get("defaults", "product", "brand") or "Імпорт"
        default_category = self.settings.get("defaults", "product", "category") or "Імпорт"
        default_unit = (self.settings.get("defaults", "product", "unit") or "pcs").strip() or "pcs"

        created = 0
        updated = 0
        skipped = 0
        errors: list[str] = []

        conn = db.get_connection()
        try:
            with db.safe_transaction(conn):
                product_rows = list(
                    conn.execute(
                        "SELECT id, sku, name, supplier_sku, brand_id, category_id, unit, is_active FROM Products"
                    )
                )
                products_by_sku = {r["sku"].lower(): dict(r) for r in product_rows if r["sku"]}
                products_by_name = {r["name"].lower(): dict(r) for r in product_rows if r["name"]}

                brands = list(conn.execute("SELECT id, name FROM Brands"))
                categories = list(conn.execute("SELECT id, name FROM Categories"))
                brands_by_id = {int(b["id"]): b for b in brands}
                categories_by_id = {int(c["id"]): c for c in categories}
                brands_by_name = {b["name"].lower(): b for b in brands}
                categories_by_name = {c["name"].lower(): c for c in categories}

                def next_sort_order(parent_id=None):
                    row = conn.execute(
                        "SELECT COALESCE(MAX(sort_order),0) FROM Categories WHERE parent_id IS ?", (parent_id,)
                    ).fetchone()
                    return int(row[0]) + 1

                def resolve_brand(name: str, brand_id: Optional[int], row_idx: int) -> Optional[int]:
                    if brand_id:
                        if brand_id in brands_by_id:
                            return brand_id
                        errors.append(f"Рядок {row_idx}: ID бренду {brand_id} не знайдено")
                        return None
                    clean = (name or "").strip()
                    if not clean:
                        return None
                    existing = brands_by_name.get(clean.lower())
                    if existing:
                        return int(existing["id"])
                    if not create_missing:
                        errors.append(f"Рядок {row_idx}: Бренд '{clean}' не знайдено")
                        return None
                    cur = conn.execute("INSERT INTO Brands (name) VALUES (?)", (clean,))
                    brand_id = cur.lastrowid
                    brand_row = {"id": brand_id, "name": clean}
                    brands_by_id[int(brand_id)] = brand_row
                    brands_by_name[clean.lower()] = brand_row
                    return int(brand_id)

                def resolve_category(name: str, category_id: Optional[int], row_idx: int) -> Optional[int]:
                    if category_id:
                        if category_id in categories_by_id:
                            return category_id
                        errors.append(f"Рядок {row_idx}: Категорію з ID {category_id} не знайдено")
                        return None
                    clean = (name or "").strip()
                    if not clean:
                        return None
                    existing = categories_by_name.get(clean.lower())
                    if existing:
                        return int(existing["id"])
                    if not create_missing:
                        errors.append(f"Рядок {row_idx}: Категорію '{clean}' не знайдено")
                        return None
                    sort_order = next_sort_order(None)
                    cur = conn.execute(
                        "INSERT INTO Categories (name, parent_id, sort_order, is_service, is_hidden) VALUES (?, NULL, ?, 0, 0)",
                        (clean, sort_order),
                    )
                    cat_id = cur.lastrowid
                    cat_row = {"id": cat_id, "name": clean}
                    categories_by_id[int(cat_id)] = cat_row
                    categories_by_name[clean.lower()] = cat_row
                    return int(cat_id)

                for idx, row in enumerate(rows, start=1):
                    sku = (row.get("sku") or "").strip()
                    name = (row.get("name") or "").strip()
                    supplier_sku = (row.get("supplier_sku") or "").strip()
                    unit = (row.get("unit") or "").strip()
                    is_active = row.get("is_active")
                    brand_val = (row.get("brand") or "").strip()
                    category_val = (row.get("category") or "").strip()
                    extra_categories = list(row.get("extra_categories") or [])

                    matched_by = None
                    product = None
                    if sku:
                        product = products_by_sku.get(sku.lower())
                        matched_by = "sku" if product else None
                    if not product and mode in {"update", "upsert"} and not sku and name:
                        product = products_by_name.get(name.lower())
                        matched_by = "name" if product else None

                    if mode == "create" and product:
                        errors.append(f"Рядок {idx}: SKU або назва вже існує")
                        skipped += 1
                        continue

                    if mode == "update" and not product:
                        skipped += 1
                        continue

                    brand_id = resolve_brand(brand_val, row.get("brand_id"), idx)
                    category_id = resolve_category(category_val, row.get("category_id"), idx)

                    if mode in {"create", "upsert"} and not product:
                        final_sku = sku.strip()
                        final_name = name.strip()
                        if not final_sku or not final_name:
                            errors.append(f"Рядок {idx}: Потрібні SKU і назва для створення")
                            skipped += 1
                            continue
                        if final_sku.lower() in products_by_sku:
                            errors.append(f"Рядок {idx}: SKU '{final_sku}' вже існує")
                            skipped += 1
                            continue
                        if final_name.lower() in products_by_name:
                            errors.append(f"Рядок {idx}: Назва '{final_name}' вже існує")
                            skipped += 1
                            continue

                        if not brand_id:
                            brand_id = resolve_brand(default_brand, None, idx) or None
                        if not category_id:
                            category_id = resolve_category(default_category, None, idx)
                        if not category_id or not brand_id:
                            skipped += 1
                            continue

                        try:
                            cur = conn.execute(
                                "INSERT INTO Products (sku, supplier_sku, name, brand_id, category_id, unit, is_active) VALUES (?, ?, ?, ?, ?, ?, ?)",
                                (
                                    final_sku,
                                    supplier_sku or None,
                                    final_name,
                                    brand_id,
                                    category_id,
                                    unit or default_unit,
                                    1 if (is_active is None or is_active) else 0,
                                ),
                            )
                        except sqlite3.IntegrityError as exc:
                            errors.append(f"Рядок {idx}: Конфлікт унікальності ({exc})")
                            skipped += 1
                            continue

                        product_id = int(cur.lastrowid)
                        product_row = {
                            "id": product_id,
                            "sku": final_sku,
                            "name": final_name,
                            "supplier_sku": supplier_sku,
                            "brand_id": brand_id,
                            "category_id": category_id,
                            "unit": unit or default_unit,
                            "is_active": 1 if (is_active is None or is_active) else 0,
                        }
                        products_by_sku[final_sku.lower()] = product_row
                        products_by_name[final_name.lower()] = product_row
                        product = product_row
                        created += 1
                    elif product:
                        old_name = product.get("name")
                        updates: dict[str, object] = {}
                        if supplier_sku:
                            updates["supplier_sku"] = supplier_sku
                        if unit:
                            updates["unit"] = unit
                        if brand_id:
                            updates["brand_id"] = brand_id
                        if category_id:
                            updates["category_id"] = category_id
                        if is_active is not None:
                            updates["is_active"] = 1 if is_active else 0

                        can_rename = matched_by == "sku" and update_name and name
                        if can_rename:
                            name_exists = products_by_name.get(name.lower())
                            if name_exists and int(name_exists.get("id")) != int(product["id"]):
                                errors.append(f"Рядок {idx}: Назва '{name}' вже використовується")
                                skipped += 1
                                continue
                            updates["name"] = name

                        if updates:
                            set_clause = ", ".join(f"{k}=?" for k in updates.keys())
                            conn.execute(
                                f"UPDATE Products SET {set_clause} WHERE id=?",
                                (*updates.values(), product["id"]),
                            )
                            product.update(updates)
                            if updates.get("name"):
                                if old_name:
                                    products_by_name.pop(old_name.lower(), None)
                                products_by_name[name.lower()] = product
                            updated += 1

                    if not product:
                        continue

                    if extra_mode != "none":
                        if extra_mode == "replace":
                            conn.execute("DELETE FROM ProductCategoryLinks WHERE product_id=?", (product["id"],))
                        if extra_categories:
                            for cat_name in extra_categories:
                                cat_id = resolve_category(cat_name, None, idx) if cat_name else None
                                if not cat_id or cat_id == product.get("category_id"):
                                    continue
                                conn.execute(
                                    "INSERT OR IGNORE INTO ProductCategoryLinks (product_id, category_id) VALUES (?, ?)",
                                    (product["id"], cat_id),
                                )
                    if extra_mode != "none" and extra_mode == "replace" and not extra_categories:
                        # Already cleared links above.
                        pass
        finally:
            conn.close()

        summary = [
            f"Режим: {mode}",
            f"Створено: {created}",
            f"Оновлено: {updated}",
            f"Пропущено: {skipped}",
            f"Помилок: {len(errors)}",
        ]
        if errors:
            summary.append("Перші помилки:\n" + "\n".join(errors[:10]))
        return "\n".join(summary)

    # Counterparties
    def create_counterparties_tab(self) -> None:
        columns = [
            ("name", "Назва", 200),
            ("type", "Тип", 120),
            ("phone", "Телефон", 120),
            ("email", "Email", 170),
            ("address", "Адреса", 200),
            ("note", "Нотатка", 200),
        ]

        def make_table(parent: tk.Widget) -> TableFrame:
            table = TableFrame(parent, columns)
            table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
            table.on_double_click(self.edit_counterparty)
            table.register_context_menu(self.edit_counterparty, self.delete_counterparty)
            return table

        self.counterparty_notebook = ttk.Notebook(self.counterparties_frame)
        self.counterparty_notebook.pack(fill=tk.BOTH, expand=True)

        supplier_tab = ttk.Frame(self.counterparty_notebook)
        customer_tab = ttk.Frame(self.counterparty_notebook)
        all_tab = ttk.Frame(self.counterparty_notebook)

        self.counterparty_notebook.add(supplier_tab, text="Постачальники")
        self.counterparty_notebook.add(customer_tab, text="Покупці")
        self.counterparty_notebook.add(all_tab, text="Всі/Інші")

        self.counterparty_tables = {
            "suppliers": make_table(supplier_tab),
            "customers": make_table(customer_tab),
            "all": make_table(all_tab),
        }
        self.counterparty_tab_frames = {
            "suppliers": supplier_tab,
            "customers": customer_tab,
            "all": all_tab,
        }

        btns = ttk.Frame(self.counterparties_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Додати", command=self.add_counterparty).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_counterparty).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_counterparty).pack(side=tk.LEFT, padx=4)

    def get_active_counterparty_selection(self) -> tuple[str, TableFrame | None, int | None]:
        if not hasattr(self, "counterparty_notebook"):
            return "suppliers", None, None
        current_tab = self.counterparty_notebook.select()
        active_key = "suppliers"
        for key, frame in self.counterparty_tab_frames.items():
            if str(frame) == current_tab:
                active_key = key
                break
        table = self.counterparty_tables.get(active_key)
        selected_id = table.selected_id() if table else None
        return active_key, table, selected_id

    def add_counterparty(self) -> None:
        active_key, _, _ = self.get_active_counterparty_selection()
        default_types = {"suppliers": "supplier", "customers": "customer", "all": "other"}
        values = counterparty_prompt(default_type=default_types.get(active_key))
        if not values:
            return
        try:
            db.add_counterparty(*values)
            self.refresh_counterparties()
        except sqlite3.IntegrityError:
            show_error("Контрагенти", "Контрагент з такою назвою вже існує.")
        except Exception:
            logging.exception("Add counterparty error")
            show_error("Контрагенти", "Не вдалося додати контрагента.")

    def edit_counterparty(self) -> None:
        _, _, counterparty_id = self.get_active_counterparty_selection()
        if not counterparty_id:
            show_error("Контрагенти", "Оберіть контрагента.")
            return
        rows = [c for c in db.list_counterparties() if c["id"] == counterparty_id]
        if not rows:
            return
        c = rows[0]
        values = counterparty_prompt(
            (c["name"], c["type"], c["phone"], c["email"], c["address"], c["note"])
        )
        if not values:
            return
        try:
            db.update_counterparty(counterparty_id, *values)
            self.refresh_counterparties()
        except sqlite3.IntegrityError:
            show_error("Контрагенти", "Контрагент з такою назвою вже існує.")
        except Exception:
            logging.exception("Edit counterparty error")
            show_error("Контрагенти", "Не вдалося змінити контрагента.")

    def delete_counterparty(self) -> None:
        _, _, counterparty_id = self.get_active_counterparty_selection()
        if not counterparty_id:
            show_error("Контрагенти", "Оберіть контрагента для видалення.")
            return
        if not messagebox.askyesno("Підтвердження", "Видалити контрагента?"):
            return
        try:
            db.delete_counterparty(counterparty_id)
            self.refresh_counterparties()
        except ValueError as exc:
            show_error("Контрагенти", str(exc))
        except Exception:
            logging.exception("Delete counterparty error")
            show_error("Контрагенти", "Не вдалося видалити контрагента.")

    # Warehouses
    def create_warehouses_tab(self) -> None:
        columns = [("name", "Назва", 200), ("description", "Опис", 260), ("is_active", "Активний", 100)]
        self.warehouse_table = TableFrame(self.warehouses_frame, columns)
        self.warehouse_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.warehouse_table.on_double_click(self.edit_warehouse)
        self.warehouse_table.register_context_menu(self.edit_warehouse, self.delete_warehouse)
        btns = ttk.Frame(self.warehouses_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Додати", command=self.add_warehouse).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_warehouse).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_warehouse).pack(side=tk.LEFT, padx=4)

    def add_warehouse(self) -> None:
        values = warehouse_prompt()
        if not values:
            return
        name, description, is_active = values
        try:
            db.add_warehouse(name, description, is_active)
            self.refresh_warehouses()
        except sqlite3.IntegrityError:
            show_error("Склади", "Склад з такою назвою вже існує.")
        except Exception:
            logging.exception("Add warehouse error")
            show_error("Склади", "Не вдалося додати склад.")

    def edit_warehouse(self) -> None:
        warehouse_id = self.warehouse_table.selected_id()
        if not warehouse_id:
            show_error("Склади", "Оберіть склад.")
            return
        rows = [w for w in db.list_warehouses() if w["id"] == warehouse_id]
        if not rows:
            return
        w = rows[0]
        values = warehouse_prompt((w["name"], w["description"] or "", bool(w["is_active"])))
        if not values:
            return
        name, description, is_active = values
        try:
            db.update_warehouse(warehouse_id, name, description, is_active)
            self.refresh_warehouses()
        except sqlite3.IntegrityError:
            show_error("Склади", "Склад з такою назвою вже існує.")
        except Exception:
            logging.exception("Edit warehouse error")
            show_error("Склади", "Не вдалося змінити склад.")

    def delete_warehouse(self) -> None:
        warehouse_id = self.warehouse_table.selected_id()
        if not warehouse_id:
            show_error("Склади", "Оберіть склад для видалення.")
            return
        if not messagebox.askyesno("Підтвердження", "Видалити склад?"):
            return
        try:
            db.delete_warehouse(warehouse_id)
            self.refresh_warehouses()
        except Exception as exc:
            logging.exception("Delete warehouse error")
            show_error("Склади", str(exc))

    # Channels
    def create_channels_tab(self) -> None:
        columns = [("name", "Назва", 240), ("is_active", "Активний", 100)]
        self.channel_table = TableFrame(self.channels_frame, columns)
        self.channel_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.channel_table.on_double_click(self.edit_channel)
        self.channel_table.register_context_menu(self.edit_channel, self.delete_channel)
        btns = ttk.Frame(self.channels_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Додати", command=self.add_channel).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_channel).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_channel).pack(side=tk.LEFT, padx=4)

    def add_channel(self) -> None:
        values = channel_prompt()
        if not values:
            return
        name, is_active = values
        try:
            db.add_channel(name, is_active)
            self.refresh_channels()
        except sqlite3.IntegrityError:
            show_error("Канали", "Канал з такою назвою вже існує.")
        except Exception:
            logging.exception("Add channel error")
            show_error("Канали", "Не вдалося додати канал.")

    def edit_channel(self) -> None:
        channel_id = self.channel_table.selected_id()
        if not channel_id:
            show_error("Канали", "Оберіть канал.")
            return
        rows = [c for c in db.list_channels() if c["id"] == channel_id]
        if not rows:
            return
        c = rows[0]
        values = channel_prompt((c["name"], bool(c["is_active"])))
        if not values:
            return
        name, is_active = values
        try:
            db.update_channel(channel_id, name, is_active)
            self.refresh_channels()
        except sqlite3.IntegrityError:
            show_error("Канали", "Канал з такою назвою вже існує.")
        except Exception:
            logging.exception("Edit channel error")
            show_error("Канали", "Не вдалося змінити канал.")

    def delete_channel(self) -> None:
        channel_id = self.channel_table.selected_id()
        if not channel_id:
            show_error("Канали", "Оберіть канал для видалення.")
            return
        if not messagebox.askyesno("Підтвердження", "Видалити канал?"):
            return
        try:
            db.delete_channel(channel_id)
            self.refresh_channels()
        except Exception as exc:
            logging.exception("Delete channel error")
            show_error("Канали", str(exc))

    # Currencies
    def create_currencies_tab(self) -> None:
        top = ttk.Frame(self.currencies_frame)
        top.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        ttk.Label(top, text="Довідник валют").pack(anchor="w")
        self.base_currency_label = ttk.Label(top, text=self._format_base_currency_label())
        self.base_currency_label.pack(anchor="w", pady=(0, 4))
        curr_columns = [("code", "Код", 80), ("name", "Назва", 200), ("decimals", "Знаків", 60), ("is_active", "Активна", 80)]
        self.currency_table = TableFrame(top, curr_columns, height=6)
        self.currency_table.pack(fill=tk.X, pady=4)
        self.currency_table.on_double_click(self.edit_currency)
        self.currency_table.register_context_menu(self.edit_currency, self.delete_currency)

        btns = ttk.Frame(top)
        btns.pack(pady=4, anchor="w")
        ttk.Button(btns, text="Додати валюту", command=self.add_currency).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_currency).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_currency).pack(side=tk.LEFT, padx=4)

        ttk.Label(top, text="Курси валют").pack(anchor="w", pady=(10, 0))
        rate_columns = [("rate_date", "Дата", 120), ("rate", "Курс до базової", 160)]
        self.rate_table = TableFrame(top, rate_columns, height=6)
        self.rate_table.pack(fill=tk.X, pady=4)

        rate_btns = ttk.Frame(top)
        rate_btns.pack(pady=4, anchor="w")
        ttk.Button(rate_btns, text="Додати курс", command=self.add_rate).pack(side=tk.LEFT, padx=4)
        ttk.Button(rate_btns, text="Змінити", command=self.edit_rate).pack(side=tk.LEFT, padx=4)
        ttk.Button(rate_btns, text="Видалити", command=self.delete_rate).pack(side=tk.LEFT, padx=4)

        self.currency_table.on_select(self.refresh_rates)
        self.rate_table.on_double_click(self.edit_rate)
        self.rate_table.register_context_menu(self.edit_rate, self.delete_rate)

    def add_currency(self) -> None:
        values = simple_prompt("Нова валюта", ["Код", "Назва", "Знаків після коми"], ["USD", "Долар США", "2"])
        if not values:
            return
        try:
            decimals = int(values[2]) if len(values) > 2 else 2
            db.add_currency(values[0], values[1], decimals)
            self.refresh_currencies()
        except sqlite3.IntegrityError:
            show_error("Валюти", "Валюта з таким кодом вже існує")
        except Exception as exc:
            logging.exception("Add currency error")
            show_error("Валюти", str(exc))

    def edit_currency(self) -> None:
        code = self.currency_table.selected_id()
        if not code:
            show_error("Валюти", "Оберіть валюту")
            return
        rows = [c for c in db.list_currencies(active_only=False) if c["code"] == code]
        if not rows:
            return
        cur = rows[0]
        values = simple_prompt("Змінити валюту", ["Код", "Назва", "Знаків після коми", "Активна (1/0)"], [cur["code"], cur["name"], str(cur["decimals"]), str(cur["is_active"]),])
        if not values:
            return
        try:
            decimals = int(values[2]) if len(values) > 2 else 2
            is_active = values[3].strip() != "0" if len(values) > 3 else True
            db.update_currency(values[0], values[1], decimals, is_active)
            self.refresh_currencies()
        except Exception as exc:
            logging.exception("Edit currency error")
            show_error("Валюти", str(exc))

    def delete_currency(self) -> None:
        code = self.currency_table.selected_id()
        if not code:
            show_error("Валюти", "Оберіть валюту")
            return
        if not messagebox.askyesno("Валюти", "Видалити валюту?"):
            return
        try:
            db.delete_currency(code)
            self.refresh_currencies()
        except Exception as exc:
            logging.exception("Delete currency error")
            show_error("Валюти", str(exc))

    def refresh_currencies(self) -> None:
        rows = db.list_currencies(active_only=False)
        self.currency_table.set_rows(
            [
                {
                    "id": row["code"],
                    "code": row["code"],
                    "name": row["name"],
                    "decimals": row["decimals"],
                    "is_active": "Так" if row["is_active"] else "Ні",
                }
                for row in rows
            ]
        )
        self.update_base_currency_label()
        self.refresh_rates()

    def _format_base_currency_label(self) -> str:
        return (
            f"Базова валюта: {get_base_currency_code()} — "
            f"{get_base_currency_name()} ({get_base_currency_decimals()} знаків)"
        )

    def update_base_currency_label(self) -> None:
        if hasattr(self, "base_currency_label"):
            self.base_currency_label.configure(text=self._format_base_currency_label())

    def refresh_rates(self) -> None:
        code = self.currency_table.selected_id()
        code = code or (db.list_currencies(active_only=True)[0]["code"] if db.list_currencies(active_only=True) else None)
        if not code:
            self.rate_table.set_rows([])
            return
        rates = db.list_currency_rates(code)
        self.rate_table.set_rows(
            [
                {
                    "id": r["id"],
                    "rate_date": r["rate_date"],
                    "rate": f"{r['rate']:.4f}",
                }
                for r in rates
            ]
        )

    def add_rate(self) -> None:
        code = self.currency_table.selected_id()
        if not code:
            show_error("Курси", "Оберіть валюту")
            return
        defaults = [datetime.now().strftime("%Y-%m-%d"), "1"]
        values = simple_prompt("Новий курс", ["Дата", "Курс до базової валюти"], defaults)
        if not values:
            return
        try:
            rate = float(values[1])
            db.add_currency_rate(code, values[0], rate)
            self.refresh_rates()
        except Exception as exc:
            logging.exception("Add rate error")
            show_error("Курси", str(exc))

    def edit_rate(self) -> None:
        rate_id = self.rate_table.selected_id()
        if not rate_id:
            show_error("Курси", "Оберіть курс")
            return
        rates = [r for r in db.list_currency_rates(self.currency_table.selected_id()) if r["id"] == rate_id]
        if not rates:
            return
        current = rates[0]
        values = simple_prompt(
            "Змінити курс",
            ["Дата", "Курс до базової валюти"],
            [current["rate_date"], f"{current['rate']:.4f}"],
        )
        if not values:
            return
        try:
            rate = float(values[1])
            db.update_currency_rate(rate_id, values[0], rate)
            self.refresh_rates()
        except Exception as exc:
            logging.exception("Edit rate error")
            show_error("Курси", str(exc))

    def delete_rate(self) -> None:
        rate_id = self.rate_table.selected_id()
        if not rate_id:
            show_error("Курси", "Оберіть курс")
            return
        if not messagebox.askyesno("Курси", "Видалити курс?"):
            return
        try:
            db.delete_currency_rate(rate_id)
            self.refresh_rates()
        except Exception as exc:
            logging.exception("Delete rate error")
            show_error("Курси", str(exc))

    # Purchases
    def create_purchases_tab(self) -> None:
        filters = ttk.Frame(self.purchases_frame)
        filters.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(filters, text="Статус:").pack(side=tk.LEFT)
        self.purchase_status_var = tk.StringVar(value="Усі")
        status_combo = ttk.Combobox(filters, textvariable=self.purchase_status_var, values=["Усі", "Чернетка", "Проведений"], state="readonly", width=14)
        status_combo.pack(side=tk.LEFT, padx=4)
        ttk.Label(filters, text="Дата з:").pack(side=tk.LEFT)
        self.purchase_date_from_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.purchase_date_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(filters, text="по:").pack(side=tk.LEFT)
        self.purchase_date_to_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.purchase_date_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(filters, text="Фільтр", command=self.refresh_purchases).pack(side=tk.LEFT, padx=6)

        columns = [
            ("doc_date", "Дата", 90),
            ("supplier", "Постачальник", 200),
            ("warehouse", "Склад", 160),
            ("status", "Статус", 90),
            ("currency", "Валюта", 80),
            ("rate", "Курс", 80),
            ("total_doc", "Сума (вал)", 110),
            ("total", "Сума (база)", 110),
            ("total_extra", "Супутні (база)", 120),
            ("comment", "Коментар", 240),
        ]
        self.purchase_table = TableFrame(self.purchases_frame, columns)
        self.purchase_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.purchase_table.on_double_click(self.edit_purchase)
        self.purchase_table.register_context_menu_actions(
            [
                ("Редагувати", self.edit_purchase),
                ("Видалити", self.delete_purchase),
                ("---", None),
                ("Друк етикеток…", self.print_purchase_labels),
            ]
        )

        btns = ttk.Frame(self.purchases_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Нова закупівля", command=self.new_purchase).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_purchase).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_purchase).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Провести", command=self.post_purchase_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Відмінити проведення", command=self.unpost_purchase_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Імпорт із файлу", command=self.import_purchases_from_file).pack(side=tk.LEFT, padx=4)

    def _selected_purchase(self):
        doc_id = self.purchase_table.selected_id()
        if not doc_id:
            show_error("Закупівлі", "Оберіть документ")
            return None
        return doc_id

    def refresh_purchases(self) -> None:
        status_filter = self.purchase_status_var.get()
        status_value = "draft" if status_filter == "Чернетка" else "posted" if status_filter == "Проведений" else None
        rows = db.list_purchases(status_value, self.purchase_date_from_var.get().strip() or None, self.purchase_date_to_var.get().strip() or None)
        self.purchase_table.set_rows(
            [
                {
                    "id": r["id"],
                    "doc_date": r["doc_date"],
                    "supplier": r["supplier"] or "-",
                    "warehouse": r["warehouse"] or "-",
                    "status": "Чернетка" if r["status"] == "draft" else "Проведений",
                    "currency": r["currency_code"],
                    "rate": f"{r['exchange_rate']:.4f}",
                    "total_doc": f"{r['total_doc']:.2f}",
                    "total": f"{r['total']:.2f}",
                    "total_extra": f"{r['total_extra_base']:.2f}",
                    "comment": r["comment"] or "",
                }
                for r in rows
            ]
        )

    def new_purchase(self) -> None:
        products = db.list_products()
        warehouses = db.list_warehouses(active_only=True)
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        result = document_prompt("purchase", products, counterparties, warehouses, [], currencies, settings=self.settings)
        if not result:
            return
        info, lines = result
        try:
            doc_id = db.create_purchase(
                info["doc_date"],
                info["counterparty_id"],
                info["warehouse_id"],
                "",
                info["comment"],
                info["currency"],
                info["rate"],
            )
            db.replace_purchase_lines(doc_id, lines, info["rate"])
            self.refresh_purchases()
        except Exception:
            logging.exception("Create purchase error")
            show_error("Закупівлі", "Не вдалося створити документ")

    def edit_purchase(self) -> None:
        doc_id = self._selected_purchase()
        if not doc_id:
            return
        doc = db.get_purchase(doc_id)
        if not doc:
            return
        lines = db.list_purchase_lines(doc_id)
        products = db.list_products()
        warehouses = db.list_warehouses(active_only=False)
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        result = document_prompt(
            "purchase",
            products,
            counterparties,
            warehouses,
            [],
            currencies,
            doc=doc,
            lines=lines,
            settings=self.settings,
        )
        if not result:
            return
        info, new_lines = result
        try:
            if doc["status"] == "draft":
                db.update_purchase(
                    doc_id,
                    info["doc_date"],
                    info["counterparty_id"],
                    info["warehouse_id"],
                    "",
                    info["comment"],
                    info["currency"],
                    info["rate"],
                )
                db.replace_purchase_lines(doc_id, new_lines, info["rate"])
            else:
                db.update_purchase(
                    doc_id,
                    doc["doc_date"],
                    doc["supplier_id"],
                    doc["warehouse_id"],
                    "",
                    info["comment"],
                    doc["currency_code"],
                    doc["exchange_rate"],
                )
            self.refresh_purchases()
        except Exception:
            logging.exception("Edit purchase error")
            show_error("Закупівлі", "Не вдалося змінити документ")

    def delete_purchase(self) -> None:
        doc_id = self._selected_purchase()
        if not doc_id:
            return
        if not messagebox.askyesno("Підтвердження", "Видалити документ?"):
            return
        try:
            with db.get_connection() as conn:
                with db.safe_transaction(conn):
                    status = conn.execute("SELECT status FROM PurchaseDocuments WHERE id=?", (doc_id,)).fetchone()
                    if status and status[0] == "posted":
                        raise ValueError("Видаляти можна лише чернетки")
                    conn.execute("DELETE FROM PurchaseDocuments WHERE id=?", (doc_id,))
            self.refresh_purchases()
        except Exception as exc:
            logging.exception("Delete purchase error")
            show_error("Закупівлі", str(exc))

    def print_purchase_labels(self) -> None:
        doc_id = self.purchase_table.selected_id()
        if not doc_id:
            show_error("Закупівлі", "Оберіть документ")
            return
        try:
            lines = db.list_purchase_lines(doc_id)
        except Exception:
            logging.exception("Failed to load purchase lines for labels")
            show_error("Етикетки", "Не вдалося завантажити позиції документа")
            return
        aggregated: dict[int, dict[str, object]] = {}
        for line in lines:
            try:
                pid = int(line["product_id"])
            except (TypeError, ValueError):
                continue
            qty = float(line["quantity"] or 0)
            if pid not in aggregated:
                aggregated[pid] = {
                    "product_id": pid,
                    "sku": line["sku"] or "",
                    "name": line["product_name"] or "",
                    "qty": 0.0,
                }
            aggregated[pid]["qty"] = float(aggregated[pid].get("qty", 0.0) or 0.0) + qty
        items = [item for item in aggregated.values() if float(item.get("qty", 0) or 0) > 0]
        if not items:
            messagebox.showwarning("Етикетки", "У документі немає позицій з кількістю > 0.")
            return
        open_labels_print_dialog(self, items, "Кількість з документа закупівлі")

    def post_purchase_action(self) -> None:
        doc_id = self._selected_purchase()
        if not doc_id:
            return
        try:
            db.post_purchase(doc_id)
            self.refresh_purchases()
            self.refresh_stock()
            self.refresh_cash()
        except Exception as exc:
            logging.exception("Post purchase error")
            show_error("Закупівлі", str(exc))

    def unpost_purchase_action(self) -> None:
        doc_id = self._selected_purchase()
        if not doc_id:
            return
        try:
            db.unpost_purchase(doc_id)
            self.refresh_purchases()
            self.refresh_stock()
            self.refresh_cash()
        except Exception as exc:
            logging.exception("Unpost purchase error")
            show_error("Закупівлі", str(exc))

    def import_purchases_from_file(self) -> None:
        file_path = filedialog.askopenfilename(
            title="Файл закупівель",
            filetypes=[("CSV", "*.csv"), ("Excel", "*.xlsx *.xls"), ("Усі файли", "*.*")],
            initialdir=str(self.default_workdir()),
        )
        if not file_path:
            return

        try:
            raw_rows, headers = parse_import_file(Path(file_path), encoding=self.settings.get("files", "encoding") or "utf-8")
        except Exception as exc:
            logging.exception("Не вдалося прочитати файл імпорту закупівель")
            show_error(
                "Імпорт закупівель",
                "Не вдалося прочитати файл. Перевірте формат, кодування та структуру даних.\n" + str(exc),
            )
            return

        if not raw_rows:
            messagebox.showinfo("Імпорт закупівель", "У файлі не знайдено рядків із товарами.")
            return

        warehouses = db.list_warehouses(active_only=True)
        if not warehouses:
            show_error("Імпорт закупівель", "Спочатку створіть хоча б один склад.")
            return

        counterparties = [dict(c) for c in db.list_counterparties()]
        suppliers = sorted(
            (c for c in counterparties if c["type"] in {"supplier", "both", "other"}),
            key=lambda c: c.get("name", ""),
        )
        if not suppliers:
            show_error("Імпорт закупівель", "Спочатку додайте постачальника у контрагенти.")
            return
        currencies = db.list_currencies()

        dialog = PurchasesImportDialog(self, raw_rows, headers, warehouses, suppliers, currencies, self.settings)
        result = dialog.result
        if not result:
            return

        try:
            summary = self._process_purchase_import(result["orders"], result["options"])
        except Exception:
            logging.exception("Помилка під час імпорту закупівель")
            show_error("Імпорт закупівель", "Імпорт перервано помилкою. Деталі у логах.")
            return

        messagebox.showinfo("Імпорт закупівель", summary)
        self.refresh_purchases()
        self.refresh_cash()
        self.refresh_stock()

    def _process_purchase_import(self, orders: list[dict], options: dict) -> str:
        warehouse_id = options["warehouse_id"]
        mode = options.get("mode", "draft")
        create_products = bool(options.get("create_products", True))
        selected_supplier_id = options.get("supplier_id")
        currency_code = (options.get("currency_code") or get_base_currency_code()).strip().upper()
        exchange_rate = float(options.get("exchange_rate") or 1.0)

        product_rows = db.list_products()
        products_by_sku = {p["sku"].lower(): dict(p) for p in product_rows if p["sku"]}
        products_by_name = {p["name"].lower(): dict(p) for p in product_rows if p["name"]}
        products_by_supplier_sku = {p["supplier_sku"].lower(): dict(p) for p in product_rows if p.get("supplier_sku")}

        default_brand = self.settings.get("defaults", "product", "brand") or "Імпорт"
        default_category = self.settings.get("defaults", "product", "category") or "Імпорт"
        default_unit = (self.settings.get("defaults", "product", "unit") or "pcs").strip() or "pcs"
        brand_id, category_id = db.ensure_import_defaults(default_brand, default_category)

        created_products = 0
        skipped_lines = 0
        posted_docs = 0
        draft_docs = 0
        total_docs = 0

        doc_date = next((r.get("doc_date") for r in orders if r.get("doc_date")), None) or datetime.now().strftime(
            "%Y-%m-%d"
        )
        supplier_id = selected_supplier_id
        if not supplier_id:
            raise ValueError("Не вказано постачальника для імпорту закупівель")
        with db.get_connection() as conn:
            supplier_code_rows = conn.execute(
                """
                SELECT lower(psc.supplier_sku) AS k, p.id, p.sku, p.supplier_sku, p.name, p.brand_id, p.category_id, p.unit, p.is_active
                FROM ProductSupplierCodes psc
                JOIN Products p ON p.id = psc.product_id
                WHERE psc.supplier_id = ?
                """,
                (supplier_id,),
            ).fetchall()
        products_by_supplier_code = {
            row["k"]: {key: row[key] for key in row.keys() if key != "k"} for row in supplier_code_rows
        }
        purchase_lines: list[tuple[int, float, float]] = []
        comments: list[str] = []

        for row in orders:
            comment_val = (row.get("comment") or "").strip()
            if comment_val:
                comments.append(comment_val)

            sku = (row.get("sku") or "").strip()
            supplier_sku = (row.get("supplier_sku") or "").strip()
            name = (row.get("product_name") or sku or "Без назви").strip()
            qty = float(row.get("quantity") or 0)
            price = float(row.get("price") or 0)
            amount = float(row.get("amount") or 0)
            if not price and qty and amount:
                price = amount / qty

            product_row = products_by_supplier_code.get(supplier_sku.lower()) if supplier_sku else None
            if not product_row:
                product_row = products_by_supplier_sku.get(supplier_sku.lower()) if supplier_sku else None
            if not product_row:
                product_row = products_by_sku.get(sku.lower()) if sku else None
            if not product_row and name:
                product_row = products_by_name.get(name.lower())
            if not product_row and create_products:
                final_sku = sku or self._generate_unique_sku(name, set(products_by_sku.keys()))
                product_id = db.add_product(
                    final_sku,
                    name,
                    brand_id,
                    category_id,
                    unit=default_unit,
                    supplier_sku=supplier_sku,
                )
                product_row = {
                    "id": product_id,
                    "sku": final_sku,
                    "name": name,
                    "supplier_sku": supplier_sku,
                    "brand_id": brand_id,
                    "category_id": category_id,
                    "unit": default_unit,
                    "is_active": True,
                }
                if supplier_sku:
                    try:
                        db.replace_product_supplier_codes(
                            product_id,
                            [
                                {
                                    "supplier_id": supplier_id,
                                    "supplier_sku": supplier_sku,
                                    "is_primary": True,
                                }
                            ],
                        )
                        products_by_supplier_code[supplier_sku.lower()] = product_row
                    except sqlite3.IntegrityError as exc:
                        logging.warning("Не вдалося зберегти артикул постачальника для імпорту: %s", exc)
                products_by_sku[final_sku.lower()] = product_row
                if supplier_sku:
                    products_by_supplier_sku[supplier_sku.lower()] = product_row
                products_by_name[name.lower()] = product_row
                created_products += 1

            if not product_row or qty <= 0:
                skipped_lines += 1
                continue

            purchase_lines.append((int(product_row["id"]), qty, price))

        if not purchase_lines:
            return "Не знайдено жодного рядка з товарами для створення закупівлі."

        total_docs = 1
        try:
            if currency_code != get_base_currency_code():
                try:
                    db.add_currency_rate(currency_code, doc_date, exchange_rate)
                except Exception as exc:
                    logging.warning("Не вдалося зберегти курс %s на %s: %s", currency_code, doc_date, exc)

            purchase_id = db.create_purchase(
                doc_date, supplier_id, warehouse_id, "", "; ".join(dict.fromkeys(comments)), currency_code, exchange_rate
            )
            db.replace_purchase_lines(purchase_id, purchase_lines, exchange_rate)
        except Exception as exc:
            logging.warning("Не вдалося створити закупівлю: %s", exc)
            skipped_lines += len(purchase_lines)
            total_docs = 0
            posted_docs = 0
            draft_docs = 0
        else:
            if mode == "post":
                try:
                    db.post_purchase(purchase_id)
                    posted_docs = 1
                    draft_docs = 0
                except Exception as exc:
                    logging.warning("Проведення закупівлі #%s завершилось помилкою: %s", purchase_id, exc)
                    posted_docs = 0
                    draft_docs = 1
            else:
                posted_docs = 0
                draft_docs = 1

        lines_msg = f"Пропущено рядків: {skipped_lines}" if skipped_lines else "Без пропусків"
        created_parts = []
        if created_products:
            created_parts.append(f"створено товарів: {created_products}")
        created_msg = ", ".join(created_parts) if created_parts else "без нових довідників"
        return (
            f"Опрацьовано документів: {total_docs}. Проведено: {posted_docs}, чернеток: {draft_docs}. "
            f"{lines_msg}; {created_msg}."
        )

    # Extra costs
    def create_extra_costs_tab(self) -> None:
        filters = ttk.Frame(self.extra_costs_frame)
        filters.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(filters, text="Статус:").pack(side=tk.LEFT)
        self.extra_status_var = tk.StringVar(value="Усі")
        ttk.Combobox(filters, textvariable=self.extra_status_var, values=["Усі", "Чернетка", "Проведений"], state="readonly", width=14).pack(side=tk.LEFT, padx=4)
        ttk.Label(filters, text="Дата з:").pack(side=tk.LEFT)
        self.extra_date_from_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.extra_date_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(filters, text="по:").pack(side=tk.LEFT)
        self.extra_date_to_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.extra_date_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(filters, text="Фільтр", command=self.refresh_extra_costs).pack(side=tk.LEFT, padx=6)

        columns = [
            ("doc_date", "Дата", 90),
            ("partner", "Контрагент", 180),
            ("currency", "Валюта", 80),
            ("rate", "Курс", 80),
            ("total_doc", "Сума (вал)", 110),
            ("total_base", "Сума (база)", 110),
            ("status", "Статус", 100),
            ("comment", "Коментар", 220),
        ]
        self.extra_table = TableFrame(self.extra_costs_frame, columns)
        self.extra_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.extra_table.on_double_click(self.edit_extra_cost)
        self.extra_table.register_context_menu(self.edit_extra_cost, self.delete_extra_cost)

        btns = ttk.Frame(self.extra_costs_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Новий документ", command=self.new_extra_cost).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_extra_cost).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_extra_cost).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Провести", command=self.post_extra_cost_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Відмінити проведення", command=self.unpost_extra_cost_action).pack(side=tk.LEFT, padx=4)

    def _selected_extra_cost(self):
        doc_id = self.extra_table.selected_id()
        if not doc_id:
            show_error("Супутні витрати", "Оберіть документ")
            return None
        return doc_id

    def refresh_extra_costs(self) -> None:
        status_filter = self.extra_status_var.get()
        status_value = "draft" if status_filter == "Чернетка" else "posted" if status_filter == "Проведений" else None
        rows = db.list_extra_cost_documents(status_value, self.extra_date_from_var.get().strip() or None, self.extra_date_to_var.get().strip() or None)
        self.extra_table.set_rows(
            [
                {
                    "id": r["id"],
                    "doc_date": r["doc_date"],
                    "partner": r["partner"] or "-",
                    "currency": r["currency_code"],
                    "rate": f"{r['exchange_rate']:.4f}",
                    "total_doc": f"{r['total_amount_doc']:.2f}",
                    "total_base": f"{r['total_amount_base']:.2f}",
                    "status": "Чернетка" if (r["status"] or "draft").strip() == "draft" else "Проведений",
                    "comment": r["comment"] or "",
                }
                for r in rows
            ]
        )

    def new_extra_cost(self) -> None:
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        purchases = db.list_purchases(status="posted")
        result = extra_cost_prompt(counterparties, currencies, purchases)
        if not result:
            return
        info, lines, purchase_ids, post_now = result
        try:
            doc_id = db.create_extra_cost_document(
                info["doc_date"],
                info["currency"],
                info["rate"],
                info["partner_id"],
                info["comment"],
            )
            db.replace_extra_cost_lines(doc_id, lines, info["rate"])
            if post_now:
                db.post_extra_cost(doc_id, purchase_ids)
            self.refresh_extra_costs()
        except Exception as exc:
            logging.exception("Create extra cost error")
            show_error("Супутні витрати", str(exc))

    def edit_extra_cost(self) -> None:
        doc_id = self._selected_extra_cost()
        if not doc_id:
            return
        doc = db.get_extra_cost_document(doc_id)
        if not doc:
            return
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        purchases = db.list_purchases(status="posted")
        lines = db.list_extra_cost_lines(doc_id)
        selected_ids = [row["purchase_id"] for row in db.list_extra_cost_allocations(doc_id)]
        allow_edit = doc["status"] == "draft"
        result = extra_cost_prompt(counterparties, currencies, purchases, doc=doc, lines=lines, selected_purchase_ids=selected_ids, allow_edit=allow_edit)
        if not result or not allow_edit:
            return
        info, new_lines, purchase_ids, post_now = result
        try:
            db.update_extra_cost_document(
                doc_id,
                info["doc_date"],
                info["currency"],
                info["rate"],
                info["partner_id"],
                info["comment"],
            )
            db.replace_extra_cost_lines(doc_id, new_lines, info["rate"])
            if post_now:
                db.post_extra_cost(doc_id, purchase_ids)
            self.refresh_extra_costs()
        except Exception as exc:
            logging.exception("Edit extra cost error")
            show_error("Супутні витрати", str(exc))

    def delete_extra_cost(self) -> None:
        doc_id = self._selected_extra_cost()
        if not doc_id:
            return
        if not messagebox.askyesno("Підтвердження", "Видалити документ?"):
            return
        try:
            db.delete_extra_cost_document(doc_id)
            self.refresh_extra_costs()
        except Exception as exc:
            logging.exception("Delete extra cost error")
            show_error("Супутні витрати", str(exc))

    def post_extra_cost_action(self) -> None:
        doc_id = self._selected_extra_cost()
        if not doc_id:
            return
        doc = db.get_extra_cost_document(doc_id)
        status = (doc["status"] or "draft").strip() if doc else None
        if not doc:
            show_error("Супутні витрати", "Документ не знайдено")
            return
        if status != "draft":
            show_error("Супутні витрати", "Проводити можна лише чернетку")
            return
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        purchases = db.list_purchases(status="posted")
        lines = db.list_extra_cost_lines(doc_id)
        selected_ids = [row["purchase_id"] for row in db.list_extra_cost_allocations(doc_id)]
        result = extra_cost_prompt(
            counterparties,
            currencies,
            purchases,
            doc=doc,
            lines=lines,
            selected_purchase_ids=selected_ids,
        )
        if not result:
            return
        info, new_lines, purchase_ids, _ = result
        try:
            db.update_extra_cost_document(
                doc_id,
                info["doc_date"],
                info["currency"],
                info["rate"],
                info["partner_id"],
                info["comment"],
            )
            db.replace_extra_cost_lines(doc_id, new_lines, info["rate"])
            db.post_extra_cost(doc_id, purchase_ids)
            self.refresh_extra_costs()
            self.refresh_purchases()
            self.refresh_stock()
        except Exception as exc:
            logging.exception("Post extra cost error")
            show_error("Супутні витрати", str(exc))

    def unpost_extra_cost_action(self) -> None:
        doc_id = self._selected_extra_cost()
        if not doc_id:
            return
        try:
            db.unpost_extra_cost(doc_id)
            self.refresh_extra_costs()
            self.refresh_purchases()
            self.refresh_stock()
        except Exception as exc:
            logging.exception("Unpost extra cost error")
            show_error("Супутні витрати", str(exc))

    # Sales
    def create_sales_tab(self) -> None:
        filters = ttk.Frame(self.sales_frame)
        filters.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(filters, text="Статус:").pack(side=tk.LEFT)
        self.sales_status_var = tk.StringVar(value="Усі")
        status_combo = ttk.Combobox(filters, textvariable=self.sales_status_var, values=["Усі", "Чернетка", "Проведений"], state="readonly", width=14)
        status_combo.pack(side=tk.LEFT, padx=4)
        ttk.Label(filters, text="Дата з:").pack(side=tk.LEFT)
        self.sales_date_from_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.sales_date_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(filters, text="по:").pack(side=tk.LEFT)
        self.sales_date_to_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.sales_date_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(filters, text="Фільтр", command=self.refresh_sales).pack(side=tk.LEFT, padx=6)

        columns = [
            ("doc_date", "Дата", 90),
            ("customer", "Покупець", 180),
            ("warehouse", "Склад", 140),
            ("channel", "Канал", 100),
            ("status", "Статус", 90),
            ("currency", "Валюта", 80),
            ("rate", "Курс", 80),
            ("total_doc", "Сума (вал)", 110),
            ("total", "Сума (база)", 110),
            ("comment", "Коментар", 220),
        ]
        self.sales_table = TableFrame(self.sales_frame, columns)
        self.sales_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.sales_table.on_double_click(self.edit_sale)
        self.sales_table.register_context_menu(self.edit_sale, self.delete_sale)

        btns = ttk.Frame(self.sales_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Новий продаж", command=self.new_sale).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Змінити", command=self.edit_sale).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_sale).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Провести", command=self.post_sale_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Відмінити проведення", command=self.unpost_sale_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Імпорт із файлу", command=self.import_sales_from_file).pack(side=tk.LEFT, padx=4)

    def create_inventory_tab(self) -> None:
        filters = ttk.Frame(self.inventory_frame)
        filters.pack(fill=tk.X, padx=8, pady=4)

        ttk.Label(filters, text="Статус:").pack(side=tk.LEFT)
        self.inventory_status_var = tk.StringVar(value="Усі")
        ttk.Combobox(
            filters,
            textvariable=self.inventory_status_var,
            values=["Усі", "Чернетка", "Проведений"],
            state="readonly",
            width=14,
        ).pack(side=tk.LEFT, padx=4)

        ttk.Label(filters, text="Склад:").pack(side=tk.LEFT)
        self.inventory_warehouse_var = tk.StringVar(value="Усі")
        self.inventory_warehouse_combo = ttk.Combobox(
            filters,
            textvariable=self.inventory_warehouse_var,
            values=["Усі"],
            state="readonly",
            width=18,
        )
        self.inventory_warehouse_combo.pack(side=tk.LEFT, padx=4)

        ttk.Label(filters, text="Дата з:").pack(side=tk.LEFT)
        self.inventory_date_from_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.inventory_date_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(filters, text="по:").pack(side=tk.LEFT)
        self.inventory_date_to_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.inventory_date_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(filters, text="Фільтр/Оновити", command=self.refresh_inventory_documents).pack(
            side=tk.LEFT, padx=6
        )

        columns = [
            ("id", "ID", 60),
            ("doc_date", "Дата", 90),
            ("warehouse", "Склад", 160),
            ("status", "Статус", 90),
            ("lines_count", "Рядків", 80),
            ("diff_total", "Розбіжність", 120),
            ("comment", "Коментар", 240),
        ]
        self.inventory_table = TableFrame(self.inventory_frame, columns)
        self.inventory_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.inventory_table.on_double_click(self.edit_inventory)
        self.inventory_table.register_context_menu_actions(
            [
                ("Редагувати/Переглянути", self.edit_inventory),
                ("Видалити", self.delete_inventory),
            ]
        )

        btns = ttk.Frame(self.inventory_frame)
        btns.pack(pady=4)
        ttk.Button(btns, text="Створити…", command=self.new_inventory).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Редагувати…", command=self.edit_inventory).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Провести", command=self.post_inventory_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Розпровести", command=self.unpost_inventory_action).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Видалити", command=self.delete_inventory).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns, text="Експорт CSV", command=self.export_inventory_csv).pack(side=tk.LEFT, padx=4)

    def _selected_inventory(self) -> Optional[int]:
        doc_id = self.inventory_table.selected_id()
        if not doc_id:
            show_error("Інвентаризація", "Оберіть документ")
            return None
        return int(doc_id)

    def _refresh_inventory_warehouse_filter(self) -> None:
        warehouses = db.list_warehouses(active_only=False)
        self.inventory_warehouse_options = warehouses
        names = ["Усі"] + [w["name"] for w in warehouses]
        self.inventory_warehouse_combo.configure(values=names)
        if self.inventory_warehouse_var.get() not in names:
            self.inventory_warehouse_var.set("Усі")

    def refresh_inventory_documents(self) -> None:
        self._refresh_inventory_warehouse_filter()
        status_filter = self.inventory_status_var.get()
        status_value = "draft" if status_filter == "Чернетка" else "posted" if status_filter == "Проведений" else None
        warehouse_name = self.inventory_warehouse_var.get()
        warehouse_id = None
        if warehouse_name and warehouse_name != "Усі":
            match = next((w["id"] for w in self.inventory_warehouse_options if w["name"] == warehouse_name), None)
            warehouse_id = match
        rows = db.list_inventory_documents(
            status_value,
            self.inventory_date_from_var.get().strip() or None,
            self.inventory_date_to_var.get().strip() or None,
            warehouse_id,
        )
        self.inventory_table.set_rows(
            [
                {
                    "id": r["id"],
                    "doc_date": r["doc_date"],
                    "warehouse": r["warehouse_name"] or "-",
                    "status": "Чернетка" if r["status"] == "draft" else "Проведений",
                    "lines_count": r["lines_count"],
                    "diff_total": f"{float(r['diff_total'] or 0.0):.2f}",
                    "comment": r["comment"] or "",
                }
                for r in rows
            ]
        )

    def new_inventory(self) -> None:
        warehouses = db.list_warehouses(active_only=True)
        products = db.list_products()
        result = inventory_prompt(warehouses, products, self.settings)
        if not result:
            return
        info, lines, post_now = result
        try:
            doc_id = db.create_inventory_document(info["doc_date"], info["warehouse_id"], info["comment"])
            db.replace_inventory_lines(doc_id, lines)
            if post_now:
                db.post_inventory(doc_id)
            self.refresh_inventory_documents()
        except Exception as exc:
            logging.exception("Create inventory error")
            show_error("Інвентаризація", str(exc))

    def edit_inventory(self) -> None:
        doc_id = self._selected_inventory()
        if not doc_id:
            return
        doc = db.get_inventory_document(doc_id)
        if not doc:
            return
        lines = db.list_inventory_lines(doc_id)
        warehouses = db.list_warehouses(active_only=False)
        products = db.list_products()
        result = inventory_prompt(warehouses, products, self.settings, doc=doc, lines=lines)
        if not result:
            return
        info, new_lines, post_now = result
        try:
            if doc["status"] == "draft":
                db.update_inventory_document(doc_id, info["doc_date"], info["warehouse_id"], info["comment"])
                db.replace_inventory_lines(doc_id, new_lines)
                if post_now:
                    db.post_inventory(doc_id)
            else:
                db.update_inventory_document(doc_id, doc["doc_date"], doc["warehouse_id"], info["comment"])
            self.refresh_inventory_documents()
        except Exception as exc:
            logging.exception("Edit inventory error")
            show_error("Інвентаризація", str(exc))

    def delete_inventory(self) -> None:
        doc_id = self._selected_inventory()
        if not doc_id:
            return
        if not messagebox.askyesno("Підтвердження", "Видалити документ?"):
            return
        try:
            db.delete_inventory_document(doc_id)
            self.refresh_inventory_documents()
        except Exception as exc:
            logging.exception("Delete inventory error")
            show_error("Інвентаризація", str(exc))

    def post_inventory_action(self) -> None:
        doc_id = self._selected_inventory()
        if not doc_id:
            return
        try:
            doc = db.get_inventory_document(doc_id)
            if not doc:
                raise ValueError("Документ не знайдено")
            latest = db.get_latest_posted_stock_doc_date()
            if latest and doc["doc_date"] < latest:
                proceed = messagebox.askyesno(
                    "Підтвердження",
                    "Документ інвентаризації датований "
                    f"{doc['doc_date']}, але є проведені документи до {latest}.\n"
                    f"Проведення змінить історію залишків після {doc['doc_date']}. Продовжити?",
                )
                if not proceed:
                    return
            db.post_inventory(doc_id)
            self.refresh_inventory_documents()
        except Exception as exc:
            logging.exception("Post inventory error")
            show_error("Інвентаризація", str(exc))

    def unpost_inventory_action(self) -> None:
        doc_id = self._selected_inventory()
        if not doc_id:
            return
        try:
            db.unpost_inventory(doc_id)
            self.refresh_inventory_documents()
        except Exception as exc:
            logging.exception("Unpost inventory error")
            show_error("Інвентаризація", str(exc))

    def export_inventory_csv(self) -> None:
        doc_id = self.inventory_table.selected_id()
        default_name = "inventory_lines.csv" if doc_id else "inventory_documents.csv"
        file_path = filedialog.asksaveasfilename(
            title="Експорт CSV",
            defaultextension=".csv",
            initialfile=default_name,
            initialdir=str(self.default_workdir()),
            filetypes=[("CSV", "*.csv"), ("Усі файли", "*.*")],
        )
        if not file_path:
            return
        try:
            if doc_id:
                doc = db.get_inventory_document(int(doc_id))
                if not doc:
                    raise ValueError("Документ не знайдено")
                lines = db.list_inventory_lines(int(doc_id))
                warehouse_name = next(
                    (w["name"] for w in db.list_warehouses(active_only=False) if w["id"] == doc["warehouse_id"]),
                    "",
                )
                with open(file_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(
                        ["date", "warehouse", "sku", "name", "expected_qty", "counted_qty", "diff", "cost_override", "note"]
                    )
                    for ln in lines:
                        writer.writerow(
                            [
                                doc["doc_date"],
                                warehouse_name,
                                ln["sku"],
                                ln["name"],
                                ln["expected_qty"],
                                ln["counted_qty"],
                                ln["diff"],
                                ln["cost_override"] if ln["cost_override"] is not None else "",
                                ln["note"] or "",
                            ]
                        )
            else:
                rows = db.list_inventory_documents()
                with open(file_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(["id", "date", "warehouse", "status", "lines_count", "diff_total", "comment"])
                    for row in rows:
                        writer.writerow(
                            [
                                row["id"],
                                row["doc_date"],
                                row["warehouse_name"] or "",
                                row["status"],
                                row["lines_count"],
                                row["diff_total"],
                                row["comment"] or "",
                            ]
                        )
            messagebox.showinfo("Експорт CSV", "Дані збережено.")
        except Exception:
            logging.exception("Inventory export error")
            show_error("Інвентаризація", "Не вдалося експортувати дані.")

    def _selected_sale(self):
        doc_id = self.sales_table.selected_id()
        if not doc_id:
            show_error("Продажі", "Оберіть документ")
            return None
        return doc_id

    def refresh_sales(self) -> None:
        status_filter = self.sales_status_var.get()
        status_value = "draft" if status_filter == "Чернетка" else "posted" if status_filter == "Проведений" else None
        rows = db.list_sales(status_value, self.sales_date_from_var.get().strip() or None, self.sales_date_to_var.get().strip() or None)
        self.sales_table.set_rows(
            [
                {
                    "id": r["id"],
                    "doc_date": r["doc_date"],
                    "customer": r["customer"] or "-",
                    "warehouse": r["warehouse"] or "-",
                    "channel": r["channel"] or "-",
                    "status": "Чернетка" if r["status"] == "draft" else "Проведений",
                    "currency": r["currency_code"],
                    "rate": f"{r['exchange_rate']:.4f}",
                    "total_doc": f"{r['total_doc']:.2f}",
                    "total": f"{r['total']:.2f}",
                    "comment": r["comment"] or "",
                }
                for r in rows
            ]
        )

    def new_sale(self) -> None:
        products = db.list_products()
        warehouses = db.list_warehouses(active_only=True)
        channels = db.list_channels(active_only=True)
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        result = document_prompt("sale", products, counterparties, warehouses, channels, currencies, settings=self.settings)
        if not result:
            return
        info, lines = result
        try:
            doc_id = db.create_sale(
                info["doc_date"],
                info["counterparty_id"],
                info["warehouse_id"],
                info["channel"],
                info["comment"],
                info["currency"],
                info["rate"],
                info.get("order_expense_doc", 0.0),
            )
            db.replace_sale_lines(doc_id, lines, info["rate"], info.get("order_expense_doc", 0.0))
            self.refresh_sales()
        except Exception:
            logging.exception("Create sale error")
            show_error("Продажі", "Не вдалося створити документ")

    def edit_sale(self) -> None:
        doc_id = self._selected_sale()
        if not doc_id:
            return
        doc = db.get_sale(doc_id)
        if not doc:
            return
        lines = db.list_sale_lines(doc_id)
        products = db.list_products()
        warehouses = db.list_warehouses(active_only=False)
        channels = db.list_channels(active_only=False)
        counterparties = db.list_counterparties()
        currencies = db.list_currencies()
        result = document_prompt(
            "sale",
            products,
            counterparties,
            warehouses,
            channels,
            currencies,
            doc=doc,
            lines=lines,
            settings=self.settings,
        )
        if not result:
            return
        info, new_lines = result
        try:
            if doc["status"] == "draft":
                db.update_sale(
                    doc_id,
                    info["doc_date"],
                    info["counterparty_id"],
                    info["warehouse_id"],
                    info["channel"],
                    info["comment"],
                    info["currency"],
                    info["rate"],
                    info.get("order_expense_doc", 0.0),
                )
                db.replace_sale_lines(doc_id, new_lines, info["rate"], info.get("order_expense_doc", 0.0))
            else:
                db.update_sale(
                    doc_id,
                    doc["doc_date"],
                    doc["customer_id"],
                    doc["warehouse_id"],
                    doc["channel"] or "",
                    info["comment"],
                    doc["currency_code"],
                    doc["exchange_rate"],
                    doc.get("order_expense_doc", 0.0),
                )
            self.refresh_sales()
        except Exception:
            logging.exception("Edit sale error")
            show_error("Продажі", "Не вдалося змінити документ")

    def delete_sale(self) -> None:
        doc_id = self._selected_sale()
        if not doc_id:
            return
        if not messagebox.askyesno("Підтвердження", "Видалити документ?"):
            return
        try:
            with db.get_connection() as conn:
                status_row = conn.execute("SELECT status FROM SalesDocuments WHERE id=?", (doc_id,)).fetchone()
            if not status_row:
                raise ValueError("Документ не знайдено")

            was_posted = status_row[0] == "posted"
            if was_posted:
                if not messagebox.askyesno(
                    "Підтвердження", "Документ проведено. Скасувати проведення та видалити?"
                ):
                    return
                db.unpost_sale(doc_id)

            with db.get_connection() as conn:
                with db.safe_transaction(conn):
                    conn.execute("DELETE FROM SalesDocuments WHERE id=?", (doc_id,))

            self.refresh_sales()
            if was_posted:
                self.refresh_stock()
                self.refresh_cash()
        except Exception as exc:
            logging.exception("Delete sale error")
            show_error("Продажі", str(exc))

    def post_sale_action(self) -> None:
        doc_id = self._selected_sale()
        if not doc_id:
            return
        try:
            db.post_sale(doc_id)
            self.refresh_sales()
            self.refresh_stock()
            self.refresh_cash()
        except Exception as exc:
            logging.exception("Post sale error")
            show_error("Продажі", str(exc))

    def unpost_sale_action(self) -> None:
        doc_id = self._selected_sale()
        if not doc_id:
            return
        try:
            db.unpost_sale(doc_id)
            self.refresh_sales()
            self.refresh_stock()
            self.refresh_cash()
        except Exception as exc:
            logging.exception("Unpost sale error")
            show_error("Продажі", str(exc))

    def import_sales_from_file(self) -> None:
        file_path = filedialog.askopenfilename(
            title="Файл замовлень",
            filetypes=[("CSV", "*.csv"), ("Excel", "*.xlsx *.xls"), ("Усі файли", "*.*")],
            initialdir=str(self.default_workdir()),
        )
        if not file_path:
            return

        try:
            raw_rows, headers = parse_sales_file(Path(file_path), encoding=self.settings.get("files", "encoding") or "utf-8")
        except Exception as exc:
            logging.exception("Не вдалося прочитати файл імпорту")
            show_error(
                "Імпорт продажів",
                "Не вдалося прочитати файл. Перевірте формат, кодування та структуру даних.\n" + str(exc),
            )
            return

        if not raw_rows:
            messagebox.showinfo("Імпорт продажів", "У файлі не знайдено рядків із товарами.")
            return

        warehouses = db.list_warehouses(active_only=True)
        if not warehouses:
            show_error("Імпорт продажів", "Спочатку створіть хоча б один склад.")
            return
        channels = db.list_channels(active_only=False)

        dialog = SalesImportDialog(self, raw_rows, headers, warehouses, channels, self.settings)
        result = dialog.result
        if not result:
            return

        try:
            summary = self._process_sales_import(result["orders"], result["options"])
        except Exception:
            logging.exception("Помилка під час імпорту продажів")
            show_error("Імпорт продажів", "Імпорт перервано помилкою. Деталі у логах.")
            return

        messagebox.showinfo("Імпорт продажів", summary)
        self.refresh_sales()
        self.refresh_cash()
        self.refresh_stock()

    def _process_sales_import(self, orders: list[dict], options: dict) -> str:
        warehouse_id = options["warehouse_id"]
        channel_override = options.get("channel", "")
        mode = options.get("mode", "draft")
        allow_negative = bool(options.get("allow_negative"))
        create_products = bool(options.get("create_products", True))
        create_customers = bool(options.get("create_customers"))
        use_file_channel = bool(options.get("use_file_channel"))

        product_rows = db.list_products()
        products_by_sku = {p["sku"].lower(): dict(p) for p in product_rows if p["sku"]}
        products_by_name = {p["name"].lower(): dict(p) for p in product_rows if p["name"]}
        products_by_supplier_sku = {p["supplier_sku"].lower(): dict(p) for p in product_rows if p.get("supplier_sku")}
        counterparties = db.list_counterparties()
        allowed_customer_types = {"customer", "both", "other"}
        customers_by_name = {
            c["name"].lower(): c for c in counterparties if c["type"] in allowed_customer_types and c["name"]
        }

        default_brand = self.settings.get("defaults", "product", "brand") or "Імпорт"
        default_category = self.settings.get("defaults", "product", "category") or "Імпорт"
        default_unit = (self.settings.get("defaults", "product", "unit") or "pcs").strip() or "pcs"
        brand_id, category_id = db.ensure_import_defaults(default_brand, default_category)
        stock_map = db.stock_on_hand(warehouse_id)

        created_products = 0
        created_customers = 0
        skipped_lines = 0
        posted_docs = 0
        draft_docs = 0
        total_docs = 0

        grouped: dict[str, list[dict]] = defaultdict(list)
        for idx, row in enumerate(orders):
            key = row.get("order_no") or f"#{idx+1}"
            grouped[key].append(row)

        for order_no, lines in grouped.items():
            doc_date = lines[0].get("doc_date") or datetime.now().strftime("%Y-%m-%d")
            customer_name = lines[0].get("customer", "").strip()
            phone = lines[0].get("phone", "").strip()
            email = lines[0].get("email", "").strip()
            customer_id = None

            if customer_name:
                existing = customers_by_name.get(customer_name.lower()) or db.find_counterparty_by_name(
                    customer_name, allowed_customer_types
                )
                if existing:
                    customer_id = existing["id"]
                    customers_by_name[customer_name.lower()] = dict(existing)
                elif create_customers:
                    customer_id = db.add_counterparty(customer_name, "customer", phone, email, "", "Імпортований клієнт")
                    new_cp = {
                        "id": customer_id,
                        "name": customer_name,
                        "type": "customer",
                        "phone": phone,
                        "email": email,
                        "address": "",
                        "note": "Імпортований клієнт",
                    }
                    counterparties.append(new_cp)
                    customers_by_name[customer_name.lower()] = new_cp
                    created_customers += 1

            sale_lines: list[tuple[int, float, float]] = []
            comment = lines[0].get("comment", "").strip()
            line_channel = lines[0].get("channel", "").strip()
            channel_value = line_channel if (use_file_channel and line_channel) else channel_override

            for row in lines:
                sku = (row.get("sku") or "").strip()
                supplier_sku = (row.get("supplier_sku") or "").strip()
                name = (row.get("product_name") or sku or "Без назви").strip()
                qty = float(row.get("quantity") or 0)
                price = float(row.get("price") or 0)
                amount = float(row.get("amount") or 0)
                if not price and qty and amount:
                    price = amount / qty

                product_row = products_by_supplier_sku.get(supplier_sku.lower()) if supplier_sku else None
                if not product_row:
                    product_row = products_by_sku.get(sku.lower()) if sku else None
                if not product_row and name:
                    product_row = products_by_name.get(name.lower())
                if not product_row and create_products:
                    final_sku = sku or self._generate_unique_sku(name, set(products_by_sku.keys()))
                    product_id = db.add_product(
                        final_sku,
                        name,
                        brand_id,
                        category_id,
                        unit=default_unit,
                        supplier_sku=supplier_sku,
                    )
                    product_row = {
                        "id": product_id,
                        "sku": final_sku,
                        "name": name,
                        "supplier_sku": supplier_sku,
                    }
                    products_by_sku[final_sku.lower()] = product_row
                    if supplier_sku:
                        products_by_supplier_sku[supplier_sku.lower()] = product_row
                    products_by_name[name.lower()] = product_row
                    created_products += 1

                if not product_row or qty <= 0:
                    skipped_lines += 1
                    continue

                sale_lines.append((int(product_row["id"]), qty, price, 0.0))

            if not sale_lines:
                skipped_lines += len(lines)
                continue

            total_docs += 1
            try:
                sale_id = db.create_sale(doc_date, customer_id, warehouse_id, channel_value, comment, "UAH", 1.0)
                db.replace_sale_lines(sale_id, sale_lines, 1.0, 0.0)
            except Exception as exc:
                logging.warning("Не вдалося створити продаж %s: %s", order_no, exc)
                skipped_lines += len(sale_lines)
                continue

            should_post = mode == "post"
            if mode == "in_stock":
                enough = True
                for pid, qty, _, _ in sale_lines:
                    current_qty = stock_map.get(pid, db.get_stock_quantity(pid, warehouse_id))
                    if qty > current_qty:
                        enough = False
                        break
                should_post = enough

            if should_post:
                try:
                    db.post_sale(sale_id, allow_negative=allow_negative)
                    posted_docs += 1
                    for pid, qty, _, _ in sale_lines:
                        stock_map[pid] = stock_map.get(pid, db.get_stock_quantity(pid, warehouse_id)) - qty
                except Exception as exc:
                    logging.warning("Проведення продажу #%s завершилось помилкою: %s", sale_id, exc)
                    try:
                        db.unpost_sale(sale_id)
                    except Exception:
                        logging.exception("Не вдалося скасувати проведення після помилки імпорту")
                    draft_docs += 1
            else:
                draft_docs += 1

        lines_msg = f"Пропущено рядків: {skipped_lines}" if skipped_lines else "Без пропусків"
        created_parts = []
        if created_products:
            created_parts.append(f"створено товарів: {created_products}")
        if created_customers:
            created_parts.append(f"створено клієнтів: {created_customers}")
        created_msg = ", ".join(created_parts) if created_parts else "без нових довідників"
        return (
            f"Опрацьовано документів: {total_docs}. Проведено: {posted_docs}, чернеток: {draft_docs}. "
            f"{lines_msg}; {created_msg}."
        )

    def _generate_unique_sku(self, base: str, existing: set[str]) -> str:
        clean = (base or "AUTO").upper().replace(" ", "")
        if len(clean) < 3:
            clean = f"AUTO{clean}"
        candidate = clean[:20] or "AUTO"
        idx = 1
        while candidate.lower() in existing:
            idx += 1
            candidate = f"{clean[:15]}-{idx}"
        return candidate

    # Cash
    def create_cash_tab(self) -> None:
        top = ttk.Frame(self.cash_frame)
        top.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(top, text="Дата з:").pack(side=tk.LEFT)
        self.cash_date_from_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.cash_date_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(top, text="по:").pack(side=tk.LEFT)
        self.cash_date_to_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.cash_date_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="Фільтр", command=self.refresh_cash).pack(side=tk.LEFT, padx=6)
        ttk.Button(top, text="Додати рух", command=self.add_cash).pack(side=tk.LEFT, padx=6)

        columns = [
            ("date", "Дата", 90),
            ("type", "Тип", 140),
            ("amount", "Сума", 100),
            ("counterparty", "Контрагент", 160),
            ("channel", "Канал", 120),
            ("related", "Документ", 120),
            ("comment", "Коментар", 260),
        ]
        self.cash_table = TableFrame(self.cash_frame, columns)
        self.cash_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

    def refresh_cash(self) -> None:
        rows = db.list_cash(self.cash_date_from_var.get().strip() or None, self.cash_date_to_var.get().strip() or None)
        type_labels = {
            "sale_payment": "Оплата від клієнта",
            "purchase_payment": "Оплата постачальнику",
            "other_income": "Інший дохід",
            "other_variable_expense": "Змінна витрата",
            "fixed_expense": "Постійна витрата",
        }
        self.cash_table.set_rows(
            [
                {
                    "id": r["id"],
                    "date": r["date"],
                    "type": type_labels.get(r["type"], r["type"]),
                    "amount": f"{r['amount']:.2f}",
                    "counterparty": r["counterparty"] or "-",
                    "channel": r["channel"] or "-",
                    "related": f"{r['related_doc_type'] or ''} #{r['related_doc_id'] or ''}",
                    "comment": r["comment"] or "",
                }
                for r in rows
            ]
        )

    def add_cash(self) -> None:
        counterparties = db.list_counterparties()
        channels = db.list_channels(active_only=True)
        transactions = cash_prompt(counterparties, channels)
        if not transactions:
            return
        try:
            for tx in transactions:
                db.add_cash_transaction(**tx)
            self.refresh_cash()
        except Exception:
            logging.exception("Add cash error")
            show_error("Каса", "Не вдалося зберегти рух коштів")

    # Stock
    def create_stock_tab(self) -> None:
        top = ttk.Frame(self.stock_frame)
        top.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(top, text="Пошук товару:").pack(side=tk.LEFT)
        self.stock_search_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.stock_search_var, width=30).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="Оновити", command=self.on_search_stock).pack(side=tk.LEFT)
        ttk.Button(top, text="Перерахувати залишки", command=self.recalc_stock).pack(side=tk.LEFT, padx=6)

        columns = [("name", "Товар", 240), ("sku", "SKU", 120), ("warehouse", "Склад", 160), ("quantity", "Кількість", 100), ("average_cost", "Сер. собівартість", 140)]
        self.stock_table = TableFrame(self.stock_frame, columns, selectmode="extended")
        self.stock_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self._stock_row_meta: dict[str, dict[str, object]] = {}
        self.stock_table.register_context_menu_actions(
            [
                ("Друк етикеток…", self.print_stock_labels),
                ("Масові дії з товарами…", self.bulk_actions_from_stock),
            ]
        )

    def on_search_stock(self) -> None:
        self.refresh_stock(self.stock_search_var.get())

    def refresh_stock(self, search: str | None = None) -> None:
        rows = db.list_stock(search)
        meta: dict[str, dict[str, object]] = {}
        table_rows = []
        for r in rows:
            iid = f"{r['product_id']}-{r['warehouse_id'] if r['warehouse_id'] else '0'}"
            qty = float(r["quantity"] or 0)
            meta[iid] = {
                "product_id": int(r["product_id"]),
                "warehouse_id": int(r["warehouse_id"]) if r["warehouse_id"] is not None else 0,
                "qty": qty,
                "sku": r["sku"],
                "name": r["name"],
            }
            table_rows.append(
                {
                    "id": iid,
                    "name": r["name"],
                    "sku": r["sku"],
                    "warehouse": r["warehouse"] or "-",
                    "quantity": f"{qty:.2f}",
                    "average_cost": f"{r['average_cost']:.2f}",
                }
            )
        self._stock_row_meta = meta
        self.stock_table.set_rows(table_rows)

    def recalc_stock(self) -> None:
        if not messagebox.askyesno("Підтвердження", "Перерахувати усі залишки? Це використовує рухи товарів"):
            return
        try:
            db.recalc_stock()
            self.refresh_stock()
            messagebox.showinfo("Залишки", "Перерахунок виконано")
        except Exception as exc:
            logging.exception("Recalc stock error")
            show_error("Залишки", str(exc))

    def bulk_actions_from_stock(self) -> None:
        selections = self.stock_table.tree.selection()
        if not selections:
            messagebox.showwarning("Залишки", "Оберіть хоча б одну позицію.")
            return
        product_ids: list[int] = []
        seen: set[int] = set()
        for iid in selections:
            meta = self._stock_row_meta.get(iid) if hasattr(self, "_stock_row_meta") else None
            if not meta:
                continue
            try:
                pid = int(meta.get("product_id"))
            except (TypeError, ValueError):
                continue
            if pid in seen:
                continue
            seen.add(pid)
            product_ids.append(pid)
        if not product_ids:
            messagebox.showwarning("Залишки", "Не вдалося визначити товари у виборі.")
            return
        open_products_bulk_actions_dialog(self, db.get_connection(), None, product_ids)

    def print_stock_labels(self) -> None:
        selections = self.stock_table.tree.selection()
        if not selections:
            messagebox.showwarning("Залишки", "Оберіть позиції для друку етикеток.")
            return
        aggregated: dict[int, dict[str, object]] = {}
        for iid in selections:
            meta = self._stock_row_meta.get(iid) if hasattr(self, "_stock_row_meta") else None
            if not meta:
                continue
            try:
                pid = int(meta.get("product_id"))
            except (TypeError, ValueError):
                continue
            qty = float(meta.get("qty", 0) or 0)
            if pid not in aggregated:
                aggregated[pid] = {
                    "product_id": pid,
                    "sku": str(meta.get("sku") or ""),
                    "name": str(meta.get("name") or ""),
                    "qty": 0.0,
                }
            aggregated[pid]["qty"] = float(aggregated[pid].get("qty", 0.0) or 0.0) + qty
        items = [item for item in aggregated.values() if float(item.get("qty", 0) or 0) > 0]
        if not items:
            messagebox.showwarning("Етикетки", "Немає позицій з кількістю > 0 для друку.")
            return
        open_labels_print_dialog(self, items, "Кількість із залишків")

    # Reports
    def create_reports_tab(self) -> None:
        notebook = ttk.Notebook(self.reports_frame)
        notebook.pack(fill=tk.BOTH, expand=True)

        # Dashboard
        dashboard_tab = ttk.Frame(notebook)
        notebook.add(dashboard_tab, text="Дашборд")
        dash_filters = ttk.Frame(dashboard_tab)
        dash_filters.pack(fill=tk.X, padx=8, pady=6)
        ttk.Label(dash_filters, text="Дата з:").pack(side=tk.LEFT)
        self.dashboard_from_var = tk.StringVar()
        ttk.Entry(dash_filters, textvariable=self.dashboard_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(dash_filters, text="по:").pack(side=tk.LEFT)
        self.dashboard_to_var = tk.StringVar()
        ttk.Entry(dash_filters, textvariable=self.dashboard_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(dash_filters, text="Оновити", command=self.refresh_dashboard).pack(side=tk.LEFT, padx=6)

        cards_frame = ttk.Frame(dashboard_tab)
        cards_frame.pack(fill=tk.X, padx=8, pady=(2, 6))
        self.dashboard_cards: dict[str, dict[str, object]] = {}
        for key, title, hint in [
            ("turnover", "Оборот", "оплачені продажі"),
            ("gross_profit", "Валовий прибуток", "дохід мінус собівартість та витрати"),
            ("margin_pct", "Маржа", "%"),
            ("stock_value", "Вартість залишків", "на зараз"),
        ]:
            frame = ttk.LabelFrame(cards_frame, text=title)
            frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=4)
            value_var = tk.StringVar(value="0.00")
            extra_var = tk.StringVar(value=hint)
            ttk.Label(frame, textvariable=value_var, font=("Segoe UI", 13, "bold")).pack(anchor="w", padx=8, pady=(6, 0))
            ttk.Label(frame, textvariable=extra_var, foreground="#555").pack(anchor="w", padx=8, pady=(0, 2))
            canvas = tk.Canvas(frame, height=42, bg="white", highlightthickness=1, highlightbackground="#e5e7eb")
            canvas.pack(fill=tk.X, padx=6, pady=(2, 6))
            self.dashboard_cards[key] = {"value": value_var, "extra": extra_var, "canvas": canvas}

        trend_frame = ttk.LabelFrame(dashboard_tab, text="Тренд обороту / валового прибутку")
        trend_frame.pack(fill=tk.BOTH, expand=False, padx=8, pady=(0, 6))
        self.dashboard_trend_canvas = tk.Canvas(
            trend_frame, height=240, background="white", highlightthickness=1, highlightbackground="#d9d9d9"
        )
        self.dashboard_trend_canvas.pack(fill=tk.BOTH, expand=True, padx=4, pady=6)

        dash_columns = [("name", "Показник", 260), ("value", "Значення", 200), ("extra", "Примітка", 200)]
        self.dashboard_table = TableFrame(dashboard_tab, dash_columns)
        self.dashboard_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        # Sales analysis
        sales_tab = ttk.Frame(notebook)
        notebook.add(sales_tab, text="Продажі")
        sales_filters = ttk.Frame(sales_tab)
        sales_filters.pack(fill=tk.X, padx=8, pady=6)
        ttk.Label(sales_filters, text="Дата з:").pack(side=tk.LEFT)
        self.sales_from_var = tk.StringVar()
        ttk.Entry(sales_filters, textvariable=self.sales_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(sales_filters, text="по:").pack(side=tk.LEFT)
        self.sales_to_var = tk.StringVar()
        ttk.Entry(sales_filters, textvariable=self.sales_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(sales_filters, text="Оновити", command=self.refresh_sales_analysis).pack(side=tk.LEFT, padx=6)
        charts_frame = ttk.Frame(sales_tab)
        charts_frame.pack(fill=tk.X, padx=8, pady=(4, 2))
        channel_box = ttk.LabelFrame(charts_frame, text="Структура по каналах")
        channel_box.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 4))
        self.sales_channel_pie = tk.Canvas(channel_box, height=240, bg="white", highlightthickness=1, highlightbackground="#d9d9d9")
        self.sales_channel_pie.pack(fill=tk.BOTH, expand=True, padx=4, pady=6)
        category_box = ttk.LabelFrame(charts_frame, text="Структура по категоріях")
        category_box.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(4, 0))
        self.sales_category_pie = tk.Canvas(category_box, height=240, bg="white", highlightthickness=1, highlightbackground="#d9d9d9")
        self.sales_category_pie.pack(fill=tk.BOTH, expand=True, padx=4, pady=6)

        sales_tables = ttk.Frame(sales_tab)
        sales_tables.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
        channel_columns = [
            ("name", "Канал", 180),
            ("revenue", "Дохід", 140),
            ("gross_profit", "Валовий прибуток", 160),
            ("qty", "К-сть", 100),
        ]
        category_columns = [
            ("name", "Категорія", 200),
            ("revenue", "Дохід", 140),
            ("gross_profit", "Валовий прибуток", 160),
            ("qty", "К-сть", 100),
        ]
        ttk.Label(sales_tables, text="По каналах", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        self.sales_channels_table = TableFrame(sales_tables, channel_columns)
        self.sales_channels_table.pack(fill=tk.BOTH, expand=True, pady=4)
        ttk.Label(sales_tables, text="По категоріях", font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(10, 0))
        self.sales_categories_table = TableFrame(sales_tables, category_columns)
        self.sales_categories_table.pack(fill=tk.BOTH, expand=True, pady=4)

        # ABC/XYZ
        abc_tab = ttk.Frame(notebook)
        notebook.add(abc_tab, text="ABC/XYZ")
        abc_filters = ttk.Frame(abc_tab)
        abc_filters.pack(fill=tk.X, padx=8, pady=6)
        ttk.Label(abc_filters, text="Дата з:").pack(side=tk.LEFT)
        self.abc_from_var = tk.StringVar()
        ttk.Entry(abc_filters, textvariable=self.abc_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(abc_filters, text="по:").pack(side=tk.LEFT)
        self.abc_to_var = tk.StringVar()
        ttk.Entry(abc_filters, textvariable=self.abc_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Button(abc_filters, text="Оновити", command=self.refresh_abc_xyz).pack(side=tk.LEFT, padx=6)
        abc_columns = [
            ("name", "Товар", 220),
            ("sku", "SKU", 100),
            ("category", "Категорія", 160),
            ("revenue", "Дохід", 120),
            ("abc", "ABC", 60),
            ("xyz", "XYZ", 60),
        ]
        abc_body = ttk.Frame(abc_tab)
        abc_body.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
        self.abc_table = TableFrame(abc_body, abc_columns)
        self.abc_table.pack(fill=tk.BOTH, expand=True, side=tk.LEFT, padx=(0, 6), pady=4)
        self.abc_table.tag_configure("abc_a", background="#ecfdf3")
        self.abc_table.tag_configure("abc_c", background="#fff1f2")
        self.abc_table.tag_configure("xyz_x", foreground="#166534")
        self.abc_table.tag_configure("xyz_z", foreground="#991b1b")

        legend = ttk.LabelFrame(abc_body, text="Легенда")
        legend.pack(fill=tk.Y, side=tk.LEFT, padx=(6, 0), pady=4)
        ttk.Label(legend, text="A/X — лідери та стабільні", foreground="#166534").pack(anchor="w", padx=8, pady=(6, 2))
        ttk.Label(legend, text="C/Z — дрібні та волатильні", foreground="#991b1b").pack(anchor="w", padx=8, pady=2)
        self.abc_distribution_canvas = tk.Canvas(
            legend, width=220, height=140, bg="white", highlightthickness=1, highlightbackground="#d9d9d9"
        )
        self.abc_distribution_canvas.pack(fill=tk.BOTH, expand=True, padx=6, pady=8)

        # Cash flow
        cash_tab = ttk.Frame(notebook)
        notebook.add(cash_tab, text="Каса")
        cash_filters = ttk.Frame(cash_tab)
        cash_filters.pack(fill=tk.X, padx=8, pady=6)
        ttk.Label(cash_filters, text="Дата з:").pack(side=tk.LEFT)
        self.cash_from_var = tk.StringVar()
        ttk.Entry(cash_filters, textvariable=self.cash_from_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(cash_filters, text="по:").pack(side=tk.LEFT)
        self.cash_to_var = tk.StringVar()
        ttk.Entry(cash_filters, textvariable=self.cash_to_var, width=10).pack(side=tk.LEFT, padx=2)
        ttk.Label(cash_filters, text="Тип:").pack(side=tk.LEFT, padx=(10, 2))
        self.cash_type_var = tk.StringVar()
        self.cash_type_combo = ttk.Combobox(
            cash_filters,
            textvariable=self.cash_type_var,
            values=["", "sale_payment", "purchase_payment", "income", "expense", "transfer"],
            width=15,
        )
        self.cash_type_combo.pack(side=tk.LEFT)
        ttk.Label(cash_filters, text="Канал:").pack(side=tk.LEFT, padx=(10, 2))
        self.cash_channel_var = tk.StringVar()
        ttk.Entry(cash_filters, textvariable=self.cash_channel_var, width=12).pack(side=tk.LEFT, padx=2)
        ttk.Label(cash_filters, text="Контрагент:").pack(side=tk.LEFT, padx=(10, 2))
        self.cash_counterparty_var = tk.StringVar()
        self.cash_counterparty_combo = ttk.Combobox(cash_filters, textvariable=self.cash_counterparty_var, width=25)
        self.cash_counterparty_combo.pack(side=tk.LEFT)
        ttk.Button(cash_filters, text="Показати", command=self.refresh_cash_flow_report).pack(side=tk.LEFT, padx=6)

        cash_summary_box = ttk.LabelFrame(cash_tab, text="Рух коштів за типами")
        cash_summary_box.pack(fill=tk.X, padx=8, pady=(0, 6))
        self.cash_summary_canvas = tk.Canvas(
            cash_summary_box, height=200, bg="white", highlightthickness=1, highlightbackground="#d9d9d9"
        )
        self.cash_summary_canvas.pack(fill=tk.BOTH, expand=True, padx=6, pady=6)

        cash_columns = [
            ("date", "Дата", 90),
            ("type", "Тип", 140),
            ("amount", "Сума", 100),
            ("counterparty", "Контрагент", 180),
            ("channel", "Канал", 120),
            ("comment", "Коментар", 200),
        ]
        self.cash_report_table = TableFrame(cash_tab, cash_columns)
        self.cash_report_table.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.refresh_dashboard()
        self.refresh_sales_analysis()
        self.refresh_abc_xyz()
        self.refresh_cash_counterparties()
        self.refresh_cash_flow_report()

    def refresh_dashboard(self) -> None:
        metrics = db.dashboard_metrics(
            self.dashboard_from_var.get().strip() or None, self.dashboard_to_var.get().strip() or None
        )
        if hasattr(self, "dashboard_cards"):
            self.dashboard_cards["turnover"]["value"].set(f"{metrics['turnover']:.2f}")
            self.dashboard_cards["gross_profit"]["value"].set(f"{metrics['gross_profit']:.2f}")
            self.dashboard_cards["margin_pct"]["value"].set(f"{metrics['margin_pct']:.2f}%")
            self.dashboard_cards["stock_value"]["value"].set(f"{metrics['stock_value']:.2f}")
        self.dashboard_table.set_rows(
            [
                {"id": 1, "name": "Оборот", "value": f"{metrics['turnover']:.2f}", "extra": "оплачені продажі"},
                {
                    "id": 2,
                    "name": "Валовий прибуток",
                    "value": f"{metrics['gross_profit']:.2f}",
                    "extra": "дохід мінус собівартість та витрати продажів",
                },
                {"id": 3, "name": "Маржа", "value": f"{metrics['margin_pct']:.2f}%", "extra": ""},
                {"id": 4, "name": "Вартість залишків", "value": f"{metrics['stock_value']:.2f}", "extra": "на зараз"},
            ]
        )
        self.update_dashboard_trend()

    def refresh_sales_analysis(self) -> None:
        analysis = db.sales_analysis(self.sales_from_var.get().strip() or None, self.sales_to_var.get().strip() or None)
        self.sales_channels_table.set_rows(
            [
                {
                    "id": idx,
                    "name": r["name"],
                    "revenue": f"{r['revenue']:.2f}",
                    "gross_profit": f"{r['gross_profit']:.2f}",
                    "qty": f"{r['qty']:.2f}",
                }
                for idx, r in enumerate(analysis["channels"], 1)
            ]
        )
        self.sales_categories_table.set_rows(
            [
                {
                    "id": idx,
                    "name": r["name"],
                    "revenue": f"{r['revenue']:.2f}",
                    "gross_profit": f"{r['gross_profit']:.2f}",
                    "qty": f"{r['qty']:.2f}",
                }
                for idx, r in enumerate(analysis["categories"], 1)
            ]
        )
        self.draw_pie_chart(
            self.sales_channel_pie,
            [(r["name"], r["revenue"]) for r in analysis["channels"]],
            "Канали продажу",
        )
        self.draw_pie_chart(
            self.sales_category_pie,
            [(r["name"], r["revenue"]) for r in analysis["categories"]],
            "Категорії",
        )

    def refresh_abc_xyz(self) -> None:
        rows = db.abc_xyz_report(self.abc_from_var.get().strip() or None, self.abc_to_var.get().strip() or None)
        table_rows = []
        distribution: defaultdict[str, int] = defaultdict(int)
        for r in rows:
            tags = []
            if r["abc"] == "A":
                tags.append("abc_a")
            if r["abc"] == "C":
                tags.append("abc_c")
            if r["xyz"] == "X":
                tags.append("xyz_x")
            if r["xyz"] == "Z":
                tags.append("xyz_z")
            combo = f"{r['abc']}/{r['xyz']}"
            distribution[combo] += 1
            table_rows.append(
                {
                    "id": r["product_id"],
                    "name": r["name"],
                    "sku": r["sku"],
                    "category": r["category"],
                    "revenue": f"{r['revenue']:.2f}",
                    "abc": r["abc"],
                    "xyz": r["xyz"],
                    "tags": tags,
                }
            )
        self.abc_table.set_rows(table_rows)
        self.draw_bar_chart(
            self.abc_distribution_canvas,
            sorted(distribution.items()),
            title="Кількість товарів за групами",
            bar_color="#2563eb",
        )

    def refresh_cash_counterparties(self) -> None:
        counterparts = db.list_counterparties()
        names = ["Усі"] + [c["name"] for c in counterparts]
        self.cash_counterparty_combo["values"] = names
        if not self.cash_counterparty_var.get():
            self.cash_counterparty_combo.current(0) if names else None

    def refresh_cash_flow_report(self) -> None:
        cp_name = self.cash_counterparty_var.get().strip()
        cp_id = None
        if cp_name and cp_name != "Усі":
            for c in db.list_counterparties():
                if c["name"] == cp_name:
                    cp_id = c["id"]
                    break
        rows = db.cash_flow_detailed(
            self.cash_from_var.get().strip() or None,
            self.cash_to_var.get().strip() or None,
            self.cash_type_var.get().strip() or None,
            self.cash_channel_var.get().strip() or None,
            cp_id,
        )
        self.cash_report_table.set_rows(
            [
                {
                    "id": r["id"],
                    "date": r["date"],
                    "type": r["type"],
                    "amount": f"{r['amount']:.2f}",
                    "counterparty": r["counterparty"],
                    "channel": r["channel"],
                    "comment": r["comment"],
                }
                for r in rows
            ]
        )
        summary = db.cash_flow_summary(self.cash_from_var.get().strip() or None, self.cash_to_var.get().strip() or None)
        self.draw_bar_chart(
            self.cash_summary_canvas,
            [(r["type"], r["total"]) for r in summary],
            title="Сума за обраний період",
            bar_color="#0ea5e9",
        )

    # Charts & visualizations
    def update_dashboard_trend(self) -> None:
        trend = db.dashboard_trends(
            self.dashboard_from_var.get().strip() or None, self.dashboard_to_var.get().strip() or None
        )
        periods = [row["period"] for row in trend]
        turnover = [row["turnover"] for row in trend]
        gross_profit = [row["gross_profit"] for row in trend]
        self.draw_line_chart(
            self.dashboard_trend_canvas,
            periods,
            [
                ("Оборот", turnover, self.chart_palette[0]),
                ("Валовий прибуток", gross_profit, self.chart_palette[1]),
            ],
        )
        if trend:
            self.draw_sparkline(self.dashboard_cards["turnover"]["canvas"], turnover, self.chart_palette[0])
            self.draw_sparkline(self.dashboard_cards["gross_profit"]["canvas"], gross_profit, self.chart_palette[1])
        else:
            self.draw_sparkline(self.dashboard_cards["turnover"]["canvas"], [], self.chart_palette[0])
            self.draw_sparkline(self.dashboard_cards["gross_profit"]["canvas"], [], self.chart_palette[1])
        self.draw_sparkline(self.dashboard_cards["margin_pct"]["canvas"], [], self.chart_palette[2])
        self.draw_sparkline(self.dashboard_cards["stock_value"]["canvas"], [], self.chart_palette[3])

    def _canvas_size(self, canvas: tk.Canvas) -> tuple[int, int]:
        canvas.update_idletasks()
        width = max(int(canvas.winfo_width() or canvas.cget("width")), 200)
        height = max(int(canvas.winfo_height() or canvas.cget("height")), 120)
        return width, height

    def draw_line_chart(self, canvas: tk.Canvas, labels: list[str], series: list[tuple[str, list[float], str]]) -> None:
        canvas.delete("all")
        width, height = self._canvas_size(canvas)
        margin = 40
        if not labels or not any(vals for _, vals, _ in series):
            canvas.create_text(width / 2, height / 2, text="Немає даних", fill="#6b7280")
            return
        max_val = max((max(vals) if vals else 0) for _, vals, _ in series)
        min_val = min((min(vals) if vals else 0) for _, vals, _ in series)
        span = max(max_val - min_val, 1)
        plot_height = height - 2 * margin
        plot_width = width - 2 * margin
        x_step = plot_width / max(len(labels) - 1, 1)
        canvas.create_line(margin, height - margin, width - margin, height - margin, fill="#9ca3af")
        canvas.create_line(margin, margin, margin, height - margin, fill="#9ca3af")
        for i, label in enumerate(labels):
            x = margin + i * x_step
            canvas.create_text(x, height - margin + 12, text=label, anchor="n", font=("Segoe UI", 8))
        for idx, (name, vals, color) in enumerate(series):
            if not vals:
                continue
            points = []
            for i, val in enumerate(vals):
                x = margin + i * x_step
                y = height - margin - ((val - min_val) / span * plot_height)
                points.extend([x, y])
            canvas.create_line(points, fill=color, width=2, smooth=True)
            canvas.create_text(width - margin + 6, margin + 14 * idx, anchor="w", text=name, fill=color)
        canvas.create_text(margin, margin - 10, text=f"макс {max_val:.2f}", anchor="w", fill="#6b7280", font=("Segoe UI", 8))
        canvas.create_text(margin, height - margin + 4, text=f"мін {min_val:.2f}", anchor="w", fill="#6b7280", font=("Segoe UI", 8))

    def draw_pie_chart(self, canvas: tk.Canvas, data: list[tuple[str, float]], title: str) -> None:
        canvas.delete("all")
        width, height = self._canvas_size(canvas)
        total = sum(val for _, val in data)
        if total <= 0:
            canvas.create_text(width / 2, height / 2, text="Немає даних", fill="#6b7280")
            return
        radius = min(width, height) // 4
        cx, cy = width // 3, height // 2
        start_angle = 0.0
        for idx, (label, value) in enumerate(data):
            if value <= 0:
                continue
            extent = value / total * 360
            color = self.chart_palette[idx % len(self.chart_palette)]
            canvas.create_arc(cx - radius, cy - radius, cx + radius, cy + radius, start=start_angle, extent=extent, fill=color, outline="white")
            canvas.create_rectangle(width * 0.6, 20 + idx * 20, width * 0.6 + 12, 20 + idx * 20 + 12, fill=color, outline=color)
            canvas.create_text(width * 0.6 + 16, 20 + idx * 20 + 6, anchor="w", text=f"{label} ({value:.2f})")
            start_angle += extent
        canvas.create_text(cx, 14, text=title, font=("Segoe UI", 10, "bold"))

    def draw_bar_chart(
        self, canvas: tk.Canvas, data: list[tuple[str, float]], title: str = "", bar_color: str | None = None
    ) -> None:
        canvas.delete("all")
        width, height = self._canvas_size(canvas)
        if not data:
            canvas.create_text(width / 2, height / 2, text="Немає даних", fill="#6b7280")
            return
        max_val = max(abs(v) for _, v in data) or 1
        margin = 30
        bar_space = (width - 2 * margin) / max(len(data), 1)
        bar_width = bar_space * 0.6
        for idx, (label, value) in enumerate(data):
            color = bar_color or self.chart_palette[idx % len(self.chart_palette)]
            x0 = margin + idx * bar_space
            x1 = x0 + bar_width
            y_base = height - margin
            y_val = y_base - (abs(value) / max_val) * (height - 2 * margin)
            canvas.create_rectangle(x0, y_val, x1, y_base, fill=color, outline=color)
            canvas.create_text((x0 + x1) / 2, y_val - 8, text=f"{value:.2f}", font=("Segoe UI", 8))
            canvas.create_text((x0 + x1) / 2, height - margin + 12, text=label, font=("Segoe UI", 8))
        if title:
            canvas.create_text(margin, margin - 12, text=title, anchor="w", font=("Segoe UI", 10, "bold"))

    def draw_sparkline(self, canvas: tk.Canvas, values: list[float], color: str) -> None:
        canvas.delete("all")
        width, height = self._canvas_size(canvas)
        if not values:
            canvas.create_text(width / 2, height / 2, text="—", fill="#9ca3af")
            return
        max_val = max(values)
        min_val = min(values)
        span = max(max_val - min_val, 1)
        margin = 6
        plot_width = width - 2 * margin
        plot_height = height - 2 * margin
        x_step = plot_width / max(len(values) - 1, 1)
        points = []
        for i, v in enumerate(values):
            x = margin + i * x_step
            y = height - margin - ((v - min_val) / span * plot_height)
            points.extend([x, y])
        canvas.create_line(points, fill=color, width=2, smooth=True)

    # Export
    def create_export_tab(self) -> None:
        ttk.Label(self.export_frame, text="Експорт таблиць у CSV", font=("Segoe UI", 10, "bold")).pack(pady=10)
        tables = [
            ("Brands", "Бренди"),
            ("Categories", "Категорії"),
            ("Products", "Товари"),
            ("AdditionalProductCategories", "Додаткові категорії"),
            ("Counterparties", "Контрагенти"),
            ("Warehouses", "Склади"),
            ("SalesChannels", "Канали"),
            ("PurchaseDocuments", "Закупівлі"),
            ("PurchaseLines", "Рядки закупівель"),
            ("SalesDocuments", "Продажі"),
            ("SalesLines", "Рядки продажів"),
            ("StockBalances", "Залишки"),
            ("StockMoves", "Рухи товарів"),
            ("CashTransactions", "Каса"),
        ]
        for table, label in tables:
            ttk.Button(self.export_frame, text=f"Експорт {label}", command=lambda t=table: self.export_csv(t)).pack(pady=4)

    def export_csv(self, table: str) -> None:
        file_path = get_data_dir() / f"{table.lower()}_export.csv"
        try:
            db.export_table_to_csv(table, file_path)
            messagebox.showinfo("Експорт", f"Файл збережено: {file_path}")
        except Exception:
            logging.exception("Export error")
            show_error("Експорт", "Не вдалося експортувати таблицю.")

    # About
    def show_about(self) -> None:
        about_window = tk.Toplevel(self)
        about_window.title("Про програму")
        about_window.resizable(False, False)
        about_window.transient(self)
        about_window.grab_set()

        frame = ttk.Frame(about_window, padding=12)
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frame, text=f"{APP_NAME} v{VERSION}", font=("Segoe UI", 12, "bold")).pack(pady=(0, 8))
        ttk.Label(
            frame,
            text=(
                "Мінімалістичний офлайн-облік товарів і каси. "
                "Касовий метод, середньозважена собівартість, мультивалюта."
            ),
            wraplength=420,
            justify=tk.CENTER,
        ).pack(pady=4)

        info_frame = ttk.Frame(frame)
        info_frame.pack(pady=6, fill=tk.X)
        ttk.Label(info_frame, text="Шлях до БД:").grid(row=0, column=0, sticky=tk.W, padx=(0, 6))
        ttk.Label(info_frame, text=str(get_db_path()), wraplength=340, justify=tk.LEFT).grid(
            row=0, column=1, sticky=tk.W
        )
        ttk.Label(info_frame, text="Тека даних:").grid(row=1, column=0, sticky=tk.W, padx=(0, 6), pady=(4, 0))
        ttk.Label(info_frame, text=str(get_data_dir()), wraplength=340, justify=tk.LEFT).grid(
            row=1, column=1, sticky=tk.W, pady=(4, 0)
        )

        actions = ttk.Frame(frame)
        actions.pack(pady=(10, 0))
        ttk.Button(actions, text="Документація/FAQ", command=self.open_docs).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="Відкрити папку даних", command=lambda: open_data_folder(get_data_dir())).pack(
            side=tk.LEFT, padx=4
        )
        ttk.Button(actions, text="Закрити", command=about_window.destroy).pack(side=tk.LEFT, padx=4)

    # Refresh helpers
    def refresh_brands(self) -> None:
        rows = db.list_brands()
        self.brand_table.set_rows([{"id": r["id"], "name": r["name"]} for r in rows])

    def refresh_categories(self) -> None:
        rows = db.list_categories(include_hidden=True)
        self.categories_index = {int(r["id"]): dict(r) for r in rows}
        search = getattr(self, "category_search_var", tk.StringVar(value="")).get().lower().strip()
        category_products, product_quantities = db.category_inventory_data()

        children_map: dict[Optional[int], list[dict]] = {}
        for row in rows:
            children_map.setdefault(row["parent_id"], []).append(dict(row))

        for lst in children_map.values():
            lst.sort(key=lambda c: (c.get("sort_order", 0), c.get("name", "")))

        subtree_counts: dict[int, int] = {}
        subtree_quantities: dict[int, float] = {}

        def calc_total(cid: int) -> set[int]:
            products = set(category_products.get(cid, set()))
            for child in children_map.get(cid, []):
                products |= calc_total(child["id"])
            subtree_counts[cid] = len(products)
            subtree_quantities[cid] = sum(product_quantities.get(pid, 0.0) for pid in products)
            return products

        for root in children_map.get(None, []):
            calc_total(root["id"])

        match_cache: dict[int, bool] = {}

        def has_match(cid: int) -> bool:
            if cid in match_cache:
                return match_cache[cid]
            cat = self.categories_index.get(cid, {})
            own_match = not search or search in cat.get("name", "").lower()
            child_match = any(has_match(child["id"]) for child in children_map.get(cid, []))
            match_cache[cid] = bool(own_match or child_match)
            return match_cache[cid]

        self.category_tree.delete(*self.category_tree.get_children())

        def render(parent_id: Optional[int], tree_parent: str) -> None:
            for cat in children_map.get(parent_id, []):
                if not has_match(cat["id"]):
                    continue
                flags = []
                if cat.get("is_service"):
                    flags.append("службова")
                if cat.get("is_hidden"):
                    flags.append("прихована")
                node_id = self.category_tree.insert(
                    tree_parent,
                    "end",
                    iid=str(cat["id"]),
                    text=cat["name"],
                    values=(
                        subtree_counts.get(cat["id"], len(category_products.get(cat["id"], set()))),
                        f"{subtree_quantities.get(cat['id'], 0.0):.2f}",
                        ", ".join(flags),
                    ),
                )
                render(cat["id"], node_id)

        render(None, "")
        self.update_category_filters()

    def flatten_categories(self) -> list[dict]:
        rows = db.list_categories(include_hidden=False)
        children_map: dict[Optional[int], list[dict]] = {}
        for row in rows:
            children_map.setdefault(row["parent_id"], []).append(dict(row))
        for lst in children_map.values():
            lst.sort(key=lambda c: (c.get("sort_order", 0), c.get("name", "")))

        result: list[dict] = []

        def walk(parent_id: Optional[int], depth: int) -> None:
            for cat in children_map.get(parent_id, []):
                label = "  " * depth + ("• " if depth else "") + cat.get("name", "")
                result.append({"id": cat["id"], "label": label, "name": cat.get("name", "")})
                walk(cat["id"], depth + 1)

        walk(None, 0)
        return result

    def update_category_filters(self) -> None:
        self.category_filter_options = self.flatten_categories()
        values = ["Усі категорії"] + [c["label"] for c in self.category_filter_options]
        if hasattr(self, "product_category_filter_combo"):
            self.product_category_filter_combo.configure(values=values)
            current = self.product_category_filter_var.get()
            if current not in values:
                self.product_category_filter_combo.current(0)

    def refresh_products(self, search: str | None = None) -> None:
        category_id = None
        selected_label = self.product_category_filter_var.get()
        for option in getattr(self, "category_filter_options", []):
            if option["label"] == selected_label:
                category_id = option["id"]
                break

        rows = db.list_products(search, category_id, self.include_subcategories_var.get())
        prefix = _sanitize_barcode_prefix(self.settings.get("defaults", "product", "barcode_prefix") or "")
        self.product_table.set_rows(
            [
                {
                    "id": r["id"],
                    "sku": r["sku"],
                    "barcode": f"{prefix}{r['sku']}",
                    "supplier_sku": r["supplier_sku"] or "",
                    "name": r["name"],
                    "brand": r["brand"],
                    "category": r["category"],
                    "extra_categories": r["extra_categories"] or "",
                    "unit": r["unit"],
                    "is_active": "Так" if r["is_active"] else "Ні",
                }
                for r in rows
            ]
        )

    def refresh_counterparties(self) -> None:
        rows = db.list_counterparties()
        type_labels = {
            "supplier": "Постачальник",
            "customer": "Покупець",
            "both": "Постачальник/Покупець",
            "other": "Інший",
        }

        suppliers: list[dict] = []
        customers: list[dict] = []
        all_rows: list[dict] = []

        for r in rows:
            mapped = {
                "id": r["id"],
                "name": r["name"],
                "type": type_labels.get(r["type"], r["type"]),
                "phone": r["phone"] or "",
                "email": r["email"] or "",
                "address": r["address"] or "",
                "note": r["note"] or "",
            }
            all_rows.append(mapped)
            if r["type"] in ("supplier", "both"):
                suppliers.append(mapped)
            if r["type"] in ("customer", "both"):
                customers.append(mapped)

        self.counterparty_tables["suppliers"].set_rows(suppliers)
        self.counterparty_tables["customers"].set_rows(customers)
        self.counterparty_tables["all"].set_rows(all_rows)

    def refresh_warehouses(self) -> None:
        rows = db.list_warehouses()
        self.warehouse_table.set_rows(
            [
                {
                    "id": r["id"],
                    "name": r["name"],
                    "description": r["description"] or "",
                    "is_active": "Так" if r["is_active"] else "Ні",
                }
                for r in rows
            ]
        )

    def refresh_channels(self) -> None:
        rows = db.list_channels()
        self.channel_table.set_rows(
            [
                {
                    "id": r["id"],
                    "name": r["name"],
                    "is_active": "Так" if r["is_active"] else "Ні",
                }
                for r in rows
            ]
        )

    def refresh_all(self) -> None:
        self.refresh_brands()
        self.refresh_categories()
        self.refresh_products()
        self.refresh_counterparties()
        self.refresh_warehouses()
        self.refresh_channels()
        self.refresh_currencies()
        self.refresh_purchases()
        self.refresh_extra_costs()
        self.refresh_sales()
        self.refresh_inventory_documents()
        self.refresh_cash()
        self.refresh_stock()
        self.refresh_cash_counterparties()
        self.refresh_dashboard()
        self.refresh_sales_analysis()
        self.refresh_abc_xyz()
        self.refresh_cash_flow_report()


# Dialogs

def _parse_date_value(raw: str) -> str:
    raw = (raw or "").strip()
    if not raw:
        return datetime.now().strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return datetime.now().strftime("%Y-%m-%d")


def _parse_float_value(raw: str) -> float:
    if raw is None:
        return 0.0
    text = str(raw).replace(" ", "").replace(",", ".").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def _format_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).strftime("%Y-%m-%d")
    return str(value)


def _normalize_header(value: str) -> str:
    """Strip whitespace/BOM from column headers to avoid mapping typos."""

    return (value or "").strip().lstrip("\ufeff")


SalesField = tuple[str, str, tuple[str, ...]]


SALES_FIELDS: list[SalesField] = [
    ("order_no", "Замовлення", ("номер", "замовлення", "order", "order_id", "order_no", "id")),
    ("doc_date", "Дата", ("дата", "date", "order_date", "дата оформлення")),
    ("customer", "Клієнт", ("клієнт", "покупець", "customer", "контрагент")),
    ("phone", "Телефон", ("телефон", "phone")),
    ("email", "Email", ("email", "e-mail")),
    ("sku", "SKU", ("sku", "артикул", "код")),
    ("supplier_sku", "Артикул постачальника", ("артикул постачальника", "supplier_sku", "vendor_sku", "vendor code")),
    ("product_name", "Товар", ("товар", "product", "назва", "item")),
    ("quantity", "Кількість", ("кількість", "к-сть", "qty", "quantity", "шт")),
    ("price", "Ціна", ("ціна", "price", "amount")),
    ("amount", "Сума", ("сума", "amount", "total")),
    ("discount", "Знижка", ("знижка", "discount")),
    ("comment", "Коментар", ("коментар", "примітка", "comment", "note")),
    ("channel", "Канал", ("канал", "channel", "майданчик", "площадка", "platform")),
]


PURCHASE_FIELDS: list[SalesField] = [
    ("sku", "SKU", ("sku", "артикул", "код")),
    ("supplier_sku", "Артикул постачальника", ("артикул постачальника", "supplier_sku", "vendor_sku", "vendor code")),
    ("product_name", "Товар", ("товар", "product", "назва", "item")),
    ("quantity", "Кількість", ("кількість", "к-сть", "qty", "quantity", "шт")),
    ("price", "Ціна", ("ціна", "price", "amount")),
]


PRODUCT_FIELDS: list[SalesField] = [
    ("sku", "SKU", ("sku", "артикул", "код")),
    ("name", "Назва", ("назва", "name", "product", "товар", "item")),
    (
        "supplier_sku",
        "Артикул постачальника",
        ("артикул постачальника", "supplier_sku", "vendor_sku", "vendor code"),
    ),
    ("brand", "Бренд", ("бренд", "brand")),
    ("brand_id", "ID бренду", ("brand id", "id бренду")),
    ("category", "Категорія", ("категорія", "category")),
    ("category_id", "ID категорії", ("category id", "id категорії")),
    ("unit", "Одиниця", ("одиниця", "unit", "шт")),
    ("is_active", "Активний", ("активний", "is_active", "active")),
    (
        "extra_categories",
        "Додаткові категорії",
        ("додаткові категорії", "extra categories", "tags", "категорії"),
    ),
]


def _suggest_sales_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {h.lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for key, _label, aliases in SALES_FIELDS:
        for alias in aliases:
            if alias.lower() in normalized_headers:
                mapping[key] = normalized_headers[alias.lower()]
                break
        else:
            mapping[key] = ""
    return mapping


def _suggest_purchase_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {h.lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for key, _label, aliases in PURCHASE_FIELDS:
        for alias in aliases:
            if alias.lower() in normalized_headers:
                mapping[key] = normalized_headers[alias.lower()]
                break
        else:
            mapping[key] = ""
    return mapping


def _suggest_product_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {h.lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for key, _label, aliases in PRODUCT_FIELDS:
        for alias in aliases:
            if alias.lower() in normalized_headers:
                mapping[key] = normalized_headers[alias.lower()]
                break
        else:
            mapping[key] = ""
    return mapping


def _normalize_sales_records(rows: list[dict[str, object]], mapping: dict[str, str]) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        normalized = {_normalize_header(k): _format_cell_value(v).strip() for k, v in row.items()}

        def pick(field: str, parser=None):
            header = mapping.get(field, "")
            value = normalized.get(header, "") if header else ""
            return parser(value) if parser else value

        records.append(
            {
                "order_no": pick("order_no"),
                "doc_date": pick("doc_date", _parse_date_value),
                "customer": pick("customer"),
                "phone": pick("phone"),
                "email": pick("email"),
                "sku": pick("sku"),
                "supplier_sku": pick("supplier_sku"),
                "product_name": pick("product_name"),
                "quantity": pick("quantity", _parse_float_value),
                "price": pick("price", _parse_float_value),
                "amount": pick("amount", _parse_float_value),
                "discount": pick("discount", _parse_float_value),
                "comment": pick("comment"),
                "channel": pick("channel"),
            }
        )
    return records


def _parse_bool_value(raw: str) -> Optional[int]:
    text = (raw or "").strip().lower()
    if not text:
        return None
    if text in {"1", "true", "yes", "так", "y", "t", "on"}:
        return 1
    if text in {"0", "false", "no", "ні", "n", "off", "f"}:
        return 0
    return None


def _normalize_product_records(rows: list[dict[str, object]], mapping: dict[str, str]) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        normalized = {_normalize_header(k): _format_cell_value(v) for k, v in row.items()}

        def pick(field: str, parser=None):
            header = mapping.get(field, "")
            value = normalized.get(header, "") if header else ""
            return parser(value) if parser else (value.strip() if isinstance(value, str) else value)

        extra_raw = pick("extra_categories") or ""
        extra_items = []
        for part in re.split(r"[;,|]", extra_raw):
            clean = part.strip()
            if clean and clean not in extra_items:
                extra_items.append(clean)

        def parse_int(value: object) -> Optional[int]:
            try:
                num = int(str(value).strip())
                return num if num > 0 else None
            except Exception:
                return None

        records.append(
            {
                "sku": pick("sku"),
                "name": pick("name"),
                "supplier_sku": pick("supplier_sku"),
                "brand": pick("brand"),
                "brand_id": parse_int(pick("brand_id")),
                "category": pick("category"),
                "category_id": parse_int(pick("category_id")),
                "unit": pick("unit"),
                "is_active": _parse_bool_value(pick("is_active")),
                "extra_categories": extra_items,
            }
        )
    return records


def _normalize_purchase_records(
    rows: list[dict[str, object]],
    mapping: dict[str, str],
    default_supplier: str | None = None,
    default_doc_date: str | None = None,
    default_order_no: str | None = None,
    default_comment: str | None = None,
) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        normalized = {_normalize_header(k): _format_cell_value(v).strip() for k, v in row.items()}

        def pick(field: str, parser=None):
            header = mapping.get(field, "")
            value = normalized.get(header, "") if header else ""
            return parser(value) if parser else value

        records.append(
            {
                "order_no": (default_order_no or "").strip(),
                "doc_date": default_doc_date or pick("doc_date", _parse_date_value),
                "supplier": default_supplier if default_supplier else pick("supplier"),
                "sku": pick("sku"),
                "supplier_sku": pick("supplier_sku"),
                "product_name": pick("product_name"),
                "quantity": pick("quantity", _parse_float_value),
                "price": pick("price", _parse_float_value),
                "amount": pick("amount", _parse_float_value),
                "comment": (default_comment or "").strip(),
            }
        )
    return records


def _read_import_csv(path: Path, encoding: str) -> tuple[list[dict[str, object]], list[str]]:
    with path.open("r", encoding=encoding, newline="") as f:
        sample = f.read(2048)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample) if sample else csv.excel
        except Exception:
            dialect = csv.excel
        reader = csv.DictReader(f, dialect=dialect)
        raw_headers = list(reader.fieldnames or [])
        normalized_headers = [_normalize_header(h or "") for h in raw_headers]
        header_map = {raw or "": normalized for raw, normalized in zip(raw_headers, normalized_headers)}

        rows = []
        for row in reader:
            cleaned_row: dict[str, object] = {}
            for raw_key, value in row.items():
                normalized_key = header_map.get(raw_key or "", _normalize_header(raw_key or ""))
                cleaned_row[normalized_key] = value
            rows.append(cleaned_row)

        return rows, normalized_headers


def _read_import_xlsx(path: Path) -> tuple[list[dict[str, object]], list[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [_normalize_header(_format_cell_value(cell)) for cell in rows[0]]
    records: list[dict[str, object]] = []
    for row in rows[1:]:
        record: dict[str, object] = {}
        for idx, value in enumerate(row):
            header = headers[idx] if idx < len(headers) else ""
            record[header] = value
        records.append(record)
    return records, headers


def _read_import_xls(path: Path) -> tuple[list[dict[str, object]], list[str]]:
    try:
        # Деякі сервіси експортують XLSX-файли з розширенням .xls, тому
        # спершу пробуємо прочитати їх через openpyxl.
        return _read_import_xlsx(path)
    except Exception:
        # Якщо це справді старий XLS, повертаємося до xlrd.
        try:
            import xlrd
        except ImportError as exc:
            raise ImportError(
                "Для імпорту XLS-файлів потрібно встановити залежність 'xlrd'."
            ) from exc

    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return [], []

    headers = [_normalize_header(_format_cell_value(sheet.cell_value(0, col))) for col in range(sheet.ncols)]
    records: list[dict[str, object]] = []
    for row_idx in range(1, sheet.nrows):
        record: dict[str, object] = {}
        for col_idx in range(sheet.ncols):
            header = headers[col_idx] if col_idx < len(headers) else ""
            cell = sheet.cell(row_idx, col_idx)
            value: object = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                except Exception:
                    pass
            record[header] = value
        records.append(record)
    return records, headers


def parse_import_file(path: Path, encoding: str = "utf-8") -> tuple[list[dict[str, object]], list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        raw_rows, headers = _read_import_xlsx(path)
    elif suffix == ".xls":
        raw_rows, headers = _read_import_xls(path)
    else:
        raw_rows, headers = _read_import_csv(path, encoding)

    return raw_rows, headers


def parse_sales_file(path: Path, encoding: str = "utf-8") -> tuple[list[dict[str, object]], list[str]]:
    return parse_import_file(path, encoding)


class ProductsImportDialog(tk.Toplevel):
    def __init__(self, app: tk.Tk, raw_rows: list[dict[str, object]], headers: list[str], settings) -> None:
        super().__init__(app)
        self.title("Імпорт товарів")
        self.resizable(True, True)
        self.grab_set()
        self.result: Optional[dict] = None
        self.raw_rows = raw_rows
        self.headers = headers
        self.settings = settings
        self.templates: dict[str, dict[str, str]] = settings.get("product_import", "templates") or {}
        self.current_mapping = _suggest_product_mapping(headers)

        main = ttk.Frame(self, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        info = ttk.Label(main, text=f"Рядків у файлі: {len(raw_rows)}")
        info.grid(row=0, column=0, columnspan=3, sticky="w")

        self._build_template_controls(main)
        self._build_mapping_controls(main)
        self._build_options(main)
        self.preview = self._build_preview(main)
        self._refresh_preview()

        btns = ttk.Frame(main)
        btns.grid(row=11, column=0, columnspan=3, pady=8, sticky="e")
        ttk.Button(btns, text="Скасувати", command=self.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btns, text="Імпортувати", command=self._on_ok).pack(side=tk.RIGHT, padx=4)

        self.bind("<Return>", lambda _e: self._on_ok())
        self.bind("<Escape>", lambda _e: self.destroy())
        self.wait_window(self)

    def _on_ok(self) -> None:
        normalized_rows = _normalize_product_records(self.raw_rows, self.current_mapping)
        self.result = {
            "mapping": dict(self.current_mapping),
            "options": {
                "mode": self.mode_var.get(),
                "create_missing": bool(self.create_missing_var.get()),
                "extra_categories_mode": self.extra_categories_mode.get(),
                "update_name": bool(self.update_name_var.get()),
            },
            "rows": normalized_rows,
        }
        self.settings.set(self.current_template_name.get(), "product_import", "last_template")
        self.settings.save()
        self.destroy()

    def _build_template_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Шаблон співставлення:").grid(row=1, column=0, sticky="w", pady=4)
        self.current_template_name = tk.StringVar(value=self.settings.get("product_import", "last_template") or "")
        self.template_combo = ttk.Combobox(
            parent, textvariable=self.current_template_name, values=list(self.templates.keys()), state="readonly"
        )
        self.template_combo.grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Button(parent, text="Застосувати", command=self._apply_template).grid(row=1, column=2, padx=4, sticky="w")
        ttk.Button(parent, text="Зберегти", command=self._save_template).grid(row=1, column=3, padx=4, sticky="w")
        ttk.Button(parent, text="Видалити", command=self._delete_template).grid(row=1, column=4, padx=4, sticky="w")

    def _build_mapping_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Співставлення колонок:").grid(row=2, column=0, sticky="nw", pady=4)
        mapping_frame = ttk.Frame(parent)
        mapping_frame.grid(row=2, column=1, columnspan=4, sticky="ew", pady=4)
        mapping_frame.columnconfigure(1, weight=1)

        options = ["(не використовувати)"] + self.headers
        self.mapping_vars: dict[str, tk.StringVar] = {}
        for idx, (field_key, field_label, _aliases) in enumerate(PRODUCT_FIELDS):
            ttk.Label(mapping_frame, text=field_label).grid(row=idx, column=0, sticky="w", pady=2)
            var = tk.StringVar(value=self.current_mapping.get(field_key, ""))
            combo = ttk.Combobox(mapping_frame, textvariable=var, values=options, state="readonly")
            combo.grid(row=idx, column=1, sticky="ew", pady=2)
            combo.bind("<<ComboboxSelected>>", lambda _e, key=field_key, v=var: self._update_mapping(key, v.get()))
            self.mapping_vars[field_key] = var

    def _build_options(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Режим імпорту:").grid(row=3, column=0, sticky="nw", pady=4)
        mode_frame = ttk.Frame(parent)
        mode_frame.grid(row=3, column=1, sticky="w", pady=4)
        self.mode_var = tk.StringVar(value="create")
        ttk.Radiobutton(mode_frame, text="Створювати", variable=self.mode_var, value="create").pack(anchor="w")
        ttk.Radiobutton(mode_frame, text="Оновлювати", variable=self.mode_var, value="update").pack(anchor="w")
        ttk.Radiobutton(mode_frame, text="Додавати/оновлювати", variable=self.mode_var, value="upsert").pack(anchor="w")

        self.create_missing_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            parent,
            text="Створювати відсутні бренди/категорії",
            variable=self.create_missing_var,
        ).grid(row=4, column=1, sticky="w", pady=4)

        ttk.Label(parent, text="Додаткові категорії:").grid(row=5, column=0, sticky="nw", pady=4)
        extra_frame = ttk.Frame(parent)
        extra_frame.grid(row=5, column=1, sticky="w")
        self.extra_categories_mode = tk.StringVar(value="none")
        ttk.Radiobutton(extra_frame, text="Не чіпати", variable=self.extra_categories_mode, value="none").pack(anchor="w")
        ttk.Radiobutton(extra_frame, text="Додати", variable=self.extra_categories_mode, value="add").pack(anchor="w")
        ttk.Radiobutton(extra_frame, text="Замінити", variable=self.extra_categories_mode, value="replace").pack(anchor="w")

        self.update_name_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(parent, text="Оновлювати назву товару", variable=self.update_name_var).grid(
            row=6, column=1, sticky="w", pady=4
        )

    def _build_preview(self, parent: ttk.Frame) -> ttk.Treeview:
        ttk.Label(parent, text="Попередній перегляд (перші 50 рядків):").grid(row=7, column=0, columnspan=4, sticky="w", pady=6)
        preview = ttk.Treeview(
            parent,
            columns=("sku", "name", "brand", "category", "unit", "active", "extras"),
            show="headings",
            height=12,
        )
        headings = {
            "sku": ("SKU", 120),
            "name": ("Назва", 180),
            "brand": ("Бренд", 140),
            "category": ("Категорія", 140),
            "unit": ("Одиниця", 90),
            "active": ("Активний", 90),
            "extras": ("Додаткові категорії", 220),
        }
        for col, (title, width) in headings.items():
            preview.heading(col, text=title)
            preview.column(col, width=width, anchor="w")
        preview.grid(row=8, column=0, columnspan=4, sticky="nsew")
        parent.grid_rowconfigure(8, weight=1)
        parent.grid_columnconfigure(1, weight=1)
        scroll = ttk.Scrollbar(parent, orient="vertical", command=preview.yview)
        preview.configure(yscrollcommand=scroll.set)
        scroll.grid(row=8, column=4, sticky="ns")
        return preview

    def _refresh_preview(self) -> None:
        self.preview.delete(*self.preview.get_children())
        normalized = _normalize_product_records(self.raw_rows, self.current_mapping)
        for row in normalized[:50]:
            self.preview.insert(
                "",
                "end",
                values=(
                    row.get("sku"),
                    row.get("name"),
                    row.get("brand"),
                    row.get("category"),
                    row.get("unit"),
                    row.get("is_active"),
                    "; ".join(row.get("extra_categories") or []),
                ),
            )

    def _update_mapping(self, key: str, value: str) -> None:
        clean_value = "" if value == "(не використовувати)" else value
        self.current_mapping[key] = clean_value
        self._refresh_preview()

    def _apply_template(self) -> None:
        name = self.current_template_name.get().strip()
        if not name or name not in self.templates:
            return
        template = self.templates[name]
        for key, var in self.mapping_vars.items():
            var.set(template.get(key, ""))
            self.current_mapping[key] = template.get(key, "")
        self._refresh_preview()

    def _save_template(self) -> None:
        values = simple_prompt(
            "Назва шаблону",
            ["Вкажіть назву шаблону"],
            [self.current_template_name.get().strip()],
        )
        if not values:
            return

        name = values[0].strip()
        if not name:
            return

        self.templates[name] = dict(self.current_mapping)
        self.settings.set(self.templates, "product_import", "templates")
        self.settings.set(name, "product_import", "last_template")
        self.settings.save()
        self.current_template_name.set(name)
        self.template_combo.configure(values=list(self.templates.keys()))

    def _delete_template(self) -> None:
        name = self.current_template_name.get().strip()
        if not name or name not in self.templates:
            return
        if not messagebox.askyesno("Шаблони", f"Видалити шаблон '{name}'?"):
            return
        self.templates.pop(name, None)
        self.settings.set(self.templates, "product_import", "templates")
        if self.settings.get("product_import", "last_template") == name:
            self.settings.set("", "product_import", "last_template")
        self.settings.save()
        self.template_combo.configure(values=list(self.templates.keys()))
        if self.templates:
            self.current_template_name.set(list(self.templates.keys())[0])
        else:
            self.current_template_name.set("")

class PurchasesImportDialog(tk.Toplevel):
    def __init__(
        self,
        app: tk.Tk,
        raw_rows: list[dict[str, object]],
        headers: list[str],
        warehouses,
        suppliers,
        currencies,
        settings,
    ) -> None:
        super().__init__(app)
        self.title("Імпорт закупівель")
        self.resizable(True, True)
        self.grab_set()
        self.result: Optional[dict] = None
        self.raw_rows = raw_rows
        self.headers = headers
        self.warehouses = warehouses
        self.suppliers = suppliers
        self.supplier_names = list(dict.fromkeys([s["name"] for s in suppliers if s.get("name")]))
        self.currencies = currencies
        self.currency_codes = [c["code"] for c in currencies] if currencies else [get_base_currency_code()]
        self.settings = settings
        self.templates: dict[str, dict[str, str]] = settings.get("purchase_import", "templates") or {}
        self.current_mapping = _suggest_purchase_mapping(headers)

        main = ttk.Frame(self, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        info = ttk.Label(main, text=f"Рядків у файлі: {len(raw_rows)}")
        info.grid(row=0, column=0, columnspan=3, sticky="w")

        self._build_template_controls(main)
        self._build_mapping_controls(main)
        self._build_options(main)
        self.preview = self._build_preview(main)
        self._refresh_preview()

        btns = ttk.Frame(main)
        btns.grid(row=14, column=0, columnspan=3, pady=8, sticky="e")
        ttk.Button(btns, text="Скасувати", command=self.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btns, text="Імпортувати", command=self._on_ok).pack(side=tk.RIGHT, padx=4)

        self.bind("<Return>", lambda _e: self._on_ok())
        self.bind("<Escape>", lambda _e: self.destroy())
        self.wait_window(self)

    def _on_ok(self) -> None:
        warehouse = next((w for w in self.warehouses if w["name"] == self.wh_var.get()), None)
        supplier = next((s for s in self.suppliers if s["name"] == self.supplier_var.get()), None)
        if not warehouse:
            show_error("Імпорт", "Оберіть склад")
            return
        if not supplier:
            show_error("Імпорт", "Оберіть постачальника")
            return

        try:
            exchange_rate = float(self.rate_var.get())
        except ValueError:
            show_error("Імпорт", "Курс має бути числом")
            return
        if exchange_rate <= 0:
            show_error("Імпорт", "Курс має бути більшим за 0")
            return

        normalized_orders = _normalize_purchase_records(
            self.raw_rows,
            self.current_mapping,
            default_supplier=self.supplier_var.get().strip(),
            default_doc_date=self.date_picker.get(),
            default_order_no=self.order_no_var.get(),
            default_comment=self.comment_var.get(),
        )
        self.result = {
            "options": {
                "warehouse_id": warehouse["id"],
                "supplier_id": supplier["id"],
                "mode": self.mode_var.get(),
                "create_products": bool(self.create_products_var.get()),
                "currency_code": self.currency_var.get(),
                "exchange_rate": exchange_rate,
            },
            "orders": normalized_orders,
        }
        self.settings.set(self.current_template_name.get(), "purchase_import", "last_template")
        self.settings.save()
        self.destroy()

    def _build_template_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Шаблон співставлення:").grid(row=1, column=0, sticky="w", pady=4)
        self.current_template_name = tk.StringVar(value=self.settings.get("purchase_import", "last_template") or "")
        self.template_combo = ttk.Combobox(
            parent, textvariable=self.current_template_name, values=list(self.templates.keys()), state="readonly"
        )
        self.template_combo.grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Button(parent, text="Застосувати", command=self._apply_template).grid(row=1, column=2, padx=4, sticky="w")
        ttk.Button(parent, text="Зберегти як…", command=self._save_template).grid(row=1, column=3, padx=4, sticky="w")

    def _build_mapping_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Співставлення колонок:").grid(row=2, column=0, sticky="nw", pady=4)
        mapping_frame = ttk.Frame(parent)
        mapping_frame.grid(row=2, column=1, columnspan=3, sticky="ew", pady=4)
        mapping_frame.columnconfigure(1, weight=1)

        options = ["(не використовувати)"] + self.headers
        self.mapping_vars: dict[str, tk.StringVar] = {}
        for idx, (field_key, field_label, _aliases) in enumerate(PURCHASE_FIELDS):
            ttk.Label(mapping_frame, text=field_label).grid(row=idx, column=0, sticky="w", pady=2)
            var = tk.StringVar(value=self.current_mapping.get(field_key, ""))
            combo = ttk.Combobox(mapping_frame, textvariable=var, values=options, state="readonly")
            combo.grid(row=idx, column=1, sticky="ew", pady=2)
            combo.bind("<<ComboboxSelected>>", lambda _e, key=field_key, v=var: self._update_mapping(key, v.get()))
            self.mapping_vars[field_key] = var

    def _build_options(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Склад для імпорту:").grid(row=3, column=0, sticky="w", pady=4)
        self.wh_var = tk.StringVar(value=self.warehouses[0]["name"] if self.warehouses else "")
        wh_combo = ttk.Combobox(parent, textvariable=self.wh_var, values=[w["name"] for w in self.warehouses], state="readonly")
        wh_combo.grid(row=3, column=1, sticky="ew", pady=4)

        ttk.Label(parent, text="Постачальник:").grid(row=4, column=0, sticky="w", pady=4)
        self.supplier_var = tk.StringVar(value=self.supplier_names[0] if self.supplier_names else "")
        supplier_combo = ttk.Combobox(parent, textvariable=self.supplier_var, values=self.supplier_names, state="readonly")
        supplier_combo.grid(row=4, column=1, sticky="ew", pady=4)

        ttk.Label(parent, text="Дата документа:").grid(row=5, column=0, sticky="w", pady=4)
        self.date_picker = DatePicker(parent)
        self.date_picker.grid(row=5, column=1, sticky="w", pady=4)

        ttk.Label(parent, text="Валюта:").grid(row=6, column=0, sticky="w", pady=4)
        default_currency = self.currency_codes[0] if self.currency_codes else get_base_currency_code()
        self.currency_var = tk.StringVar(value=default_currency)
        currency_combo = ttk.Combobox(parent, textvariable=self.currency_var, values=self.currency_codes, state="readonly")
        currency_combo.grid(row=6, column=1, sticky="ew", pady=4)

        ttk.Label(parent, text="Курс:").grid(row=7, column=0, sticky="w", pady=4)
        try:
            default_rate = ensure_rate_for_date(self.currency_var.get(), self.date_picker.get())
        except Exception:
            default_rate = 1.0
        self.rate_var = tk.StringVar(value=f"{default_rate:.4f}")
        ttk.Entry(parent, textvariable=self.rate_var, width=14).grid(row=7, column=1, sticky="w", pady=4)

        def on_currency_change(_event=None):
            try:
                rate_val = ensure_rate_for_date(self.currency_var.get(), self.date_picker.get())
            except ValueError as exc:
                messagebox.showerror("Курс", str(exc))
                self.currency_var.set(default_currency)
                return
            self.rate_var.set(f"{rate_val:.4f}")

        currency_combo.bind("<<ComboboxSelected>>", on_currency_change)

        ttk.Label(parent, text="Замовлення/рахунок:").grid(row=8, column=0, sticky="w", pady=4)
        self.order_no_var = tk.StringVar()
        ttk.Entry(parent, textvariable=self.order_no_var).grid(row=8, column=1, sticky="ew", pady=4)

        ttk.Label(parent, text="Коментар:").grid(row=9, column=0, sticky="w", pady=4)
        self.comment_var = tk.StringVar()
        ttk.Entry(parent, textvariable=self.comment_var).grid(row=9, column=1, sticky="ew", pady=4)

        ttk.Label(parent, text="Режим проведення:").grid(row=10, column=0, sticky="nw", pady=4)
        mode_frame = ttk.Frame(parent)
        mode_frame.grid(row=10, column=1, sticky="w", pady=4)
        self.mode_var = tk.StringVar(value="post")
        ttk.Radiobutton(mode_frame, text="Провести всі", variable=self.mode_var, value="post").pack(anchor="w")
        ttk.Radiobutton(mode_frame, text="Тільки чернетки", variable=self.mode_var, value="draft").pack(anchor="w")

        self.create_products_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(parent, text="Створювати відсутні товари", variable=self.create_products_var).grid(
            row=11, column=1, sticky="w", pady=(4, 0)
        )

    def _build_preview(self, parent: ttk.Frame) -> ttk.Treeview:
        ttk.Label(parent, text="Попередній перегляд (перші 30 рядків):").grid(row=12, column=0, columnspan=3, sticky="w", pady=6)
        preview = ttk.Treeview(
            parent,
            columns=("order", "date", "supplier", "sku", "name", "qty", "price"),
            show="headings",
            height=10,
        )
        headings = {
            "order": ("Замовлення", 120),
            "date": ("Дата", 90),
            "supplier": ("Постачальник", 180),
            "sku": ("SKU", 90),
            "name": ("Товар", 200),
            "qty": ("К-сть", 70),
            "price": ("Ціна", 90),
        }
        for col, (title, width) in headings.items():
            preview.heading(col, text=title)
            preview.column(col, width=width, anchor="w")
        preview.grid(row=13, column=0, columnspan=3, sticky="nsew")
        parent.grid_rowconfigure(13, weight=1)
        parent.grid_columnconfigure(1, weight=1)
        scroll = ttk.Scrollbar(parent, orient="vertical", command=preview.yview)
        preview.configure(yscrollcommand=scroll.set)
        scroll.grid(row=13, column=3, sticky="ns")
        return preview

    def _refresh_preview(self) -> None:
        self.preview.delete(*self.preview.get_children())
        normalized = _normalize_purchase_records(
            self.raw_rows,
            self.current_mapping,
            default_supplier=self.supplier_var.get().strip(),
            default_doc_date=self.date_picker.get(),
            default_order_no=self.order_no_var.get(),
            default_comment=self.comment_var.get(),
        )
        for row in normalized[:30]:
            self.preview.insert(
                "",
                "end",
                values=(
                    row.get("order_no") or "-",
                    row.get("doc_date"),
                    row.get("supplier"),
                    row.get("sku"),
                    row.get("product_name"),
                    f"{float(row.get('quantity') or 0):.2f}",
                    f"{float(row.get('price') or 0):.2f}",
                ),
            )

    def _update_mapping(self, key: str, value: str) -> None:
        clean_value = "" if value == "(не використовувати)" else value
        self.current_mapping[key] = clean_value
        self._refresh_preview()

    def _apply_template(self) -> None:
        name = self.current_template_name.get().strip()
        if not name or name not in self.templates:
            return
        template = self.templates[name]
        for key, var in self.mapping_vars.items():
            var.set(template.get(key, ""))
            self.current_mapping[key] = template.get(key, "")
        self._refresh_preview()

    def _save_template(self) -> None:
        values = simple_prompt(
            "Назва шаблону",
            ["Вкажіть назву шаблону"],
            [self.current_template_name.get().strip()],
        )
        if not values:
            return

        name = values[0].strip()
        if not name:
            return

        self.templates[name] = dict(self.current_mapping)
        self.settings.set(self.templates, "purchase_import", "templates")
        self.settings.set(name, "purchase_import", "last_template")
        self.settings.save()
        self.current_template_name.set(name)
        self.template_combo.configure(values=list(self.templates.keys()))


class SalesImportDialog(tk.Toplevel):
    def __init__(self, app: tk.Tk, raw_rows: list[dict[str, object]], headers: list[str], warehouses, channels, settings) -> None:
        super().__init__(app)
        self.title("Імпорт продажів")
        self.resizable(True, True)
        self.grab_set()
        self.result: Optional[dict] = None
        self.raw_rows = raw_rows
        self.headers = headers
        self.warehouses = warehouses
        self.channels = channels
        self.settings = settings
        self.templates: dict[str, dict[str, str]] = settings.get("sales_import", "templates") or {}
        self.current_mapping = _suggest_sales_mapping(headers)

        main = ttk.Frame(self, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        info = ttk.Label(main, text=f"Рядків у файлі: {len(raw_rows)}")
        info.grid(row=0, column=0, columnspan=3, sticky="w")

        self._build_template_controls(main)
        self._build_mapping_controls(main)
        self._build_options(main)
        self.preview = self._build_preview(main)
        self._refresh_preview()

        btns = ttk.Frame(main)
        btns.grid(row=12, column=0, columnspan=3, pady=8, sticky="e")
        ttk.Button(btns, text="Скасувати", command=self.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btns, text="Імпортувати", command=self._on_ok).pack(side=tk.RIGHT, padx=4)

        self.bind("<Return>", lambda _e: self._on_ok())
        self.bind("<Escape>", lambda _e: self.destroy())
        self.wait_window(self)

    def _on_ok(self) -> None:
        warehouse = next((w for w in self.warehouses if w["name"] == self.wh_var.get()), None)
        if not warehouse:
            show_error("Імпорт", "Оберіть склад")
            return

        normalized_orders = _normalize_sales_records(self.raw_rows, self.current_mapping)
        self.result = {
            "options": {
                "warehouse_id": warehouse["id"],
                "channel": self.channel_var.get().strip(),
                "mode": self.mode_var.get(),
                "allow_negative": bool(self.allow_negative_var.get()),
                "create_products": bool(self.create_products_var.get()),
                "create_customers": bool(self.create_customers_var.get()),
                "use_file_channel": bool(self.use_file_channel_var.get()),
            },
            "orders": normalized_orders,
        }
        self.settings.set(self.current_template_name.get(), "sales_import", "last_template")
        self.settings.save()
        self.destroy()

    def _build_template_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Шаблон співставлення:").grid(row=1, column=0, sticky="w", pady=4)
        self.current_template_name = tk.StringVar(value=self.settings.get("sales_import", "last_template") or "")
        self.template_combo = ttk.Combobox(
            parent, textvariable=self.current_template_name, values=list(self.templates.keys()), state="readonly"
        )
        self.template_combo.grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Button(parent, text="Застосувати", command=self._apply_template).grid(row=1, column=2, padx=4, sticky="w")
        ttk.Button(parent, text="Зберегти як…", command=self._save_template).grid(row=1, column=3, padx=4, sticky="w")

    def _build_mapping_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Співставлення колонок:").grid(row=2, column=0, sticky="nw", pady=4)
        mapping_frame = ttk.Frame(parent)
        mapping_frame.grid(row=2, column=1, columnspan=3, sticky="ew", pady=4)
        mapping_frame.columnconfigure(1, weight=1)

        options = ["(не використовувати)"] + self.headers
        self.mapping_vars: dict[str, tk.StringVar] = {}
        for idx, (field_key, field_label, _aliases) in enumerate(SALES_FIELDS):
            ttk.Label(mapping_frame, text=field_label).grid(row=idx, column=0, sticky="w", pady=2)
            var = tk.StringVar(value=self.current_mapping.get(field_key, ""))
            combo = ttk.Combobox(mapping_frame, textvariable=var, values=options, state="readonly")
            combo.grid(row=idx, column=1, sticky="ew", pady=2)
            combo.bind("<<ComboboxSelected>>", lambda _e, key=field_key, v=var: self._update_mapping(key, v.get()))
            self.mapping_vars[field_key] = var

    def _build_options(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="Склад для імпорту:").grid(row=3, column=0, sticky="w", pady=4)
        self.wh_var = tk.StringVar(value=self.warehouses[0]["name"] if self.warehouses else "")
        wh_combo = ttk.Combobox(parent, textvariable=self.wh_var, values=[w["name"] for w in self.warehouses], state="readonly")
        wh_combo.grid(row=3, column=1, sticky="ew", pady=4)

        ttk.Label(parent, text="Канал (якщо не вказано у файлі):").grid(row=4, column=0, sticky="w", pady=4)
        self.channel_var = tk.StringVar()
        channel_values = [c["name"] for c in self.channels]
        ttk.Combobox(parent, textvariable=self.channel_var, values=channel_values).grid(row=4, column=1, sticky="ew", pady=4)

        self.use_file_channel_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(parent, text="Брати канал із файлу, якщо він є", variable=self.use_file_channel_var).grid(
            row=5, column=1, sticky="w"
        )

        ttk.Label(parent, text="Режим проведення:").grid(row=6, column=0, sticky="nw", pady=4)
        mode_frame = ttk.Frame(parent)
        mode_frame.grid(row=6, column=1, sticky="w", pady=4)
        self.mode_var = tk.StringVar(value="post")
        ttk.Radiobutton(mode_frame, text="Провести всі", variable=self.mode_var, value="post").pack(anchor="w")
        ttk.Radiobutton(mode_frame, text="Тільки чернетки", variable=self.mode_var, value="draft").pack(anchor="w")
        ttk.Radiobutton(
            mode_frame, text="Проводити, лише якщо є залишок", variable=self.mode_var, value="in_stock"
        ).pack(anchor="w")

        self.allow_negative_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            parent,
            text="Дозволити від'ємний залишок під час проведення",
            variable=self.allow_negative_var,
        ).grid(row=7, column=1, sticky="w")

        self.create_products_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(parent, text="Створювати відсутні товари", variable=self.create_products_var).grid(
            row=8, column=1, sticky="w", pady=(4, 0)
        )
        self.create_customers_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(parent, text="Створювати відсутніх клієнтів", variable=self.create_customers_var).grid(
            row=9, column=1, sticky="w"
        )

    def _build_preview(self, parent: ttk.Frame) -> ttk.Treeview:
        ttk.Label(parent, text="Попередній перегляд (перші 30 рядків):").grid(row=10, column=0, columnspan=3, sticky="w", pady=6)
        preview = ttk.Treeview(
            parent,
            columns=("order", "date", "customer", "sku", "name", "qty", "price"),
            show="headings",
            height=10,
        )
        headings = {
            "order": ("Замовлення", 120),
            "date": ("Дата", 90),
            "customer": ("Клієнт", 160),
            "sku": ("SKU", 90),
            "name": ("Товар", 200),
            "qty": ("К-сть", 70),
            "price": ("Ціна", 90),
        }
        for col, (title, width) in headings.items():
            preview.heading(col, text=title)
            preview.column(col, width=width, anchor="w")
        preview.grid(row=11, column=0, columnspan=3, sticky="nsew")
        parent.grid_rowconfigure(11, weight=1)
        parent.grid_columnconfigure(1, weight=1)
        scroll = ttk.Scrollbar(parent, orient="vertical", command=preview.yview)
        preview.configure(yscrollcommand=scroll.set)
        scroll.grid(row=11, column=3, sticky="ns")
        return preview

    def _refresh_preview(self) -> None:
        self.preview.delete(*self.preview.get_children())
        normalized = _normalize_sales_records(self.raw_rows, self.current_mapping)
        for row in normalized[:30]:
            self.preview.insert(
                "",
                "end",
                values=(
                    row.get("order_no") or "-",
                    row.get("doc_date"),
                    row.get("customer"),
                    row.get("sku"),
                    row.get("product_name"),
                    f"{float(row.get('quantity') or 0):.2f}",
                    f"{float(row.get('price') or 0):.2f}",
                ),
            )

    def _update_mapping(self, key: str, value: str) -> None:
        clean_value = "" if value == "(не використовувати)" else value
        self.current_mapping[key] = clean_value
        self._refresh_preview()

    def _apply_template(self) -> None:
        name = self.current_template_name.get().strip()
        if not name or name not in self.templates:
            return
        template = self.templates[name]
        for key, var in self.mapping_vars.items():
            var.set(template.get(key, ""))
            self.current_mapping[key] = template.get(key, "")
        self._refresh_preview()

    def _save_template(self) -> None:
        values = simple_prompt(
            "Назва шаблону",
            ["Вкажіть назву шаблону"],
            [self.current_template_name.get().strip()],
        )
        if not values:
            return

        name = values[0].strip()
        if not name:
            return

        self.templates[name] = dict(self.current_mapping)
        self.settings.set(self.templates, "sales_import", "templates")
        self.settings.set(name, "sales_import", "last_template")
        self.settings.save()
        self.current_template_name.set(name)
        self.template_combo.configure(values=list(self.templates.keys()))


def _find_index_by_name(items: list[str], target: str | None) -> int | None:
    if not target:
        return None
    target_lower = target.lower()
    return next((i for i, name in enumerate(items) if str(name).lower() == target_lower), None)


def _sanitize_barcode_prefix(prefix: str) -> str:
    cleaned = (prefix or "").strip()
    if not cleaned:
        return ""
    return re.sub(r"\s+", "-", cleaned)


def product_prompt(brands, categories, title: str, initial=None, settings: Settings | None = None):
    base_initial = {
        "sku": "",
        "supplier_sku": "",
        "name": "",
        "brand_id": None,
        "category_id": None,
        "unit": (settings.get("defaults", "product", "unit") if settings else None) or "pcs",
        "is_active": True,
        "extras": [],
        "supplier_codes": [],
        "barcodes": [],
    }

    if isinstance(initial, dict):
        normalized_initial = {**base_initial, **initial}
    elif initial:
        # Support tuples/lists passed by the callers
        normalized_initial = base_initial.copy()
        normalized_initial.update(
            {
                "sku": initial[0] if len(initial) > 0 else "",
                "supplier_sku": initial[1] if len(initial) > 1 else "",
                "name": initial[2] if len(initial) > 2 else "",
                "brand_id": initial[3] if len(initial) > 3 else None,
                "category_id": initial[4] if len(initial) > 4 else None,
                "unit": initial[5] if len(initial) > 5 else "pcs",
                "is_active": bool(initial[6]) if len(initial) > 6 else True,
                "extras": initial[7] if len(initial) > 7 else [],
                "supplier_codes": initial[8] if len(initial) > 8 else [],
                "barcodes": initial[9] if len(initial) > 9 else [],
            }
        )
    else:
        normalized_initial = base_initial

    dlg = tk.Toplevel()
    dlg.title(title)
    dlg.grab_set()
    dlg.columnconfigure(1, weight=1)

    ttk.Label(dlg, text="SKU").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    sku_var = tk.StringVar(value=normalized_initial.get("sku", ""))
    ttk.Entry(dlg, textvariable=sku_var, width=30).grid(row=0, column=1, padx=6, pady=4, sticky="ew")

    ttk.Label(dlg, text="Артикул постачальника (legacy)").grid(row=1, column=0, padx=6, pady=4, sticky="w")
    supplier_sku_var = tk.StringVar(value=normalized_initial.get("supplier_sku", ""))
    ttk.Entry(dlg, textvariable=supplier_sku_var, width=30).grid(row=1, column=1, padx=6, pady=4, sticky="ew")

    ttk.Label(dlg, text="Назва").grid(row=2, column=0, padx=6, pady=4, sticky="w")
    name_var = tk.StringVar(value=normalized_initial.get("name", ""))
    ttk.Entry(dlg, textvariable=name_var, width=30).grid(row=2, column=1, padx=6, pady=4, sticky="ew")

    ttk.Label(dlg, text="Бренд").grid(row=3, column=0, padx=6, pady=4, sticky="w")
    brand_var = tk.StringVar()
    brand_names = [b["name"] for b in brands]
    brand_combo = ttk.Combobox(dlg, textvariable=brand_var, state="readonly", values=brand_names)
    brand_combo.grid(row=3, column=1, padx=6, pady=4, sticky="ew")

    ttk.Label(dlg, text="Головна категорія").grid(row=4, column=0, padx=6, pady=4, sticky="w")
    category_var = tk.StringVar()
    category_combo = ttk.Combobox(dlg, textvariable=category_var, values=[c["label"] for c in categories])
    category_combo.grid(row=4, column=1, padx=6, pady=4, sticky="ew")

    def refresh_category_options(*_args):
        search = category_var.get().strip().lower()
        filtered = [c["label"] for c in categories if search in c["label"].lower()]
        category_combo["values"] = filtered or [c["label"] for c in categories]

    category_combo.bind("<KeyRelease>", refresh_category_options)

    ttk.Label(dlg, text="Додаткові категорії").grid(row=5, column=0, padx=6, pady=4, sticky="nw")
    extras_frame = ttk.Frame(dlg)
    extras_frame.grid(row=5, column=1, padx=6, pady=4, sticky="nsew")
    extras_frame.columnconfigure(0, weight=1)
    extras_frame.rowconfigure(1, weight=1)
    extras_search_var = tk.StringVar()
    ttk.Entry(extras_frame, textvariable=extras_search_var).grid(
        row=0, column=0, columnspan=2, padx=(0, 8), pady=(0, 4), sticky="ew"
    )
    extras_box = tk.Listbox(
        extras_frame, selectmode=tk.MULTIPLE, height=min(10, max(6, len(categories))), exportselection=False
    )
    extras_scroll = ttk.Scrollbar(extras_frame, orient="vertical", command=extras_box.yview)
    extras_box.configure(yscrollcommand=extras_scroll.set)
    extras_box.grid(row=1, column=0, sticky="nsew")
    extras_scroll.grid(row=1, column=1, sticky="ns")

    filtered_extra_categories = list(categories)
    initial_extra_ids = set(normalized_initial.get("extras") or [])

    def refresh_extra_list(*_args):
        selected_labels = {extras_box.get(i) for i in extras_box.curselection()}
        search = extras_search_var.get().strip().lower()
        extras_box.delete(0, tk.END)
        filtered_extra_categories.clear()
        filtered_extra_categories.extend([c for c in categories if search in c["label"].lower()])
        for idx, cat in enumerate(filtered_extra_categories):
            extras_box.insert(tk.END, cat["label"])
            if cat["label"] in selected_labels or cat["id"] in initial_extra_ids:
                extras_box.selection_set(idx)
        initial_extra_ids.difference_update({c["id"] for c in filtered_extra_categories})

    extras_search_var.trace_add("write", refresh_extra_list)
    refresh_extra_list()

    dlg.rowconfigure(5, weight=1)

    ttk.Label(dlg, text="Одиниця").grid(row=6, column=0, padx=6, pady=4, sticky="w")
    unit_var = tk.StringVar(value=normalized_initial.get("unit", "pcs"))
    ttk.Entry(dlg, textvariable=unit_var, width=12).grid(row=6, column=1, padx=6, pady=4, sticky="w")

    is_active_var = tk.BooleanVar(value=normalized_initial.get("is_active", True))
    ttk.Checkbutton(dlg, text="Активний", variable=is_active_var).grid(row=7, column=1, padx=6, pady=4, sticky="w")

    suppliers = db.list_suppliers()
    supplier_names = [s["name"] for s in suppliers]
    supplier_name_by_id = {s["id"]: s["name"] for s in suppliers}
    supplier_codes_state: list[dict] = []
    for code in normalized_initial.get("supplier_codes") or []:
        try:
            supplier_codes_state.append(
                {
                    "supplier_id": int(code.get("supplier_id")),  # type: ignore[arg-type]
                    "supplier_sku": (code.get("supplier_sku") or "").strip(),
                    "is_primary": bool(code.get("is_primary")),
                }
            )
        except (TypeError, ValueError):
            continue

    barcode_prefix = _sanitize_barcode_prefix((settings.get("defaults", "product", "barcode_prefix") if settings else "") or "")
    barcodes_state: list[dict] = []
    for code in normalized_initial.get("barcodes") or []:
        raw = (code.get("code") or "").strip()
        if not raw:
            continue
        barcodes_state.append({"code": raw, "note": (code.get("note") or "").strip()})

    supplier_codes_frame = ttk.LabelFrame(dlg, text="Артикули постачальників")
    supplier_codes_frame.grid(row=8, column=0, columnspan=2, padx=6, pady=4, sticky="nsew")
    supplier_codes_frame.columnconfigure(1, weight=1)
    dlg.rowconfigure(8, weight=1)

    ttk.Label(supplier_codes_frame, text="Постачальник:").grid(row=0, column=0, padx=4, pady=2, sticky="w")
    supplier_var = tk.StringVar()
    supplier_combo = ttk.Combobox(
        supplier_codes_frame, textvariable=supplier_var, state="readonly", values=supplier_names, width=30
    )
    supplier_combo.grid(row=0, column=1, padx=4, pady=2, sticky="ew")

    ttk.Label(supplier_codes_frame, text="Артикул:").grid(row=1, column=0, padx=4, pady=2, sticky="w")
    supplier_code_var = tk.StringVar()
    ttk.Entry(supplier_codes_frame, textvariable=supplier_code_var, width=30).grid(
        row=1, column=1, padx=4, pady=2, sticky="ew"
    )

    supplier_primary_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(supplier_codes_frame, text="Основний", variable=supplier_primary_var).grid(
        row=1, column=2, padx=4, pady=2, sticky="w"
    )

    def refresh_supplier_codes_tree() -> None:
        supplier_codes_tree.delete(*supplier_codes_tree.get_children())
        for idx, code in enumerate(supplier_codes_state):
            supplier_name = supplier_name_by_id.get(code.get("supplier_id"), str(code.get("supplier_id")))
            supplier_codes_tree.insert(
                "",
                "end",
                iid=str(idx),
                values=(supplier_name, code.get("supplier_sku", ""), "Так" if code.get("is_primary") else ""),
            )

    def add_supplier_code() -> None:
        if not supplier_names:
            messagebox.showerror("Артикули постачальників", "Створіть постачальника зі статусом постачальника.")
            return
        try:
            supplier_idx = supplier_names.index(supplier_var.get())
        except ValueError:
            messagebox.showerror("Артикули постачальників", "Оберіть постачальника")
            return
        supplier_id = suppliers[supplier_idx]["id"]
        supplier_sku = supplier_code_var.get().strip()
        if not supplier_sku:
            messagebox.showerror("Артикули постачальників", "Введіть артикул постачальника")
            return
        key = (supplier_id, supplier_sku.lower())
        if any((c.get("supplier_id"), (c.get("supplier_sku") or "").lower()) == key for c in supplier_codes_state):
            messagebox.showerror("Артикули постачальників", "Такий артикул вже додано для цього постачальника")
            return
        is_primary = bool(supplier_primary_var.get())
        if is_primary:
            for c in supplier_codes_state:
                if c.get("supplier_id") == supplier_id:
                    c["is_primary"] = False
        supplier_codes_state.append(
            {"supplier_id": supplier_id, "supplier_sku": supplier_sku, "is_primary": is_primary}
        )
        refresh_supplier_codes_tree()
        supplier_code_var.set("")
        supplier_primary_var.set(False)

    def delete_supplier_code() -> None:
        selection = supplier_codes_tree.selection()
        if not selection:
            return
        idx = int(selection[0])
        if 0 <= idx < len(supplier_codes_state):
            supplier_codes_state.pop(idx)
        refresh_supplier_codes_tree()

    def mark_primary() -> None:
        selection = supplier_codes_tree.selection()
        if not selection:
            return
        idx = int(selection[0])
        if 0 <= idx < len(supplier_codes_state):
            supplier_id = supplier_codes_state[idx].get("supplier_id")
            for i, code in enumerate(supplier_codes_state):
                if code.get("supplier_id") == supplier_id:
                    code["is_primary"] = i == idx
        refresh_supplier_codes_tree()

    ttk.Button(supplier_codes_frame, text="Додати", command=add_supplier_code).grid(
        row=0, column=2, padx=4, pady=2, sticky="w"
    )

    columns = ("supplier", "sku", "primary")
    supplier_codes_tree = ttk.Treeview(
        supplier_codes_frame, columns=columns, show="headings", selectmode="browse", height=6
    )
    supplier_codes_tree.heading("supplier", text="Постачальник")
    supplier_codes_tree.heading("sku", text="Артикул")
    supplier_codes_tree.heading("primary", text="Основний")
    supplier_codes_tree.column("supplier", width=180, anchor="w")
    supplier_codes_tree.column("sku", width=120, anchor="w")
    supplier_codes_tree.column("primary", width=90, anchor="center")
    supplier_codes_tree.grid(row=2, column=0, columnspan=3, padx=4, pady=4, sticky="nsew")
    supplier_codes_frame.rowconfigure(2, weight=1)
    supplier_codes_frame.columnconfigure(0, weight=1)
    supplier_codes_frame.columnconfigure(1, weight=1)
    scroll = ttk.Scrollbar(supplier_codes_frame, orient="vertical", command=supplier_codes_tree.yview)
    supplier_codes_tree.configure(yscrollcommand=scroll.set)
    scroll.grid(row=2, column=3, sticky="ns")

    actions_frame = ttk.Frame(supplier_codes_frame)
    actions_frame.grid(row=3, column=0, columnspan=3, padx=4, pady=(0, 4), sticky="w")
    ttk.Button(actions_frame, text="Видалити вибраний", command=delete_supplier_code).pack(side=tk.LEFT, padx=4)
    ttk.Button(actions_frame, text="Зробити основним", command=mark_primary).pack(side=tk.LEFT, padx=4)

    if supplier_names:
        supplier_combo.current(0)
    refresh_supplier_codes_tree()

    barcodes_frame = ttk.LabelFrame(dlg, text="Додаткові штрихкоди")
    barcodes_frame.grid(row=9, column=0, columnspan=2, padx=6, pady=4, sticky="nsew")
    barcodes_frame.columnconfigure(1, weight=1)
    dlg.rowconfigure(9, weight=1)

    def _main_barcode_value() -> str:
        return f"{barcode_prefix}{sku_var.get().strip()}"

    main_barcode_label = ttk.Label(barcodes_frame, text=f"Основний (SKU): {_main_barcode_value()}")
    main_barcode_label.grid(row=0, column=0, columnspan=3, padx=4, pady=2, sticky="w")

    def _update_main_barcode(*_args) -> None:
        main_barcode_label.configure(text=f"Основний (SKU): {_main_barcode_value()}")

    sku_var.trace_add("write", _update_main_barcode)

    ttk.Label(barcodes_frame, text="Штрихкод:").grid(row=1, column=0, padx=4, pady=2, sticky="w")
    barcode_code_var = tk.StringVar()
    ttk.Entry(barcodes_frame, textvariable=barcode_code_var, width=22).grid(row=1, column=1, padx=4, pady=2, sticky="ew")

    ttk.Label(barcodes_frame, text="Нотатка:").grid(row=2, column=0, padx=4, pady=2, sticky="w")
    barcode_note_var = tk.StringVar()
    ttk.Entry(barcodes_frame, textvariable=barcode_note_var, width=22).grid(row=2, column=1, padx=4, pady=2, sticky="ew")

    def refresh_barcode_tree() -> None:
        barcode_tree.delete(*barcode_tree.get_children())
        for idx, code in enumerate(barcodes_state):
            barcode_tree.insert("", "end", iid=str(idx), values=(code.get("code", ""), code.get("note", "")))

    def add_barcode() -> None:
        code_val = barcode_code_var.get().strip()
        if not code_val:
            messagebox.showerror("Штрихкоди", "Введіть штрихкод")
            return
        normalized = code_val.lower()
        if any((c.get("code") or "").lower() == normalized for c in barcodes_state):
            messagebox.showerror("Штрихкоди", "Такий штрихкод вже додано")
            return
        barcodes_state.append({"code": code_val, "note": barcode_note_var.get().strip()})
        refresh_barcode_tree()
        barcode_code_var.set("")
        barcode_note_var.set("")

    def delete_barcode() -> None:
        selection = barcode_tree.selection()
        if not selection:
            return
        idx = int(selection[0])
        if 0 <= idx < len(barcodes_state):
            barcodes_state.pop(idx)
        refresh_barcode_tree()

    ttk.Button(barcodes_frame, text="Додати", command=add_barcode).grid(row=1, column=2, padx=4, pady=2, sticky="w")

    barcode_tree = ttk.Treeview(barcodes_frame, columns=("code", "note"), show="headings", selectmode="browse", height=5)
    barcode_tree.heading("code", text="Штрихкод")
    barcode_tree.heading("note", text="Нотатка")
    barcode_tree.column("code", width=180, anchor="w")
    barcode_tree.column("note", width=180, anchor="w")
    barcode_tree.grid(row=3, column=0, columnspan=3, padx=4, pady=4, sticky="nsew")
    barcodes_frame.rowconfigure(3, weight=1)
    barcodes_frame.columnconfigure(0, weight=1)
    barcodes_frame.columnconfigure(1, weight=1)
    barcode_scroll = ttk.Scrollbar(barcodes_frame, orient="vertical", command=barcode_tree.yview)
    barcode_tree.configure(yscrollcommand=barcode_scroll.set)
    barcode_scroll.grid(row=3, column=3, sticky="ns")

    ttk.Button(barcodes_frame, text="Видалити вибраний", command=delete_barcode).grid(
        row=4, column=0, columnspan=3, padx=4, pady=(0, 4), sticky="w"
    )

    refresh_barcode_tree()

    if initial:
        brand_combo.current(next((i for i, b in enumerate(brands) if b["id"] == normalized_initial["brand_id"]), 0))
        category_var.set(next((c["label"] for c in categories if c["id"] == normalized_initial["category_id"]), ""))
        refresh_category_options()
        extras = set(normalized_initial.get("extras") or [])
        for idx, cat in enumerate(filtered_extra_categories):
            if cat["id"] in extras:
                extras_box.selection_set(idx)
    else:
        preferred_brand = _find_index_by_name(brand_names, (settings.get("defaults", "product", "brand") if settings else ""))
        if preferred_brand is not None:
            brand_combo.current(preferred_brand)
        elif brands:
            brand_combo.current(0)
        if categories:
            default_category = (settings.get("defaults", "product", "category") if settings else "") or ""
            if default_category:
                category_var.set(default_category)
            else:
                category_var.set(categories[0]["label"])
            refresh_category_options()

    result = None

    def on_ok():
        nonlocal result
        sku = sku_var.get().strip()
        name = name_var.get().strip()
        if not sku or not name:
            messagebox.showerror("Валідація", "Заповніть SKU та назву")
            return
        if not categories:
            messagebox.showerror("Валідація", "Створіть принаймні одну категорію")
            return
        try:
            brand_id = brands[brand_combo.current()]["id"]
        except IndexError:
            messagebox.showerror("Валідація", "Оберіть бренд і категорію")
            return
        selected_label = category_var.get().strip()
        matched_category = next((c for c in categories if c["label"].lower() == selected_label.lower()), None)
        if not matched_category:
            matched_category = next((c for c in categories if selected_label.lower() in c["label"].lower()), None)
        if not matched_category:
            messagebox.showerror("Валідація", "Оберіть бренд і категорію")
            return
        selected = [filtered_extra_categories[i]["id"] for i in extras_box.curselection() if i < len(filtered_extra_categories)]
        result = (
            sku,
            supplier_sku_var.get().strip(),
            name,
            brand_id,
            matched_category["id"],
            unit_var.get().strip() or "pcs",
            bool(is_active_var.get()),
            selected,
            list(supplier_codes_state),
            list(barcodes_state),
        )
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    btns = ttk.Frame(dlg)
    btns.grid(row=10, column=0, columnspan=2, pady=8, sticky="e")
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: on_cancel())
    if editable:
        scan_entry.focus_set()
    dlg.wait_window()
    return result


def open_products_bulk_actions_dialog(parent, db_conn, table_frame=None, product_ids=None) -> None:
    if product_ids is None:
        if table_frame is None:
            messagebox.showwarning("Масові дії", "Оберіть хоча б один товар.")
            try:
                db_conn.close()
            except Exception:
                pass
            return
        product_ids = table_frame.get_selected_row_ids()
    if not product_ids:
        messagebox.showwarning("Масові дії", "Оберіть хоча б один товар.")
        try:
            db_conn.close()
        except Exception:
            pass
        return

    dlg = tk.Toplevel(parent)
    dlg.title("Масові дії з товарами")
    dlg.grab_set()
    dlg.resizable(False, False)

    ttk.Label(dlg, text=f"Обрано товарів: {len(product_ids)}").grid(
        row=0, column=0, columnspan=3, padx=8, pady=(8, 4), sticky="w"
    )

    brands = db.list_brands()
    brand_names = [b["name"] for b in brands]
    categories = parent.flatten_categories()
    category_names = [c["label"] for c in categories]

    default_unit = "pcs"
    try:
        default_unit = (parent.settings.get("defaults", "product", "unit") or default_unit).strip() or "pcs"
    except Exception:
        default_unit = "pcs"

    def close_dialog() -> None:
        try:
            db_conn.close()
        except Exception:
            pass
        dlg.destroy()

    def confirm_and_apply(action_label: str, func) -> None:
        if not messagebox.askyesno("Підтвердження", f"Застосувати до {len(product_ids)} товарів?"):
            return
        try:
            affected = func()
        except Exception:
            logging.exception("Bulk products action error")
            show_error("Масові дії", "Не вдалося виконати дію.")
            return
        messagebox.showinfo("Масові дії", f"{action_label}: {affected}")
        parent.refresh_products()

    # Activation
    act_frame = ttk.LabelFrame(dlg, text="Активація")
    act_frame.grid(row=1, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Button(act_frame, text="Активувати", command=lambda: confirm_and_apply(
        "Оновлено товарів",
        lambda: db.bulk_update_products_is_active(db_conn, product_ids, 1),
    )).pack(side=tk.LEFT, padx=4, pady=4)
    ttk.Button(act_frame, text="Деактивувати", command=lambda: confirm_and_apply(
        "Оновлено товарів",
        lambda: db.bulk_update_products_is_active(db_conn, product_ids, 0),
    )).pack(side=tk.LEFT, padx=4, pady=4)

    # Brand
    brand_frame = ttk.LabelFrame(dlg, text="Встановити бренд")
    brand_frame.grid(row=2, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Label(brand_frame, text="Бренд:").pack(side=tk.LEFT, padx=4, pady=4)
    brand_var = tk.StringVar()
    brand_combo = ttk.Combobox(brand_frame, textvariable=brand_var, state="readonly", values=brand_names, width=30)
    brand_combo.pack(side=tk.LEFT, padx=4, pady=4)
    if brand_names:
        brand_combo.current(0)

    def apply_brand() -> None:
        if not brands:
            messagebox.showwarning("Бренди", "Створіть принаймні один бренд.")
            return
        confirm_and_apply(
            "Оновлено товарів",
            lambda: db.bulk_update_products_brand(db_conn, product_ids, brands[brand_combo.current()]["id"]),
        )

    ttk.Button(brand_frame, text="Застосувати бренд", command=apply_brand).pack(side=tk.LEFT, padx=4, pady=4)

    # Category
    category_frame = ttk.LabelFrame(dlg, text="Встановити категорію")
    category_frame.grid(row=3, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Label(category_frame, text="Категорія:").pack(side=tk.LEFT, padx=4, pady=4)
    category_var = tk.StringVar()
    category_combo = ttk.Combobox(category_frame, textvariable=category_var, values=category_names, width=40)
    category_combo.pack(side=tk.LEFT, padx=4, pady=4)
    if category_names:
        category_combo.current(0)

    def apply_category() -> None:
        selected_label = category_var.get().strip()
        matched = next((c for c in categories if c["label"] == selected_label), None)
        if not matched:
            matched = next((c for c in categories if selected_label.lower() in c["label"].lower()), None)
        if not matched:
            messagebox.showwarning("Категорії", "Оберіть категорію.")
            return
        confirm_and_apply(
            "Оновлено товарів",
            lambda: db.bulk_update_products_category(db_conn, product_ids, matched["id"]),
        )

    ttk.Button(category_frame, text="Застосувати категорію", command=apply_category).pack(
        side=tk.LEFT, padx=4, pady=4
    )

    # Unit
    unit_frame = ttk.LabelFrame(dlg, text="Одиниця")
    unit_frame.grid(row=4, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Label(unit_frame, text="Одиниця:").pack(side=tk.LEFT, padx=4, pady=4)
    unit_var = tk.StringVar(value=default_unit)
    ttk.Entry(unit_frame, textvariable=unit_var, width=10).pack(side=tk.LEFT, padx=4, pady=4)
    ttk.Button(
        unit_frame,
        text="Застосувати одиницю",
        command=lambda: confirm_and_apply(
            "Оновлено товарів",
            lambda: db.bulk_update_products_unit(db_conn, product_ids, unit_var.get()),
        ),
    ).pack(side=tk.LEFT, padx=4, pady=4)

    # Extra categories
    extras_frame = ttk.LabelFrame(dlg, text="Додаткові категорії")
    extras_frame.grid(row=5, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    extras_frame.columnconfigure(0, weight=1)
    extras_box = tk.Listbox(extras_frame, selectmode=tk.MULTIPLE, height=min(10, max(6, len(categories))), exportselection=False)
    for cat in categories:
        extras_box.insert(tk.END, cat["label"])
    extras_box.grid(row=0, column=0, rowspan=2, padx=4, pady=4, sticky="nsew")
    scroll = ttk.Scrollbar(extras_frame, orient="vertical", command=extras_box.yview)
    extras_box.configure(yscrollcommand=scroll.set)
    scroll.grid(row=0, column=1, rowspan=2, sticky="ns", pady=4)

    def selected_extra_ids() -> list[int]:
        return [categories[i]["id"] for i in extras_box.curselection() if i < len(categories)]

    def apply_extra_add() -> None:
        ids = selected_extra_ids()
        if not ids:
            messagebox.showwarning("Категорії", "Оберіть додаткові категорії.")
            return
        confirm_and_apply(
            "Додано зв'язків",
            lambda: db.bulk_add_product_category_links(db_conn, product_ids, ids),
        )

    def apply_extra_remove() -> None:
        ids = selected_extra_ids()
        if not ids:
            messagebox.showwarning("Категорії", "Оберіть додаткові категорії.")
            return
        confirm_and_apply(
            "Видалено зв'язків",
            lambda: db.bulk_remove_product_category_links(db_conn, product_ids, ids),
        )

    ttk.Button(extras_frame, text="Додати категорії", command=apply_extra_add).grid(
        row=0, column=2, padx=6, pady=4, sticky="n"
    )

    ttk.Button(extras_frame, text="Прибрати категорії", command=apply_extra_remove).grid(
        row=1, column=2, padx=6, pady=4, sticky="n"
    )

    labels_frame = ttk.LabelFrame(dlg, text="Етикетки (Code128)")
    labels_frame.grid(row=6, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Label(labels_frame, text="Шаблон:").pack(side=tk.LEFT, padx=4, pady=4)
    template_display_var = tk.StringVar()
    template_combo = ttk.Combobox(labels_frame, textvariable=template_display_var, state="readonly", width=26)
    template_combo.pack(side=tk.LEFT, padx=4, pady=4)
    template_map: dict[str, int] = {}
    templates_cache: dict[int, dict] = {}

    start_row_var = tk.StringVar(value="1")
    start_col_var = tk.StringVar(value="1")
    start_frame = ttk.Frame(labels_frame)
    start_row_label = ttk.Label(start_frame, text="Ряд:")
    start_row_spin = ttk.Spinbox(start_frame, from_=1, to=1, textvariable=start_row_var, width=4)
    start_col_label = ttk.Label(start_frame, text="Кол:")
    start_col_spin = ttk.Spinbox(start_frame, from_=1, to=1, textvariable=start_col_var, width=4)

    def refresh_template_choices(selected_id: int | None = None) -> None:
        nonlocal templates_cache
        template_map.clear()
        templates_cache = {}
        try:
            templates = db.list_label_templates(active_only=True)
        except Exception:
            logging.exception("Failed to load label templates")
            show_error("Етикетки", "Не вдалося завантажити шаблони")
            return
        display_values: list[str] = []
        default_display: str | None = None
        last_selected_id = None
        try:
            last_selected_id = int(parent.settings.get("print", "last_template_id") or 0)
        except Exception:
            last_selected_id = None
        for row in templates:
            tpl = dict(row)
            display = f"{tpl['title']} [{tpl['code']}]"
            display_values.append(display)
            tid = int(tpl["id"])
            template_map[display] = tid
            templates_cache[tid] = tpl
            if int(tpl.get("is_default") or 0) == 1:
                default_display = display
        template_combo.configure(values=display_values)
        target_id = selected_id or last_selected_id
        if target_id and target_id in template_map.values():
            for disp, tid in template_map.items():
                if tid == target_id:
                    template_display_var.set(disp)
                    break
        elif default_display:
            template_display_var.set(default_display)
        elif display_values:
            template_display_var.set(display_values[0])
        update_start_controls()

    def update_start_controls(*_args) -> None:
        display = template_display_var.get()
        tpl_id = template_map.get(display)
        tpl_row = templates_cache.get(tpl_id or -1)
        is_sheet = tpl_row and (tpl_row.get("kind") == "sheet")
        for widget in [start_row_label, start_row_spin, start_col_label, start_col_spin, start_frame]:
            widget.pack_forget()
        if is_sheet:
            rows = max(int(tpl_row.get("rows", 1)), 1)
            cols = max(int(tpl_row.get("cols", 1)), 1)
            start_row_spin.configure(to=rows)
            start_col_spin.configure(to=cols)
            start_row_label.pack(side=tk.LEFT, padx=2)
            start_row_spin.pack(side=tk.LEFT, padx=2)
            start_col_label.pack(side=tk.LEFT, padx=2)
            start_col_spin.pack(side=tk.LEFT, padx=2)
            start_frame.pack(side=tk.LEFT, padx=4, pady=4)
        else:
            start_row_var.set("1")
            start_col_var.set("1")

    template_combo.bind("<<ComboboxSelected>>", update_start_controls)

    def open_template_manager() -> None:
        dlg = TemplateManagerDialog(parent, settings=getattr(parent, "settings", None))
        parent.wait_window(dlg.root)
        refresh_template_choices()

    ttk.Button(labels_frame, text="Шаблони…", command=open_template_manager).pack(side=tk.LEFT, padx=4, pady=4)

    ttk.Label(labels_frame, text="К-сть етикеток на товар:").pack(side=tk.LEFT, padx=4, pady=4)
    qty_var = tk.StringVar(value="1")
    ttk.Spinbox(labels_frame, from_=1, to=999, textvariable=qty_var, width=5).pack(side=tk.LEFT, padx=4, pady=4)

    include_aliases_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(labels_frame, text="Друкувати також додаткові штрихкоди (аліаси)", variable=include_aliases_var).pack(
        side=tk.LEFT, padx=4, pady=4
    )

    refresh_template_choices()

    def generate_labels() -> None:
        try:
            qty_each = int(qty_var.get())
        except ValueError:
            show_error("Етикетки", "Вкажіть кількість етикеток числом")
            return
        if qty_each <= 0:
            show_error("Етикетки", "Кількість має бути більшою за 0")
            return

        file_path = filedialog.asksaveasfilename(
            title="Файл PDF з етикетками",
            defaultextension=".pdf",
            filetypes=[("PDF", "*.pdf"), ("Усі файли", "*.*")],
            initialdir=str(parent.default_workdir()),
        )
        if not file_path:
            return

        try:
            placeholders = ",".join("?" * len(product_ids))
            rows = db_conn.execute(
                f"SELECT id, sku, name FROM Products WHERE id IN ({placeholders})",
                tuple(product_ids),
            ).fetchall()
            if not rows:
                show_error("Етикетки", "Не знайдено жодного товару")
                return
            alias_map: dict[int, list[str]] = defaultdict(list)
            if include_aliases_var.get():
                for alias_row in db_conn.execute(
                    f"SELECT product_id, code FROM ProductBarcodes WHERE product_id IN ({placeholders})",
                    tuple(product_ids),
                ).fetchall():
                    alias_map[int(alias_row["product_id"])].append(alias_row["code"])

            prefix = _sanitize_barcode_prefix(parent.settings.get("defaults", "product", "barcode_prefix") or "")
            items = [
                {
                    "product_id": row["id"],
                    "sku": row["sku"],
                    "name": row["name"],
                    "aliases": alias_map.get(int(row["id"]), []),
                }
                for row in rows
            ]
            tpl_display = template_display_var.get()
            tpl_id = template_map.get(tpl_display)
            if not tpl_id:
                show_error("Етикетки", "Оберіть шаблон")
                return
            template_full = db.get_label_template_full(tpl_id)
            if not template_full:
                show_error("Етикетки", "Шаблон не знайдено")
                return
            try:
                start_row = max(1, int(start_row_var.get() or 1))
                start_col = max(1, int(start_col_var.get() or 1))
            except ValueError:
                start_row = start_col = 1
            labels.generate_product_labels_pdf_v2(
                Path(file_path),
                items,
                barcode_prefix=prefix,
                qty_each=qty_each,
                template_full=template_full,
                include_aliases=bool(include_aliases_var.get()),
                start_row=start_row,
                start_col=start_col,
            )
            try:
                parent.settings.set(tpl_id, "print", "last_template_id")
                parent.settings.save()
            except Exception:
                logging.exception("Failed to save template selection")
            open_file(Path(file_path))
        except Exception:
            logging.exception("Labels generation error")
            show_error("Етикетки", "Не вдалося згенерувати етикетки")

    ttk.Button(labels_frame, text="Згенерувати PDF...", command=generate_labels).pack(side=tk.LEFT, padx=6, pady=4)

    ttk.Button(dlg, text="Закрити", command=close_dialog).grid(row=7, column=0, columnspan=3, pady=8)
    dlg.protocol("WM_DELETE_WINDOW", close_dialog)


def open_labels_print_dialog(parent, items: list[dict], mode_title: str) -> None:
    if not items:
        messagebox.showwarning("Етикетки", "Немає позицій для друку.")
        return

    dlg = tk.Toplevel(parent)
    dlg.title("Друк етикеток")
    dlg.grab_set()
    dlg.resizable(False, False)

    ttk.Label(dlg, text=mode_title, font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=3, padx=8, pady=(8, 4), sticky="w")

    template_frame = ttk.LabelFrame(dlg, text="Шаблон")
    template_frame.grid(row=1, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Label(template_frame, text="Шаблон:").pack(side=tk.LEFT, padx=4, pady=4)
    template_display_var = tk.StringVar()
    template_combo = ttk.Combobox(template_frame, textvariable=template_display_var, state="readonly", width=26)
    template_combo.pack(side=tk.LEFT, padx=4, pady=4)
    template_map: dict[str, int] = {}
    templates_cache: dict[int, dict] = {}

    start_row_var = tk.StringVar(value="1")
    start_col_var = tk.StringVar(value="1")
    start_frame = ttk.Frame(template_frame)
    start_row_label = ttk.Label(start_frame, text="Ряд:")
    start_row_spin = ttk.Spinbox(start_frame, from_=1, to=1, textvariable=start_row_var, width=4)
    start_col_label = ttk.Label(start_frame, text="Кол:")
    start_col_spin = ttk.Spinbox(start_frame, from_=1, to=1, textvariable=start_col_var, width=4)

    def refresh_template_choices(selected_id: int | None = None) -> None:
        nonlocal templates_cache
        template_map.clear()
        templates_cache = {}
        try:
            templates = db.list_label_templates(active_only=True)
        except Exception:
            logging.exception("Failed to load label templates")
            show_error("Етикетки", "Не вдалося завантажити шаблони")
            return
        display_values: list[str] = []
        default_display: str | None = None
        last_selected_id = None
        try:
            last_selected_id = int(parent.settings.get("print", "last_template_id") or 0)
        except Exception:
            last_selected_id = None
        for row in templates:
            tpl = dict(row)
            display = f"{tpl['title']} [{tpl['code']}]"
            display_values.append(display)
            tid = int(tpl["id"])
            template_map[display] = tid
            templates_cache[tid] = tpl
            if int(tpl.get("is_default") or 0) == 1:
                default_display = display
        template_combo.configure(values=display_values)
        target_id = selected_id or last_selected_id
        if target_id and target_id in template_map.values():
            for disp, tid in template_map.items():
                if tid == target_id:
                    template_display_var.set(disp)
                    break
        elif default_display:
            template_display_var.set(default_display)
        elif display_values:
            template_display_var.set(display_values[0])
        update_start_controls()

    def update_start_controls(*_args) -> None:
        display = template_display_var.get()
        tpl_id = template_map.get(display)
        tpl_row = templates_cache.get(tpl_id or -1)
        is_sheet = tpl_row and (tpl_row.get("kind") == "sheet")
        for widget in [start_row_label, start_row_spin, start_col_label, start_col_spin, start_frame]:
            widget.pack_forget()
        if is_sheet:
            rows = max(int(tpl_row.get("rows", 1)), 1)
            cols = max(int(tpl_row.get("cols", 1)), 1)
            start_row_spin.configure(to=rows)
            start_col_spin.configure(to=cols)
            start_row_label.pack(side=tk.LEFT, padx=2)
            start_row_spin.pack(side=tk.LEFT, padx=2)
            start_col_label.pack(side=tk.LEFT, padx=2)
            start_col_spin.pack(side=tk.LEFT, padx=2)
            start_frame.pack(side=tk.LEFT, padx=4, pady=4)
        else:
            start_row_var.set("1")
            start_col_var.set("1")

    template_combo.bind("<<ComboboxSelected>>", update_start_controls)

    def open_template_manager() -> None:
        dlg_manager = TemplateManagerDialog(parent, settings=getattr(parent, "settings", None))
        parent.wait_window(dlg_manager.root)
        refresh_template_choices()

    ttk.Button(template_frame, text="Шаблони…", command=open_template_manager).pack(side=tk.LEFT, padx=4, pady=4)

    qty_frame = ttk.LabelFrame(dlg, text="Кількість етикеток")
    qty_frame.grid(row=2, column=0, columnspan=3, padx=8, pady=4, sticky="ew")
    ttk.Label(qty_frame, text="Округлення:").pack(side=tk.LEFT, padx=4, pady=4)
    rounding_var = tk.StringVar(value="round")
    rounding_combo = ttk.Combobox(qty_frame, textvariable=rounding_var, state="readonly", values=["round", "floor", "ceil", "skip"], width=8)
    rounding_combo.pack(side=tk.LEFT, padx=4, pady=4)

    ttk.Label(qty_frame, text="Множник:").pack(side=tk.LEFT, padx=4, pady=4)
    multiplier_var = tk.StringVar(value="1")
    ttk.Spinbox(qty_frame, from_=1, to=999, textvariable=multiplier_var, width=5).pack(side=tk.LEFT, padx=4, pady=4)

    include_aliases_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(
        qty_frame,
        text="Друкувати також додаткові штрихкоди (аліаси)",
        variable=include_aliases_var,
    ).pack(side=tk.LEFT, padx=4, pady=4)

    info_var = tk.StringVar(value="")
    info_label = ttk.Label(dlg, textvariable=info_var)
    info_label.grid(row=3, column=0, columnspan=3, padx=8, pady=(0, 4), sticky="w")

    def _apply_rounding(qty_value: float, mode: str) -> int | None:
        if mode == "floor":
            return math.floor(qty_value)
        if mode == "ceil":
            return math.ceil(qty_value)
        if mode == "skip":
            if math.isclose(qty_value, round(qty_value)):
                return int(round(qty_value))
            return None
        return int(round(qty_value))

    def compute_summary() -> tuple[int, int, int]:
        try:
            multiplier = max(1, min(999, int(multiplier_var.get())))
        except ValueError:
            multiplier = 1
        rounding_mode = rounding_var.get()
        positions = 0
        labels_count = 0
        skipped = 0
        for item in items:
            qty_value = float(item.get("qty", 0) or 0)
            qty_int = _apply_rounding(qty_value, rounding_mode)
            if qty_int is None:
                skipped += 1
                continue
            if qty_int <= 0:
                continue
            positions += 1
            labels_count += qty_int * multiplier
        return positions, labels_count, skipped

    def update_summary(*_args) -> None:
        positions, labels_count, skipped = compute_summary()
        text = f"Позицій: {positions} | Етикеток (SKU): {labels_count}"
        if include_aliases_var.get():
            text += " + аліаси"
        if rounding_var.get() == "skip" and skipped:
            text += f" | Пропущено дробових: {skipped}"
        info_var.set(text)

    rounding_combo.bind("<<ComboboxSelected>>", update_summary)
    multiplier_var.trace_add("write", update_summary)
    include_aliases_var.trace_add("write", update_summary)

    refresh_template_choices()
    update_summary()

    def generate_labels() -> None:
        try:
            multiplier = int(multiplier_var.get())
        except ValueError:
            show_error("Етикетки", "Множник має бути числом")
            return
        if multiplier < 1:
            show_error("Етикетки", "Множник має бути не менше 1")
            return
        rounding_mode = rounding_var.get()

        tpl_display = template_display_var.get()
        tpl_id = template_map.get(tpl_display)
        if not tpl_id:
            show_error("Етикетки", "Оберіть шаблон")
            return
        template_full = db.get_label_template_full(tpl_id)
        if not template_full:
            show_error("Етикетки", "Шаблон не знайдено")
            return

        positions, labels_count, skipped = compute_summary()
        if labels_count <= 0:
            messagebox.showwarning("Етикетки", "Немає позицій з кількістю > 0.")
            return
        if rounding_mode == "skip" and skipped:
            messagebox.showinfo("Етикетки", f"Пропущено позицій через дробову кількість: {skipped}")

        file_path = filedialog.asksaveasfilename(
            title="Файл PDF з етикетками",
            defaultextension=".pdf",
            filetypes=[("PDF", "*.pdf"), ("Усі файли", "*.*")],
            initialdir=str(parent.default_workdir()),
        )
        if not file_path:
            return

        try:
            start_row = max(1, int(start_row_var.get() or 1))
            start_col = max(1, int(start_col_var.get() or 1))
        except ValueError:
            start_row = start_col = 1
        include_aliases = bool(include_aliases_var.get())
        product_ids = []
        for item in items:
            try:
                pid = int(item.get("product_id"))
            except (TypeError, ValueError):
                continue
            if pid not in product_ids:
                product_ids.append(pid)

        alias_map: dict[int, list[str]] = defaultdict(list)
        if include_aliases and product_ids:
            try:
                placeholders = ",".join("?" * len(product_ids))
                with db.get_connection() as conn:
                    for alias_row in conn.execute(
                        f"SELECT product_id, code FROM ProductBarcodes WHERE product_id IN ({placeholders})",
                        tuple(product_ids),
                    ).fetchall():
                        alias_map[int(alias_row["product_id"])].append(alias_row["code"])
            except Exception:
                logging.exception("Failed to load aliases for labels")

        def build_expanded_items() -> list[dict]:
            expanded: list[dict] = []
            for item in items:
                try:
                    pid = int(item.get("product_id"))
                except (TypeError, ValueError):
                    continue
                qty_value = float(item.get("qty", 0) or 0)
                qty_int = _apply_rounding(qty_value, rounding_mode)
                if qty_int is None or qty_int <= 0:
                    continue
                total_qty = qty_int * multiplier
                if total_qty <= 0:
                    continue
                entry = {
                    "product_id": pid,
                    "sku": item.get("sku") or "",
                    "name": item.get("name") or "",
                    "aliases": alias_map.get(pid, []),
                }
                expanded.extend([entry] * total_qty)
            return expanded

        expanded_items = build_expanded_items()
        if not expanded_items:
            messagebox.showwarning("Етикетки", "Немає позицій для друку після округлення.")
            return

        try:
            prefix = _sanitize_barcode_prefix(parent.settings.get("defaults", "product", "barcode_prefix") or "")
        except Exception:
            prefix = ""

        try:
            labels.generate_product_labels_pdf_v2(
                Path(file_path),
                expanded_items,
                barcode_prefix=prefix,
                qty_each=1,
                template_full=template_full,
                include_aliases=include_aliases,
                start_row=start_row,
                start_col=start_col,
            )
            try:
                parent.settings.set(tpl_id, "print", "last_template_id")
                parent.settings.save()
            except Exception:
                logging.exception("Failed to save template selection")
            open_file(Path(file_path))
        except Exception:
            logging.exception("Labels generation error (doc/stock)")
            show_error("Етикетки", "Не вдалося згенерувати етикетки")

    ttk.Button(dlg, text="Згенерувати PDF...", command=generate_labels).grid(row=4, column=0, padx=8, pady=8, sticky="w")
    ttk.Button(dlg, text="Закрити", command=dlg.destroy).grid(row=4, column=2, padx=8, pady=8, sticky="e")
    dlg.columnconfigure(1, weight=1)
    dlg.protocol("WM_DELETE_WINDOW", dlg.destroy)

def category_prompt(title: str, initial=None):
    dlg = tk.Toplevel()
    dlg.title(title)
    dlg.grab_set()

    ttk.Label(dlg, text="Назва").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    name_var = tk.StringVar(value=initial[0] if initial else "")
    ttk.Entry(dlg, textvariable=name_var, width=30).grid(row=0, column=1, padx=6, pady=4)

    ttk.Label(dlg, text="Колір (hex)").grid(row=1, column=0, padx=6, pady=4, sticky="w")
    color_var = tk.StringVar(value=initial[1] if initial else "")
    ttk.Entry(dlg, textvariable=color_var, width=20).grid(row=1, column=1, padx=6, pady=4, sticky="w")

    ttk.Label(dlg, text="Іконка/emoji").grid(row=2, column=0, padx=6, pady=4, sticky="w")
    icon_var = tk.StringVar(value=initial[2] if initial else "")
    ttk.Entry(dlg, textvariable=icon_var, width=20).grid(row=2, column=1, padx=6, pady=4, sticky="w")

    ttk.Label(dlg, text="Типові атрибути (через кому)").grid(row=3, column=0, padx=6, pady=4, sticky="w")
    attrs_var = tk.StringVar(value=initial[3] if initial else "")
    ttk.Entry(dlg, textvariable=attrs_var, width=40).grid(row=3, column=1, padx=6, pady=4, sticky="w")

    service_var = tk.BooleanVar(value=initial[4] if initial else False)
    hidden_var = tk.BooleanVar(value=initial[5] if initial else False)
    ttk.Checkbutton(dlg, text="Службова", variable=service_var).grid(row=4, column=1, padx=6, pady=4, sticky="w")
    ttk.Checkbutton(dlg, text="Прихована", variable=hidden_var).grid(row=5, column=1, padx=6, pady=4, sticky="w")

    result = None

    def on_ok():
        nonlocal result
        name = name_var.get().strip()
        if not name:
            messagebox.showerror("Валідація", "Назва обов'язкова")
            return
        result = (name, color_var.get().strip(), icon_var.get().strip(), attrs_var.get().strip(), service_var.get(), hidden_var.get())
        dlg.destroy()

    ttk.Button(dlg, text="OK", command=on_ok).grid(row=6, column=0, padx=6, pady=8)
    ttk.Button(dlg, text="Скасувати", command=dlg.destroy).grid(row=6, column=1, padx=6, pady=8)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: dlg.destroy())
    dlg.wait_window()
    return result


def select_category_dialog(title: str, options: list[tuple[Optional[int], str]]):
    dlg = tk.Toplevel()
    dlg.title(title)
    dlg.grab_set()

    ttk.Label(dlg, text="Категорія").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    values = [opt[1] for opt in options]
    combo_var = tk.StringVar()
    combo = ttk.Combobox(dlg, textvariable=combo_var, state="readonly", values=values)
    combo.grid(row=0, column=1, padx=6, pady=4)
    combo.current(0 if values else -1)

    result: Optional[int] = None

    def on_ok():
        nonlocal result
        if not values:
            result = None
        else:
            idx = combo.current()
            result = options[idx][0]
        dlg.destroy()

    ttk.Button(dlg, text="OK", command=on_ok).grid(row=1, column=0, padx=6, pady=8)
    ttk.Button(dlg, text="Скасувати", command=dlg.destroy).grid(row=1, column=1, padx=6, pady=8)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: dlg.destroy())
    dlg.wait_window()
    return result


def counterparty_prompt(initial=None, default_type: str | None = None):
    dlg = tk.Toplevel()
    dlg.title("Контрагент")
    dlg.grab_set()

    labels = ["Назва", "Телефон", "Email", "Адреса", "Нотатка"]
    vars_ = [tk.StringVar(value=initial[i] if initial else "") for i in [0, 2, 3, 4, 5]]

    ttk.Label(dlg, text="Назва").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    ttk.Entry(dlg, textvariable=vars_[0], width=30).grid(row=0, column=1, padx=6, pady=4)

    ttk.Label(dlg, text="Тип").grid(row=1, column=0, padx=6, pady=4, sticky="w")
    type_var = tk.StringVar()
    types = ["Постачальник", "Покупець", "Постачальник/Покупець", "Інший"]
    type_values = {"Постачальник": "supplier", "Покупець": "customer", "Постачальник/Покупець": "both", "Інший": "other"}
    type_combo = ttk.Combobox(dlg, textvariable=type_var, values=types, state="readonly")
    type_combo.grid(row=1, column=1, padx=6, pady=4)
    if initial:
        inv_map = {v: k for k, v in type_values.items()}
        type_combo.set(inv_map.get(initial[1], types[0]))
    elif default_type:
        inv_map = {v: k for k, v in type_values.items()}
        type_combo.set(inv_map.get(default_type, types[0]))
    else:
        type_combo.current(0)

    for i, label in enumerate(labels[1:], start=2):
        ttk.Label(dlg, text=label).grid(row=i, column=0, padx=6, pady=4, sticky="w")
        ttk.Entry(dlg, textvariable=vars_[i - 1], width=30).grid(row=i, column=1, padx=6, pady=4)

    result = None

    def on_ok():
        nonlocal result
        name = vars_[0].get().strip()
        if not name:
            messagebox.showerror("Валідація", "Заповніть назву")
            return
        ctype = type_values.get(type_var.get(), "supplier")
        phone, email, address, note = [v.get().strip() for v in vars_[1:]]
        result = (name, ctype, phone, email, address, note)
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    btns = ttk.Frame(dlg)
    btns.grid(row=6, column=0, columnspan=2, pady=8)
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: on_cancel())
    if editable:
        scan_entry.focus_set()
    dlg.wait_window()
    return result


def warehouse_prompt(initial=None):
    dlg = tk.Toplevel()
    dlg.title("Склад")
    dlg.grab_set()

    name_var = tk.StringVar(value=initial[0] if initial else "")
    desc_var = tk.StringVar(value=initial[1] if initial else "")
    active_var = tk.BooleanVar(value=initial[2] if initial else True)

    ttk.Label(dlg, text="Назва").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    ttk.Entry(dlg, textvariable=name_var, width=30).grid(row=0, column=1, padx=6, pady=4)
    ttk.Label(dlg, text="Опис").grid(row=1, column=0, padx=6, pady=4, sticky="w")
    ttk.Entry(dlg, textvariable=desc_var, width=40).grid(row=1, column=1, padx=6, pady=4)
    ttk.Checkbutton(dlg, text="Активний", variable=active_var).grid(row=2, column=1, padx=6, pady=4, sticky="w")

    result = None

    def on_ok():
        nonlocal result
        name = name_var.get().strip()
        if not name:
            messagebox.showerror("Валідація", "Заповніть назву")
            return
        result = (name, desc_var.get().strip(), bool(active_var.get()))
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    btns = ttk.Frame(dlg)
    btns.grid(row=3, column=0, columnspan=2, pady=8)
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: on_cancel())
    if editable:
        scan_entry.focus_set()
        scan_entry.icursor("end")
        scan_entry.selection_range(0, tk.END)
    dlg.wait_window()
    return result


def channel_prompt(initial=None):
    dlg = tk.Toplevel()
    dlg.title("Канал продажу")
    dlg.grab_set()
    name_var = tk.StringVar(value=initial[0] if initial else "")
    active_var = tk.BooleanVar(value=initial[1] if initial else True)

    ttk.Label(dlg, text="Назва").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    ttk.Entry(dlg, textvariable=name_var, width=30).grid(row=0, column=1, padx=6, pady=4)
    ttk.Checkbutton(dlg, text="Активний", variable=active_var).grid(row=1, column=1, padx=6, pady=4, sticky="w")

    result = None

    def on_ok():
        nonlocal result
        name = name_var.get().strip()
        if not name:
            messagebox.showerror("Валідація", "Заповніть назву")
            return
        result = (name, bool(active_var.get()))
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    btns = ttk.Frame(dlg)
    btns.grid(row=2, column=0, columnspan=2, pady=8)
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: on_cancel())
    if editable:
        dlg.after_idle(
            lambda: (
                scan_entry.focus_set(),
                scan_entry.selection_range(0, tk.END),
                scan_entry.icursor(tk.END),
            )
        )
    dlg.wait_window()
    return result


def extra_cost_prompt(counterparties, currencies, purchases, doc=None, lines=None, selected_purchase_ids=None, allow_edit: bool = True):
    dlg = tk.Toplevel()
    dlg.title("Супутні витрати")
    dlg.grab_set()
    if doc and isinstance(doc, sqlite3.Row):
        doc = dict(doc)
    selected_purchase_ids = list(selected_purchase_ids or [])

    frame = ttk.Frame(dlg, padding=10)
    frame.grid(row=0, column=0, sticky="nsew")
    dlg.columnconfigure(0, weight=1)
    dlg.rowconfigure(0, weight=1)

    row_idx = 0
    ttk.Label(frame, text="Дата (YYYY-MM-DD)").grid(row=row_idx, column=0, sticky="e", padx=4, pady=2)
    date_var = tk.StringVar(value=doc["doc_date"] if doc else datetime.now().strftime("%Y-%m-%d"))
    ttk.Entry(frame, textvariable=date_var, width=14, state="normal" if allow_edit else "disabled").grid(row=row_idx, column=1, sticky="w")

    row_idx += 1
    ttk.Label(frame, text="Валюта").grid(row=row_idx, column=0, sticky="e", padx=4, pady=2)
    curr_codes = [c["code"] for c in currencies]
    curr_var = tk.StringVar(value=doc["currency_code"] if doc else (curr_codes[0] if curr_codes else get_base_currency_code()))
    curr_combo = ttk.Combobox(frame, textvariable=curr_var, values=curr_codes, state="readonly")
    if not allow_edit:
        curr_combo.state(["disabled"])
    curr_combo.grid(row=row_idx, column=1, sticky="w")

    last_currency = curr_var.get()

    row_idx += 1
    ttk.Label(frame, text="Курс").grid(row=row_idx, column=0, sticky="e", padx=4, pady=2)
    try:
        default_rate = doc["exchange_rate"] if doc else ensure_rate_for_date(curr_var.get(), date_var.get())
    except Exception:
        default_rate = doc["exchange_rate"] if doc else 1.0
    rate_var = tk.StringVar(value=f"{default_rate:.4f}")
    ttk.Entry(frame, textvariable=rate_var, width=12, state="normal" if allow_edit else "disabled").grid(row=row_idx, column=1, sticky="w")

    def on_currency_change(event=None):
        nonlocal last_currency
        if not allow_edit:
            return
        try:
            rate_val = ensure_rate_for_date(curr_var.get(), date_var.get())
        except ValueError as exc:
            messagebox.showerror("Курс", str(exc))
            curr_var.set(last_currency)
            return
        rate_var.set(f"{rate_val:.4f}")
        last_currency = curr_var.get()

    curr_combo.bind("<<ComboboxSelected>>", on_currency_change)

    row_idx += 1
    ttk.Label(frame, text="Контрагент").grid(row=row_idx, column=0, sticky="e", padx=4, pady=2)
    filtered_counterparties = [c for c in counterparties if c["type"] in {"supplier", "both", "other"}]
    cp_names = ["-"] + [c["name"] for c in filtered_counterparties]
    cp_var = tk.StringVar()
    cp_combo = ttk.Combobox(frame, textvariable=cp_var, values=cp_names, state="readonly")
    if not allow_edit:
        cp_combo.state(["disabled"])
    cp_combo.grid(row=row_idx, column=1, sticky="w")

    row_idx += 1
    ttk.Label(frame, text="Коментар").grid(row=row_idx, column=0, sticky="e", padx=4, pady=2)
    comment_var = tk.StringVar(value=doc["comment"] if doc else "")
    ttk.Entry(frame, textvariable=comment_var, width=40, state="normal" if allow_edit else "disabled").grid(row=row_idx, column=1, sticky="ew")

    if doc and doc.get("partner_id"):
        try:
            cp_combo.current(next(i for i, c in enumerate(filtered_counterparties, start=1) if c["id"] == doc.get("partner_id")))
        except StopIteration:
            cp_combo.current(0)
    else:
        cp_combo.current(0)

    row_idx += 1
    ttk.Label(frame, text="Рядки витрат").grid(row=row_idx, column=0, sticky="ne", padx=4, pady=4)
    line_frame = ttk.Frame(frame)
    line_frame.grid(row=row_idx, column=1, sticky="nsew")
    frame.rowconfigure(row_idx, weight=1)
    line_frame.columnconfigure(0, weight=1)

    line_columns = ("type", "amount_doc")
    line_tree = ttk.Treeview(line_frame, columns=line_columns, show="headings", height=6)
    line_tree.heading("type", text="Стаття")
    line_tree.heading("amount_doc", text="Сума (вал)")
    line_tree.column("type", width=180)
    line_tree.column("amount_doc", width=100)
    line_tree.grid(row=0, column=0, sticky="nsew")
    line_scroll = ttk.Scrollbar(line_frame, orient="vertical", command=line_tree.yview)
    line_tree.configure(yscrollcommand=line_scroll.set)
    line_scroll.grid(row=0, column=1, sticky="ns")
    line_frame.rowconfigure(0, weight=1)

    lines_data = []
    if lines:
        lines_data = [dict(cost_type=ln["cost_type"], amount_doc=float(ln["amount_doc"])) for ln in lines]

    def refresh_lines():
        line_tree.delete(*line_tree.get_children())
        for idx, ln in enumerate(lines_data):
            line_tree.insert("", "end", iid=str(idx), values=(ln["cost_type"], f"{ln['amount_doc']:.2f}"))

    if allow_edit:
        entry_row = ttk.Frame(line_frame)
        entry_row.grid(row=1, column=0, columnspan=2, sticky="ew", pady=4)
        ttk.Label(entry_row, text="Стаття").grid(row=0, column=0, padx=2)
        cost_var = tk.StringVar(value="Доставка")
        ttk.Entry(entry_row, textvariable=cost_var, width=20).grid(row=0, column=1, padx=2)
        ttk.Label(entry_row, text="Сума").grid(row=0, column=2, padx=2)
        amount_var = tk.StringVar(value="0")
        ttk.Entry(entry_row, textvariable=amount_var, width=12).grid(row=0, column=3, padx=2)

        def add_line():
            try:
                amount_val = float(amount_var.get())
            except ValueError:
                show_error("Валідація", "Невірна сума")
                return
            if amount_val == 0:
                show_error("Валідація", "Сума повинна бути більшою за 0")
                return
            lines_data.append({"cost_type": cost_var.get().strip() or "Інше", "amount_doc": amount_val})
            refresh_lines()

        def delete_line():
            sel = line_tree.selection()
            if not sel:
                return
            idx = int(sel[0])
            if 0 <= idx < len(lines_data):
                lines_data.pop(idx)
                refresh_lines()

        ttk.Button(entry_row, text="Додати", command=add_line).grid(row=0, column=4, padx=4)
        ttk.Button(entry_row, text="Видалити", command=delete_line).grid(row=0, column=5, padx=4)

    row_idx += 1
    ttk.Label(frame, text="Закупівлі").grid(row=row_idx, column=0, sticky="ne", padx=4, pady=4)
    purchase_frame = ttk.Frame(frame)
    purchase_frame.grid(row=row_idx, column=1, sticky="nsew")
    frame.rowconfigure(row_idx, weight=1)
    purchase_frame.columnconfigure(1, weight=1)

    available_tree = ttk.Treeview(purchase_frame, columns=("date", "supplier", "total"), show="headings", height=6)
    for col, title, width in [("date", "Дата", 90), ("supplier", "Постачальник", 200), ("total", "Сума (база)", 120)]:
        available_tree.heading(col, text=title)
        available_tree.column(col, width=width)
    available_tree.grid(row=0, column=0, sticky="nsew")
    avail_scroll = ttk.Scrollbar(purchase_frame, orient="vertical", command=available_tree.yview)
    available_tree.configure(yscrollcommand=avail_scroll.set)
    avail_scroll.grid(row=0, column=1, sticky="ns")

    selected_list = tk.Listbox(purchase_frame, height=6)
    selected_list.grid(row=0, column=2, padx=6, sticky="nsew")
    purchase_frame.columnconfigure(2, weight=1)

    purchase_lookup = {p["id"]: p for p in purchases}
    for p in purchases:
        available_tree.insert("", "end", iid=str(p["id"]), values=(p["doc_date"], p["supplier"] or "-", f"{p['total']:.2f}"))

    def refresh_selected():
        selected_list.delete(0, tk.END)
        for pid in selected_purchase_ids:
            p = purchase_lookup.get(pid)
            if p:
                selected_list.insert(tk.END, f"{p['id']} | {p['doc_date']} | {p['supplier'] or '-'} | {p['total']:.2f}")

    def add_purchase():
        if not allow_edit:
            return
        sel = available_tree.selection()
        if not sel:
            return
        pid = int(sel[0])
        if pid not in selected_purchase_ids:
            selected_purchase_ids.append(pid)
            refresh_selected()

    def remove_purchase():
        if not allow_edit:
            return
        idx = selected_list.curselection()
        if not idx:
            return
        pid = selected_purchase_ids[idx[0]]
        selected_purchase_ids.remove(pid)
        refresh_selected()

    purchase_btns = ttk.Frame(purchase_frame)
    purchase_btns.grid(row=1, column=0, columnspan=3, pady=4)
    ttk.Button(purchase_btns, text="Додати закупівлю", command=add_purchase).pack(side=tk.LEFT, padx=4)
    ttk.Button(purchase_btns, text="Видалити", command=remove_purchase).pack(side=tk.LEFT, padx=4)

    if selected_purchase_ids:
        refresh_selected()

    row_idx += 1
    post_var = tk.IntVar(value=0)
    if allow_edit:
        ttk.Checkbutton(frame, text="Провести після збереження", variable=post_var).grid(row=row_idx, column=1, sticky="w", pady=6)

    row_idx += 1
    btns = ttk.Frame(frame)
    btns.grid(row=row_idx, column=0, columnspan=2, pady=6)

    result: list = []

    def on_ok():
        try:
            rate_val = float(rate_var.get())
        except ValueError:
            show_error("Валідація", "Невірний курс")
            return
        if rate_val <= 0:
            show_error("Валідація", "Курс має бути більшим за 0")
            return
        partner_name = cp_var.get()
        partner_id = None
        if partner_name and partner_name != "-":
            found = next((c for c in filtered_counterparties if c["name"] == partner_name), None)
            partner_id = found["id"] if found else None
        if curr_var.get().strip().upper() != get_base_currency_code() and not db.rate_on_date(curr_var.get(), date_var.get()):
            db.add_currency_rate(curr_var.get(), date_var.get(), rate_val)
        result.append(
            (
                {
                    "doc_date": date_var.get(),
                    "currency": curr_var.get(),
                    "rate": rate_val,
                    "partner_id": partner_id,
                    "comment": comment_var.get(),
                },
                [(ln["cost_type"], ln["amount_doc"]) for ln in lines_data],
                list(selected_purchase_ids),
                bool(post_var.get()),
            )
        )
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    ttk.Button(btns, text="OK", command=on_ok, state="normal" if allow_edit else "disabled").pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)

    refresh_lines()
    dlg.wait_window()
    return result[0] if result else None


def inventory_prompt(warehouses, products, settings: Settings, doc=None, lines=None):
    dlg = tk.Toplevel()
    dlg.title("Інвентаризація")
    dlg.grab_set()
    editable = not doc or doc["status"] == "draft"

    if doc and isinstance(doc, sqlite3.Row):
        doc = dict(doc)

    dlg.columnconfigure(0, weight=1)
    dlg.rowconfigure(0, weight=1)
    content = ttk.Frame(dlg, padding=10)
    content.grid(row=0, column=0, sticky="nsew")
    content.columnconfigure(1, weight=1)

    row_idx = 0
    ttk.Label(content, text="Дата (YYYY-MM-DD)").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    if editable:
        date_picker = DatePicker(
            content,
            initial=datetime.strptime(doc["doc_date"], "%Y-%m-%d").date() if doc else date.today(),
        )
        date_picker.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
        date_entry = None
    else:
        date_var = tk.StringVar(value=doc["doc_date"] if doc else datetime.now().strftime("%Y-%m-%d"))
        date_entry = ttk.Entry(content, textvariable=date_var, width=15, state="disabled")
        date_entry.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
        date_picker = None

    row_idx += 1
    ttk.Label(content, text="Склад").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    wh_var = tk.StringVar()
    wh_names = [w["name"] for w in warehouses]
    wh_combo = ttk.Combobox(content, textvariable=wh_var, values=wh_names, state="readonly")
    if not editable:
        wh_combo.state(["disabled"])
    wh_combo.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")

    row_idx += 1
    ttk.Label(content, text="Коментар").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    comment_var = tk.StringVar(value=doc["comment"] if doc else "")
    comment_entry = ttk.Entry(content, textvariable=comment_var, width=40, state="normal" if editable else "disabled")
    comment_entry.grid(row=row_idx, column=1, padx=6, pady=4, sticky="ew")

    categories = db.list_categories(include_hidden=False)
    category_children: dict[int | None, list[dict]] = defaultdict(list)
    for cat in categories:
        category_children[cat["parent_id"]].append(cat)
    for items in category_children.values():
        items.sort(key=lambda item: item["name"].lower())

    category_display_names: list[str] = []
    cat_display_to_id: dict[str, int] = {}

    def _build_category_tree(parent_id: int | None, level: int = 0) -> None:
        for cat in category_children.get(parent_id, []):
            display = f"{'  ' * level}{cat['name']}"
            category_display_names.append(display)
            cat_display_to_id[display] = cat["id"]
            _build_category_tree(cat["id"], level + 1)

    _build_category_tree(None)
    category_values = ["Усі", *category_display_names]

    row_idx += 1
    ttk.Label(content, text="Категорія").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    cat_var = tk.StringVar(value="Усі")
    include_children_var = tk.BooleanVar(value=True)
    cat_frame = ttk.Frame(content)
    cat_frame.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
    cat_combo = ttk.Combobox(cat_frame, textvariable=cat_var, values=category_values, state="readonly")
    cat_combo.pack(side=tk.LEFT)
    include_children_check = ttk.Checkbutton(
        cat_frame,
        text="Включати підкатегорії",
        variable=include_children_var,
    )
    include_children_check.pack(side=tk.LEFT, padx=6)

    post_now_var = tk.BooleanVar(value=False)
    if editable:
        row_idx += 1
        ttk.Checkbutton(content, text="Провести одразу", variable=post_now_var).grid(
            row=row_idx, column=1, padx=6, pady=2, sticky="w"
        )

    row_idx += 1
    scan_frame = ttk.LabelFrame(content, text="Сканування")
    scan_frame.grid(row=row_idx, column=0, columnspan=2, padx=6, pady=6, sticky="ew")
    scan_frame.columnconfigure(1, weight=1)

    ttk.Label(scan_frame, text="Скан-код").grid(row=0, column=0, padx=6, pady=4, sticky="e")
    scan_var = tk.StringVar()
    scan_entry = ttk.Entry(scan_frame, textvariable=scan_var, width=30, state="normal" if editable else "disabled")
    scan_entry.grid(row=0, column=1, padx=6, pady=4, sticky="w")
    ttk.Label(scan_frame, text="К-сть при скані").grid(row=0, column=2, padx=6, pady=4, sticky="e")
    scan_qty_var = tk.IntVar(value=1)
    scan_qty = ttk.Spinbox(
        scan_frame,
        from_=1,
        to=999,
        textvariable=scan_qty_var,
        width=6,
        state="normal" if editable else "disabled",
    )
    scan_qty.grid(row=0, column=3, padx=6, pady=4, sticky="w")
    scan_status = ttk.Label(scan_frame, text="")
    scan_status.grid(row=1, column=0, columnspan=4, padx=6, pady=(0, 4), sticky="w")

    row_idx += 1
    ttk.Label(content, text="Рядки").grid(row=row_idx, column=0, padx=6, pady=4, sticky="ne")
    line_frame = ttk.Frame(content)
    line_frame.grid(row=row_idx, column=1, padx=6, pady=4, sticky="nsew")
    line_frame.grid_columnconfigure(0, weight=1)
    content.rowconfigure(row_idx, weight=1)

    columns = ["sku", "name", "expected", "counted", "diff", "cost_override", "note"]
    tree = ttk.Treeview(line_frame, columns=columns, show="headings", height=10, selectmode="browse")
    headings = {
        "sku": ("SKU", 100),
        "name": ("Назва", 220),
        "expected": ("Очікувано", 90),
        "counted": ("Факт", 90),
        "diff": ("Різниця", 90),
        "cost_override": ("Собівартість надлишку", 150),
        "note": ("Примітка", 160),
    }
    for col, (title, width) in headings.items():
        tree.heading(col, text=title)
        tree.column(col, width=width, anchor="w")
    tree.grid(row=0, column=0, sticky="nsew")
    yscroll = ttk.Scrollbar(line_frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=yscroll.set)
    yscroll.grid(row=0, column=1, sticky="ns")
    tree.tag_configure("missing_cost", background="#ffe3e3")

    product_lookup = {f"{p['sku']} — {p['name']}": p for p in products}
    products_by_id = {p["id"]: p for p in products}
    product_names = list(product_lookup.keys())
    allowed_product_ids: set[int] = set()

    line_data: list[dict] = []
    if lines:
        for ln in lines:
            line_data.append(
                {
                    "product_id": ln["product_id"],
                    "sku": ln["sku"],
                    "name": ln["name"],
                    "expected_qty": float(ln["expected_qty"] or 0.0),
                    "counted_qty": float(ln["counted_qty"] or 0.0),
                    "cost_override": float(ln["cost_override"]) if ln["cost_override"] is not None else None,
                    "note": ln["note"] or "",
                }
            )

    if doc:
        if doc["warehouse_id"]:
            try:
                wh_combo.current(next(i for i, w in enumerate(warehouses) if w["id"] == doc["warehouse_id"]))
            except StopIteration:
                wh_combo.set(warehouses[0]["name"] if warehouses else "")
    elif warehouses:
        wh_combo.current(0)

    selected_idx: list[int] = []
    balance_cache: dict[int, tuple[float, float]] = {}

    def _current_date() -> str:
        if date_picker:
            return date_picker.get()
        if date_entry:
            return date_entry.get().strip()
        return datetime.now().strftime("%Y-%m-%d")

    def _current_warehouse_id() -> Optional[int]:
        name = wh_var.get()
        if not name:
            return None
        match = next((w for w in warehouses if w["name"] == name), None)
        return match["id"] if match else None

    def _rebuild_allowed_products() -> None:
        allowed_product_ids.clear()
        display = (cat_var.get() or "Усі").strip()
        if display == "Усі":
            return
        cat_id = cat_display_to_id.get(display)
        if not cat_id:
            return
        rows = db.list_products(
            category_id=cat_id,
            include_subcategories=include_children_var.get(),
        )
        allowed_product_ids.update(int(row["id"]) for row in rows)

    def _is_allowed_product(product_id: int) -> bool:
        return (not allowed_product_ids) or (product_id in allowed_product_ids)

    def _get_balance_cached(product_id: int) -> tuple[float, float]:
        if product_id not in balance_cache:
            wh_id = _current_warehouse_id()
            if not wh_id:
                balance_cache[product_id] = (0.0, 0.0)
            else:
                balance_cache[product_id] = db.get_stock_balance(product_id, wh_id)
        return balance_cache[product_id]

    def _needs_cost(line: dict) -> bool:
        eps = 1e-9
        if line["counted_qty"] <= line["expected_qty"] + eps:
            return False
        if line.get("cost_override") and line["cost_override"] > eps:
            return False
        current_qty, avg_cost = _get_balance_cached(line["product_id"])
        if current_qty > eps and avg_cost > eps:
            return False
        wh_id = _current_warehouse_id()
        if not wh_id:
            return True
        last_price = db.get_last_purchase_price(line["product_id"], wh_id, _current_date())
        return last_price <= eps

    summary_var = tk.StringVar()
    summary_label = ttk.Label(line_frame, textvariable=summary_var)
    summary_label.grid(row=1, column=0, columnspan=2, sticky="w", pady=(4, 0))

    def refresh_lines() -> None:
        tree.delete(*tree.get_children())
        diff_total_negative = 0.0
        diff_total_positive = 0.0
        diff_count = 0
        eps = 1e-9
        for idx, ln in enumerate(line_data):
            diff = ln["counted_qty"] - ln["expected_qty"]
            if diff < -eps:
                diff_total_negative += -diff
                diff_count += 1
            elif diff > eps:
                diff_total_positive += diff
                diff_count += 1
            tags = ("missing_cost",) if _needs_cost(ln) else ()
            tree.insert(
                "",
                "end",
                iid=str(idx),
                values=(
                    ln["sku"],
                    ln["name"],
                    f"{ln['expected_qty']:.2f}",
                    f"{ln['counted_qty']:.2f}",
                    f"{diff:.2f}",
                    f"{ln['cost_override']:.2f}" if ln.get("cost_override") is not None else "",
                    ln.get("note", ""),
                ),
                tags=tags,
            )
        summary_var.set(
            "Рядків: {lines} | Розбіжності: {diffs} | Недостача (шт): {missing:.2f} | Надлишок (шт): {extra:.2f}".format(
                lines=len(line_data),
                diffs=diff_count,
                missing=diff_total_negative,
                extra=diff_total_positive,
            )
        )

    def _select_line(event=None) -> None:
        selected_idx.clear()
        sel = tree.selection()
        if not sel:
            return
        selected_idx.append(int(sel[0]))

    tree.bind("<<TreeviewSelect>>", _select_line)

    def _pick_product() -> Optional[dict]:
        def _filtered_product_names() -> list[str]:
            if not allowed_product_ids:
                return product_names
            return [
                name
                for name, product in product_lookup.items()
                if _is_allowed_product(int(product["id"]))
            ]

        filtered_names = _filtered_product_names()
        if not filtered_names:
            messagebox.showinfo("Товари", "Список товарів порожній.")
            return None
        picker = tk.Toplevel(dlg)
        picker.title("Оберіть товар")
        picker.grab_set()
        picker.columnconfigure(0, weight=1)
        picker.rowconfigure(1, weight=1)
        search_var = tk.StringVar()
        ttk.Entry(picker, textvariable=search_var).grid(row=0, column=0, padx=8, pady=6, sticky="ew")
        listbox = tk.Listbox(picker, height=12)
        listbox.grid(row=1, column=0, padx=8, pady=6, sticky="nsew")
        for name in filtered_names:
            listbox.insert(tk.END, name)

        result: dict[str, object] = {}

        def _filter(*_args) -> None:
            text = search_var.get().lower().strip()
            allowed_names = _filtered_product_names()
            listbox.delete(0, tk.END)
            for name in allowed_names:
                if text in name.lower():
                    listbox.insert(tk.END, name)

        def _confirm(_event=None) -> None:
            selection = listbox.curselection()
            if not selection:
                return
            chosen = listbox.get(selection[0])
            result["product"] = product_lookup.get(chosen)
            picker.destroy()

        search_var.trace_add("write", _filter)
        listbox.bind("<Double-1>", _confirm)
        ttk.Button(picker, text="OK", command=_confirm).grid(row=2, column=0, padx=8, pady=(0, 8))
        picker.wait_window()
        return result.get("product")

    def _add_line(product: dict, counted_delta: float = 0.0) -> None:
        for ln in line_data:
            if ln["product_id"] == product["id"]:
                ln["counted_qty"] += counted_delta
                refresh_lines()
                return
        wh_id = _current_warehouse_id()
        expected_qty = db.get_stock_quantity(product["id"], wh_id) if wh_id else 0.0
        line_data.append(
            {
                "product_id": product["id"],
                "sku": product["sku"],
                "name": product["name"],
                "expected_qty": expected_qty,
                "counted_qty": max(0.0, counted_delta),
                "cost_override": None,
                "note": "",
            }
        )
        refresh_lines()

    def _add_product() -> None:
        if not editable:
            return
        product = _pick_product()
        if not product:
            return
        _add_line(product, counted_delta=0.0)

    def _edit_line() -> None:
        if not selected_idx:
            show_error("Інвентаризація", "Оберіть рядок")
            return
        idx = selected_idx[0]
        ln = line_data[idx]
        editor = tk.Toplevel(dlg)
        editor.title("Рядок інвентаризації")
        editor.grab_set()
        ttk.Label(editor, text=f"{ln['sku']} — {ln['name']}").grid(
            row=0, column=0, columnspan=2, padx=8, pady=(8, 4), sticky="w"
        )
        ttk.Label(editor, text="Очікувано").grid(row=1, column=0, padx=8, pady=4, sticky="e")
        ttk.Label(editor, text=f"{ln['expected_qty']:.2f}").grid(row=1, column=1, padx=8, pady=4, sticky="w")
        ttk.Label(editor, text="Факт").grid(row=2, column=0, padx=8, pady=4, sticky="e")
        counted_var = tk.StringVar(value=f"{ln['counted_qty']:.2f}")
        ttk.Entry(editor, textvariable=counted_var, width=12, state="normal" if editable else "disabled").grid(
            row=2, column=1, padx=8, pady=4, sticky="w"
        )
        ttk.Label(editor, text="Собівартість надлишку").grid(row=3, column=0, padx=8, pady=4, sticky="e")
        cost_var = tk.StringVar(value=f"{ln['cost_override']:.2f}" if ln.get("cost_override") is not None else "")
        ttk.Entry(editor, textvariable=cost_var, width=12, state="normal" if editable else "disabled").grid(
            row=3, column=1, padx=8, pady=4, sticky="w"
        )
        ttk.Label(editor, text="Примітка").grid(row=4, column=0, padx=8, pady=4, sticky="e")
        note_var = tk.StringVar(value=ln.get("note", ""))
        ttk.Entry(editor, textvariable=note_var, width=30, state="normal" if editable else "disabled").grid(
            row=4, column=1, padx=8, pady=4, sticky="w"
        )

        def _save_line() -> None:
            if not editable:
                editor.destroy()
                return
            try:
                counted_val = float(counted_var.get() or 0)
                cost_val = cost_var.get().strip()
                cost_val_num = float(cost_val) if cost_val else None
            except ValueError:
                messagebox.showerror("Валідація", "Невірні числові значення")
                return
            if counted_val < 0:
                messagebox.showerror("Валідація", "Фактична кількість не може бути від'ємною")
                return
            if cost_val_num is not None and cost_val_num < 0:
                messagebox.showerror("Валідація", "Собівартість не може бути від'ємною")
                return
            ln["counted_qty"] = counted_val
            ln["cost_override"] = cost_val_num
            ln["note"] = note_var.get().strip()
            refresh_lines()
            editor.destroy()

        btn_frame = ttk.Frame(editor)
        btn_frame.grid(row=5, column=0, columnspan=2, pady=8)
        if editable:
            ttk.Button(btn_frame, text="Зберегти", command=_save_line).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_frame, text="Закрити", command=editor.destroy).pack(side=tk.LEFT, padx=4)

    def _delete_line() -> None:
        if not editable:
            return
        if not selected_idx:
            show_error("Інвентаризація", "Оберіть рядок")
            return
        idx = selected_idx[0]
        del line_data[idx]
        selected_idx.clear()
        refresh_lines()

    def _fill_counted() -> None:
        if not editable:
            return
        for ln in line_data:
            ln["counted_qty"] = ln["expected_qty"]
        refresh_lines()

    def _clear_counted() -> None:
        if not editable:
            return
        for ln in line_data:
            ln["counted_qty"] = 0.0
        refresh_lines()

    def _add_all_stock() -> None:
        if not editable:
            return
        wh_id = _current_warehouse_id()
        if not wh_id:
            show_error("Інвентаризація", "Оберіть склад.")
            return
        stock = db.stock_on_hand(wh_id)
        if not stock:
            messagebox.showinfo("Інвентаризація", "На складі немає залишків (qty>0).")
            refresh_lines()
            return
        added_any = False
        for product_id, qty in stock.items():
            if qty <= 0:
                continue
            if not _is_allowed_product(product_id):
                continue
            added_any = True
            existing = next((ln for ln in line_data if ln["product_id"] == product_id), None)
            if existing:
                existing["expected_qty"] = qty
                continue
            product = products_by_id.get(product_id)
            if not product:
                continue
            line_data.append(
                {
                    "product_id": product_id,
                    "sku": product["sku"],
                    "name": product["name"],
                    "expected_qty": float(qty),
                    "counted_qty": 0.0,
                    "cost_override": None,
                    "note": "",
                }
            )
        if (cat_var.get() or "Усі").strip() != "Усі" and not added_any:
            messagebox.showinfo("Інвентаризація", "У вибраній категорії немає залишків (qty>0).")
        refresh_lines()

    def _import_lines_csv() -> None:
        if not editable:
            return
        wh_id = _current_warehouse_id()
        if not wh_id:
            show_error("Інвентаризація", "Оберіть склад.")
            return
        file_path = filedialog.askopenfilename(
            title="Імпорт CSV",
            defaultextension=".csv",
            initialdir=str(get_data_dir()),
            filetypes=[("CSV", "*.csv"), ("Усі файли", "*.*")],
        )
        if not file_path:
            return
        not_found_codes: list[str] = []
        not_allowed_codes: list[str] = []
        bad_rows: list[str] = []
        updated_products: set[int] = set()
        new_products: set[int] = set()
        existing_product_ids = {ln["product_id"] for ln in line_data}
        prefix = _sanitize_barcode_prefix(settings.get("defaults", "product", "barcode_prefix") or "")
        try:
            def _read_csv_content(path: str) -> str:
                last_error = None
                for encoding in ("utf-8-sig", "cp1251"):
                    try:
                        with open(path, newline="", encoding=encoding) as f:
                            return f.read()
                    except UnicodeDecodeError as exc:
                        last_error = exc
                if last_error:
                    raise last_error
                raise OSError("CSV read failed")

            content = _read_csv_content(file_path)
            sample = content[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
                reader = csv.DictReader(io.StringIO(content), dialect=dialect)
            except csv.Error:
                reader = csv.DictReader(io.StringIO(content), delimiter=",")
            if not reader.fieldnames:
                messagebox.showerror("Імпорт CSV", "Файл не містить заголовків.")
                return
            field_map = {name.strip().lower(): name for name in reader.fieldnames if name}
            code_field = field_map.get("code") or field_map.get("sku")
            qty_field = field_map.get("qty") or field_map.get("counted_qty")
            if not code_field or not qty_field:
                messagebox.showerror("Імпорт CSV", "Потрібні колонки code/sku та qty/counted_qty.")
                return
            for row in reader:
                code = (row.get(code_field) or "").strip()
                if not code:
                    bad_rows.append("порожній код")
                    continue
                qty_raw = (row.get(qty_field) or "").strip()
                if "," in qty_raw and "." not in qty_raw:
                    qty_raw = qty_raw.replace(",", ".")
                try:
                    qty = float(qty_raw)
                except (TypeError, ValueError):
                    bad_rows.append(code)
                    continue
                if qty < 0:
                    bad_rows.append(code)
                    continue
                product = db.find_product_by_scan_code(code, barcode_prefix=prefix)
                if not product:
                    not_found_codes.append(code)
                    continue
                if not _is_allowed_product(int(product["id"])):
                    not_allowed_codes.append(code)
                    continue
                existing = next((ln for ln in line_data if ln["product_id"] == product["id"]), None)
                if existing:
                    existing["counted_qty"] += qty
                    if product["id"] in existing_product_ids:
                        updated_products.add(product["id"])
                    else:
                        new_products.add(product["id"])
                else:
                    expected_qty = db.get_stock_quantity(product["id"], wh_id)
                    line_data.append(
                        {
                            "product_id": product["id"],
                            "sku": product["sku"],
                            "name": product["name"],
                            "expected_qty": expected_qty,
                            "counted_qty": qty,
                            "cost_override": None,
                            "note": "",
                        }
                    )
                    new_products.add(product["id"])
        except Exception:
            logging.exception("Inventory CSV import error")
            show_error("Імпорт CSV", "Не вдалося імпортувати дані.")
            return
        refresh_lines()
        preview = ", ".join(not_found_codes[:20])
        summary = (
            "Імпорт завершено: "
            f"нових позицій: {len(new_products)}; "
            f"оновлено позицій: {len(updated_products)}; "
            f"рядків з помилками: {len(bad_rows)}; "
            f"не знайдено кодів: {len(not_found_codes)}."
        )
        if preview:
            summary += f"\nПроблемні коди: {preview}"
        messagebox.showinfo("Імпорт CSV", summary)
        if not_allowed_codes:
            skipped_preview = ", ".join(not_allowed_codes[:20])
            messagebox.showinfo(
                "Імпорт CSV",
                f"Пропущено (поза категорією): {skipped_preview}",
            )

    def _export_lines_csv() -> None:
        if not line_data:
            messagebox.showinfo("Експорт CSV", "Немає рядків для експорту.")
            return
        file_path = filedialog.asksaveasfilename(
            title="Експорт CSV",
            defaultextension=".csv",
            initialfile="inventory_lines.csv",
            initialdir=str(get_data_dir()),
            filetypes=[("CSV", "*.csv"), ("Усі файли", "*.*")],
        )
        if not file_path:
            return
        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["sku", "name", "expected_qty", "counted_qty", "diff", "cost_override", "note"])
            for ln in line_data:
                diff = ln["counted_qty"] - ln["expected_qty"]
                writer.writerow(
                    [
                        ln["sku"],
                        ln["name"],
                        ln["expected_qty"],
                        ln["counted_qty"],
                        diff,
                        ln["cost_override"] if ln.get("cost_override") is not None else "",
                        ln.get("note", ""),
                    ]
                )
        messagebox.showinfo("Експорт CSV", "Дані збережено.")

    btn_row = ttk.Frame(line_frame)
    btn_row.grid(row=2, column=0, columnspan=2, pady=(4, 0), sticky="w")
    ttk.Button(btn_row, text="Додати товар…", command=_add_product, state="normal" if editable else "disabled").pack(
        side=tk.LEFT, padx=4
    )
    ttk.Button(
        btn_row,
        text="Редагувати рядок…",
        command=_edit_line,
        state="normal" if editable else "disabled",
    ).pack(side=tk.LEFT, padx=4)
    ttk.Button(btn_row, text="Видалити рядок", command=_delete_line, state="normal" if editable else "disabled").pack(
        side=tk.LEFT, padx=4
    )
    ttk.Button(
        btn_row,
        text="Додати всі залишки (qty>0)",
        command=_add_all_stock,
        state="normal" if editable else "disabled",
    ).pack(side=tk.LEFT, padx=4)
    ttk.Button(btn_row, text="Імпорт CSV…", command=_import_lines_csv, state="normal" if editable else "disabled").pack(
        side=tk.LEFT, padx=4
    )
    ttk.Button(
        btn_row,
        text="Заповнити факт = очікувано",
        command=_fill_counted,
        state="normal" if editable else "disabled",
    ).pack(side=tk.LEFT, padx=4)
    ttk.Button(btn_row, text="Очистити факт", command=_clear_counted, state="normal" if editable else "disabled").pack(
        side=tk.LEFT, padx=4
    )
    ttk.Button(btn_row, text="Експорт CSV…", command=_export_lines_csv).pack(side=tk.LEFT, padx=4)

    def _on_scan(event=None) -> None:
        if not editable:
            return
        code = scan_var.get().strip()
        if not code:
            return
        prefix = _sanitize_barcode_prefix(settings.get("defaults", "product", "barcode_prefix") or "")
        product = db.find_product_by_scan_code(code, barcode_prefix=prefix)
        if not product:
            scan_status.config(text="Товар не знайдено", foreground="#b91c1c")
            scan_entry.bell()
            return
        if not _is_allowed_product(int(product["id"])):
            scan_status.config(text="Товар поза вибраною категорією", foreground="#b91c1c")
            scan_entry.bell()
            scan_var.set("")
            scan_entry.focus_set()
            return
        qty = float(scan_qty_var.get() or 1)
        _add_line(product, counted_delta=qty)
        scan_var.set("")
        scan_status.config(text=f"OK: {product['sku']} — {product['name']}", foreground="#15803d")
        scan_entry.focus_set()

    scan_entry.bind("<Return>", _on_scan)

    def _on_warehouse_change(event=None) -> None:
        balance_cache.clear()
        if not editable:
            return
        wh_id = _current_warehouse_id()
        if wh_id is None:
            return
        for ln in line_data:
            ln["expected_qty"] = db.get_stock_quantity(ln["product_id"], wh_id)
        refresh_lines()

    wh_combo.bind("<<ComboboxSelected>>", _on_warehouse_change)
    cat_combo.bind("<<ComboboxSelected>>", lambda _event: _rebuild_allowed_products())
    include_children_check.configure(command=_rebuild_allowed_products)
    _rebuild_allowed_products()

    refresh_lines()
    if editable:
        scan_entry.focus_set()

    result: dict[str, object] = {}

    def _validate_lines() -> bool:
        missing = [ln for ln in line_data if _needs_cost(ln)]
        if missing:
            messagebox.showerror("Валідація", "Для надлишку потрібна собівартість.")
            return False
        for ln in line_data:
            if ln["counted_qty"] < 0:
                messagebox.showerror("Валідація", "Фактична кількість не може бути від'ємною.")
                return False
            if ln.get("cost_override") is not None and ln["cost_override"] < 0:
                messagebox.showerror("Валідація", "Собівартість не може бути від'ємною.")
                return False
        return True

    def _on_ok() -> None:
        if not editable:
            dlg.destroy()
            return
        wh_id = _current_warehouse_id()
        if not wh_id:
            messagebox.showerror("Валідація", "Оберіть склад.")
            return
        if not _validate_lines():
            return
        info = {
            "doc_date": _parse_date_value(_current_date()),
            "warehouse_id": wh_id,
            "comment": comment_var.get().strip(),
        }
        result["info"] = info
        result["lines"] = [
            {
                "product_id": ln["product_id"],
                "expected_qty": ln["expected_qty"],
                "counted_qty": ln["counted_qty"],
                "cost_override": ln.get("cost_override"),
                "note": ln.get("note", ""),
            }
            for ln in line_data
        ]
        result["post_now"] = bool(post_now_var.get())
        dlg.destroy()

    def _on_cancel() -> None:
        dlg.destroy()

    btns = ttk.Frame(content)
    btns.grid(row=row_idx + 1, column=0, columnspan=2, pady=8)
    if editable:
        ttk.Button(btns, text="OK", command=_on_ok).pack(side=tk.LEFT, padx=6)
        ttk.Button(btns, text="Скасувати", command=_on_cancel).pack(side=tk.LEFT, padx=6)
    else:
        ttk.Button(btns, text="Закрити", command=_on_cancel).pack(side=tk.LEFT, padx=6)

    dlg.wait_window()
    if "info" not in result:
        return None
    return result["info"], result["lines"], result["post_now"]


def document_prompt(
    doc_type: str,
    products,
    counterparties,
    warehouses,
    channels,
    currencies,
    doc=None,
    lines=None,
    settings: Settings | None = None,
):
    dlg = tk.Toplevel()
    dlg.title("Документ")
    dlg.grab_set()
    editable = not doc or doc["status"] == "draft"

    if doc and isinstance(doc, sqlite3.Row):
        doc = dict(doc)

    dlg.columnconfigure(0, weight=1)
    dlg.rowconfigure(0, weight=1)
    content = ttk.Frame(dlg, padding=10)
    content.grid(row=0, column=0, sticky="nsew")
    content.columnconfigure(1, weight=1)

    row_idx = 0
    ttk.Label(content, text="Тип").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    doc_type_label = "Закупівля" if doc_type == "purchase" else "Продаж"
    ttk.Label(content, text=doc_type_label).grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")

    row_idx += 1
    ttk.Label(content, text="Дата (YYYY-MM-DD)").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    date_var = tk.StringVar(value=doc["doc_date"] if doc else datetime.now().strftime("%Y-%m-%d"))
    ttk.Entry(content, textvariable=date_var, width=15, state="normal" if editable else "disabled").grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")

    row_idx += 1
    ttk.Label(content, text="Валюта").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    curr_var = tk.StringVar(
        value=doc["currency_code"]
        if doc
        else (currencies[0]["code"] if currencies else get_base_currency_code())
    )
    curr_codes = [c["code"] for c in currencies] if currencies else [get_base_currency_code()]
    curr_combo = ttk.Combobox(content, textvariable=curr_var, values=curr_codes, state="readonly")
    if not editable:
        curr_combo.state(["disabled"])
    curr_combo.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
    last_currency = curr_var.get()

    row_idx += 1
    ttk.Label(content, text="Курс до базової").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    try:
        default_rate = doc["exchange_rate"] if doc else ensure_rate_for_date(curr_var.get(), date_var.get())
    except Exception:
        default_rate = doc["exchange_rate"] if doc else 1.0
    rate_var = tk.StringVar(value=f"{default_rate:.4f}")
    rate_entry = ttk.Entry(content, textvariable=rate_var, width=12, state="normal" if editable else "disabled")
    rate_entry.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")

    row_idx += 1
    ttk.Label(content, text="Склад").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    wh_var = tk.StringVar()
    wh_names = [w["name"] for w in warehouses]
    wh_combo = ttk.Combobox(content, textvariable=wh_var, values=wh_names, state="readonly")
    wh_combo.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")

    row_idx += 1
    ch_var = tk.StringVar()
    ch_combo = None
    if doc_type == "sale":
        ttk.Label(content, text="Канал").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
        ch_names = [c["name"] for c in channels]
        ch_combo = ttk.Combobox(content, textvariable=ch_var, values=ch_names, state="readonly")
        ch_combo.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
        row_idx += 1

    ttk.Label(content, text="Контрагент").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    allowed_types = {"purchase": {"supplier", "both", "other"}, "sale": {"customer", "both", "other"}}[doc_type]
    filtered_counterparties = [c for c in counterparties if c["type"] in allowed_types]
    cp_names = ["-"] + [c["name"] for c in filtered_counterparties]
    cp_var = tk.StringVar()
    cp_combo = ttk.Combobox(content, textvariable=cp_var, values=cp_names, state="readonly", width=25)
    cp_combo.grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
    row_idx += 1

    ttk.Label(content, text="Коментар").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
    comment_var = tk.StringVar(value=doc["comment"] if doc else "")
    ttk.Entry(content, textvariable=comment_var, width=40).grid(row=row_idx, column=1, padx=6, pady=4, sticky="ew")

    row_idx += 1
    order_expense_var = tk.StringVar(value=f"{float(doc.get('order_expense_doc', 0.0)):.2f}" if doc else "0")
    if doc_type == "sale":
        ttk.Label(content, text="Витрати замовлення").grid(row=row_idx, column=0, padx=6, pady=4, sticky="e")
        ttk.Entry(content, textvariable=order_expense_var, width=20).grid(row=row_idx, column=1, padx=6, pady=4, sticky="w")
        row_idx += 1

    scan_frame = ttk.LabelFrame(content, text="Сканування")
    scan_frame.grid(row=row_idx, column=0, columnspan=2, padx=6, pady=6, sticky="ew")
    scan_frame.columnconfigure(1, weight=1)

    ttk.Label(scan_frame, text="Скан-код").grid(row=0, column=0, padx=6, pady=4, sticky="e")
    scan_var = tk.StringVar()
    scan_entry = ttk.Entry(scan_frame, textvariable=scan_var, width=30, state="normal" if editable else "disabled")
    scan_entry.grid(row=0, column=1, padx=6, pady=4, sticky="w")
    ttk.Label(scan_frame, text="К-сть при скані").grid(row=0, column=2, padx=6, pady=4, sticky="e")
    scan_qty_var = tk.IntVar(value=1)
    scan_qty_spin = ttk.Spinbox(
        scan_frame,
        from_=1,
        to=999,
        textvariable=scan_qty_var,
        width=6,
        state="normal" if editable else "disabled",
    )
    scan_qty_spin.grid(row=0, column=3, padx=6, pady=4, sticky="w")
    scan_plus_one_var = tk.BooleanVar(value=False)
    scan_plus_one = ttk.Checkbutton(
        scan_frame,
        text="Кожен скан = +1",
        variable=scan_plus_one_var,
        state="normal" if editable else "disabled",
    )
    scan_plus_one.grid(row=0, column=4, padx=6, pady=4, sticky="w")
    scan_use_avg_cost_var = tk.BooleanVar(value=True)
    ttk.Checkbutton(
        scan_frame,
        text="Ціна зі середньої собівартості",
        variable=scan_use_avg_cost_var,
    ).grid(row=1, column=1, padx=6, pady=(0, 4), sticky="w")
    scan_status = ttk.Label(scan_frame, text="")
    scan_status.grid(row=2, column=0, columnspan=5, padx=6, pady=(0, 4), sticky="w")

    def _apply_scan_mode() -> None:
        if scan_plus_one_var.get():
            scan_qty_var.set(1)
            scan_qty_spin.config(state="disabled")
        else:
            scan_qty_spin.config(state="normal" if editable else "disabled")

    scan_plus_one.config(command=_apply_scan_mode)
    _apply_scan_mode()

    row_idx += 1
    ttk.Label(content, text="Рядки").grid(row=row_idx, column=0, padx=6, pady=4, sticky="ne")
    line_frame = ttk.Frame(content)
    line_frame.grid(row=row_idx, column=1, padx=6, pady=4, sticky="nsew")
    line_frame.grid_columnconfigure(0, weight=1)
    content.rowconfigure(row_idx, weight=1)

    columns = ["product", "quantity", "price", "expense", "amount"]
    tree = ttk.Treeview(line_frame, columns=columns, show="headings", height=8)
    headings = {
        "product": ("Товар", 200),
        "quantity": ("Кількість", 90),
        "price": ("Ціна", 90),
        "expense": ("Витрата/од.", 110),
        "amount": ("Сума", 90),
    }
    for col, (title, width) in headings.items():
        tree.heading(col, text=title)
        tree.column(col, width=width, anchor="w")
    tree.grid(row=0, column=0, sticky="nsew")
    yscroll = ttk.Scrollbar(line_frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=yscroll.set)
    yscroll.grid(row=0, column=1, sticky="ns")
    line_frame.grid_rowconfigure(0, weight=1)

    def refresh_currency_ui() -> None:
        price_label.config(text=f"Ціна ({curr_var.get()})")
        tree.heading("price", text=f"Ціна ({curr_var.get()})")
        expense_label.config(text=f"Витрата/од. ({curr_var.get()})")
        tree.heading("expense", text=f"Витрата/од. ({curr_var.get()})")
        tree.heading("amount", text=f"Сума ({curr_var.get()})")

    def on_currency_change(event=None):
        nonlocal last_currency
        if editable:
            try:
                rate_val = ensure_rate_for_date(curr_var.get(), date_var.get())
            except ValueError as exc:
                messagebox.showerror("Курс", str(exc))
                curr_var.set(last_currency)
                return
            rate_var.set(f"{rate_val:.4f}")
            last_currency = curr_var.get()
        refresh_currency_ui()

    curr_combo.bind("<<ComboboxSelected>>", on_currency_change)

    product_lookup = {f"{p['name']} ({p['sku']})": p["id"] for p in products}
    products_by_id = {p["id"]: f"{p['name']} ({p['sku']})" for p in products}
    product_names = list(product_lookup.keys())

    row_idx += 1
    entry_frame = ttk.Frame(content)
    entry_frame.grid(row=row_idx, column=0, columnspan=2, padx=6, pady=4, sticky="ew")
    entry_frame.columnconfigure(1, weight=1)
    ttk.Label(entry_frame, text="Товар").grid(row=0, column=0, padx=4, pady=2, sticky="e")
    product_var = tk.StringVar()
    product_combo_state = "normal" if editable else "readonly"
    product_combo = ttk.Combobox(entry_frame, textvariable=product_var, values=product_names, state=product_combo_state, width=40)
    product_combo.grid(row=0, column=1, padx=4, pady=2, sticky="ew")
    if product_lookup:
        product_combo.current(0)

    def filter_products(event=None):
        if not editable:
            return
        text = product_var.get().lower()
        matches = [name for name in product_names if text in name.lower()]
        product_combo["values"] = matches if matches else product_names

    product_combo.bind("<KeyRelease>", filter_products)

    ttk.Label(entry_frame, text="Кількість").grid(row=0, column=2, padx=4, pady=2, sticky="e")
    qty_var = tk.StringVar(value="1")
    ttk.Entry(entry_frame, textvariable=qty_var, width=10).grid(row=0, column=3, padx=4, pady=2, sticky="w")

    price_label = ttk.Label(entry_frame, text="Ціна")
    price_label.grid(row=0, column=4, padx=4, pady=2, sticky="e")
    price_var = tk.StringVar(value="0")
    ttk.Entry(entry_frame, textvariable=price_var, width=10).grid(row=0, column=5, padx=4, pady=2, sticky="w")

    expense_label = ttk.Label(entry_frame, text="Витрата/од.")
    expense_label.grid(row=0, column=6, padx=4, pady=2, sticky="e")
    expense_var = tk.StringVar(value="0")
    ttk.Entry(entry_frame, textvariable=expense_var, width=10).grid(row=0, column=7, padx=4, pady=2, sticky="w")

    line_data = []
    if lines:
        for ln in lines:
            price_field = "purchase_price" if doc_type == "purchase" else "sale_price"
            expense_value = 0.0
            if doc_type == "sale" and "unit_expense_doc" in ln.keys():
                try:
                    expense_value = float(ln["unit_expense_doc"])
                except Exception:
                    expense_value = 0.0
            line_data.append(
                {
                    "product_id": ln["product_id"],
                    "product_name": ln["product_name"],
                    "quantity": float(ln["quantity"]),
                    "price": float(ln[price_field]),
                    "expense": expense_value,
                    "amount": float(ln["quantity"]) * float(ln[price_field]),
                }
            )

    if doc:
        if doc["warehouse_id"]:
            try:
                wh_combo.current(next(i for i, w in enumerate(warehouses) if w["id"] == doc["warehouse_id"]))
            except StopIteration:
                wh_combo.set(warehouses[0]["name"] if warehouses else "")
        if curr_codes:
            try:
                curr_combo.current(curr_codes.index(doc.get("currency_code", curr_codes[0])))
            except ValueError:
                curr_combo.current(0)
        if doc_type == "sale" and ch_combo:
            if doc["channel"]:
                try:
                    ch_combo.current(next(i for i, c in enumerate(channels) if c["name"] == doc["channel"]))
                except StopIteration:
                    ch_combo.set(channels[0]["name"] if channels else "")
            elif channels:
                ch_combo.current(0)
        if doc.get("supplier_id"):
            target = next((c["name"] for c in filtered_counterparties if c["id"] == doc.get("supplier_id")), "-")
            cp_var.set(target)
        if doc.get("customer_id"):
            target = next((c["name"] for c in filtered_counterparties if c["id"] == doc.get("customer_id")), "-")
            cp_var.set(target)
    else:
        if warehouses:
            wh_combo.current(0)
        if ch_combo and channels:
            ch_combo.current(0)
        if curr_codes:
            curr_combo.current(0)
        cp_var.set("-")

    selected_idx: list[int] = []

    refresh_currency_ui()

    def refresh_lines():
        tree.delete(*tree.get_children())
        for idx, ln in enumerate(line_data):
            tree.insert(
                "",
                "end",
                iid=str(idx),
                values=(
                    ln["product_name"],
                    f"{ln['quantity']:.2f}",
                    f"{ln['price']:.2f}",
                    f"{ln.get('expense', 0.0):.2f}",
                    f"{ln['amount']:.2f}",
                ),
            )

    def _scan_get_counterparty_id() -> Optional[int]:
        cp_name = cp_var.get()
        if not cp_name or cp_name == "-":
            return None
        match = next((c for c in filtered_counterparties if c["name"] == cp_name), None)
        return match["id"] if match else None

    def _current_warehouse_id() -> int | None:
        name = (wh_var.get() or "").strip()
        w = next((x for x in warehouses if x["name"] == name), None)
        return int(w["id"]) if w else None

    def _scan_resolve_product(code: str) -> Optional[sqlite3.Row]:
        if settings is None:
            prefix = ""
        elif isinstance(settings, Settings):
            prefix = settings.get("defaults", "product", "barcode_prefix") or ""
        else:
            prefix = (settings.get("defaults", {}).get("product", {}).get("barcode_prefix") or "")
        prefix = prefix.strip()
        if doc_type == "purchase":
            counterparty_id = _scan_get_counterparty_id()
            if counterparty_id:
                product = db.get_product_by_supplier_code(counterparty_id, code)
                if product:
                    return product
        product = db.find_product_by_scan_code(code, barcode_prefix=prefix)
        if product:
            return product
        product = db.find_product_by_sku_or_name(None, None, supplier_sku=code)
        if product:
            return product
        supplier_lookup = getattr(db, "get_product_by_supplier_sku", None)
        if callable(supplier_lookup):
            return supplier_lookup(code)
        return None

    def _scan_add_line(product_row: sqlite3.Row, qty_delta: float) -> None:
        for idx, ln in enumerate(line_data):
            if ln["product_id"] == product_row["id"]:
                ln["quantity"] += qty_delta
                ln["amount"] = ln["quantity"] * ln["price"]
                refresh_lines()
                tree.selection_set(str(idx))
                tree.focus(str(idx))
                tree.see(str(idx))
                return
        price0 = 0.0
        if scan_use_avg_cost_var.get():
            wh_id = _current_warehouse_id()
            if wh_id:
                _, avg_cost = db.get_stock_balance(int(product_row["id"]), int(wh_id))
                if doc_type == "sale":
                    try:
                        rate = float(rate_var.get() or 1)
                    except ValueError:
                        rate = 1.0
                    if rate <= 0:
                        rate = 1.0
                    price0 = avg_cost / rate
                else:
                    price0 = avg_cost
        if price0 <= 0:
            try:
                price0 = float(price_var.get() or 0)
            except ValueError:
                price0 = 0.0
        try:
            exp0 = float(expense_var.get() or 0)
        except ValueError:
            exp0 = 0.0
        product_name = f"{product_row['name']} ({product_row['sku']})"
        line_data.append(
            {
                "product_id": product_row["id"],
                "product_name": product_name,
                "quantity": qty_delta,
                "price": price0,
                "expense": exp0,
                "amount": qty_delta * price0,
            }
        )
        refresh_lines()
        idx = len(line_data) - 1
        tree.selection_set(str(idx))
        tree.focus(str(idx))
        tree.see(str(idx))

    def _on_scan_commit(event=None):
        if not editable:
            return "break"
        code = scan_var.get().strip()
        if not code:
            return "break"
        try:
            qty = 1 if scan_plus_one_var.get() else int(scan_qty_var.get() or 1)
        except (TypeError, ValueError):
            qty = 1
        product = _scan_resolve_product(code)
        if not product:
            scan_status.config(text=f"Не знайдено: {code}", foreground="#b91c1c")
            dlg.bell()
            scan_entry.focus_set()
            scan_var.set("")
            return "break"
        _scan_add_line(product, qty)
        scan_var.set("")
        scan_status.config(text=f"OK: {product['sku']} — {product['name']} (+{qty})", foreground="#15803d")
        scan_entry.focus_set()
        return "break"

    scan_entry.bind("<Return>", _on_scan_commit)
    scan_entry.bind("<KP_Enter>", _on_scan_commit)

    def on_select(event=None):
        selected_idx.clear()
        sel = tree.selection()
        if sel:
            idx = int(sel[0])
            selected_idx.append(idx)
            ln = line_data[idx]
            product_name = products_by_id.get(ln["product_id"], ln["product_name"])
            product_var.set(product_name)
            qty_var.set(str(ln["quantity"]))
            price_var.set(str(ln["price"]))
            expense_var.set(str(ln.get("expense", 0.0)))

    tree.bind("<<TreeviewSelect>>", on_select)

    def add_or_update_line():
        if not editable:
            return
        try:
            qty = float(qty_var.get())
            price = float(price_var.get())
            expense_value = float(expense_var.get() or 0)
        except ValueError:
            messagebox.showerror("Валідація", "Невірні числові значення")
            return
        if qty <= 0:
            messagebox.showerror("Валідація", "Кількість повинна бути більшою за 0")
            return
        if expense_value < 0:
            messagebox.showerror("Валідація", "Витрати не можуть бути від'ємними")
            return
        product_name = product_var.get().strip()
        product_id = product_lookup.get(product_name)
        if not product_id:
            if not editable:
                return
            if not product_name:
                messagebox.showerror("Валідація", "Введіть назву товару")
                return
            created = add_new_product(product_name)
            if not created:
                return
            product_id, product_name = created
        data = {
            "product_id": product_id,
            "product_name": product_name,
            "quantity": qty,
            "price": price,
            "expense": expense_value,
            "amount": qty * price,
        }
        if selected_idx:
            line_data[selected_idx[0]] = data
        else:
            line_data.append(data)
        refresh_lines()
        selected_idx.clear()

    def delete_line():
        if not editable:
            return
        if not selected_idx:
            return
        line_data.pop(selected_idx[0])
        selected_idx.clear()
        refresh_lines()

    def add_new_product(default_name: str = ""):
        if not editable:
            return None
        brands = db.list_brands()
        categories = [
            {
                **c,
                "label": c.get("label") or ("    " * c.get("depth", 0) + c.get("name", "")),
            }
            for c in db.list_categories_tree()
        ]
        if not brands or not categories:
            messagebox.showerror(
                "Товари",
                "Додайте принаймні один бренд і категорію у вкладці \"Товари\", щоб створювати нові позиції.",
            )
            return None

        dlg_product = tk.Toplevel(dlg)
        dlg_product.title("Новий товар")
        dlg_product.grab_set()

        ttk.Label(dlg_product, text="Артикул").grid(row=0, column=0, padx=6, pady=4, sticky="e")
        sku_var = tk.StringVar(value=default_name)
        ttk.Entry(dlg_product, textvariable=sku_var, width=30).grid(row=0, column=1, padx=6, pady=4, sticky="w")

        ttk.Label(dlg_product, text="Назва").grid(row=1, column=0, padx=6, pady=4, sticky="e")
        name_var = tk.StringVar(value=default_name)
        ttk.Entry(dlg_product, textvariable=name_var, width=30).grid(row=1, column=1, padx=6, pady=4, sticky="w")

        ttk.Label(dlg_product, text="Бренд").grid(row=2, column=0, padx=6, pady=4, sticky="e")
        brand_var = tk.StringVar()
        brand_combo = ttk.Combobox(
            dlg_product,
            textvariable=brand_var,
            values=[b["name"] for b in brands],
            state="readonly",
            width=28,
        )
        brand_combo.grid(row=2, column=1, padx=6, pady=4, sticky="w")

        preferred_brand = _find_index_by_name([b["name"] for b in brands], self.settings.get("defaults", "product", "brand"))
        if preferred_brand is not None:
            brand_combo.current(preferred_brand)
        else:
            brand_combo.current(0)

        ttk.Label(dlg_product, text="Категорія").grid(row=3, column=0, padx=6, pady=4, sticky="e")
        category_var = tk.StringVar()
        category_combo = ttk.Combobox(
            dlg_product,
            textvariable=category_var,
            values=[c["label"] for c in categories],
            state="readonly",
            width=28,
        )
        category_combo.grid(row=3, column=1, padx=6, pady=4, sticky="w")
        default_category = self.settings.get("defaults", "product", "category") or ""
        preferred_category = _find_index_by_name([c["label"] for c in categories], default_category)
        if preferred_category is not None:
            category_combo.current(preferred_category)
        else:
            category_combo.current(0)

        default_unit = self.settings.get("defaults", "product", "unit") or "pcs"
        ttk.Label(dlg_product, text="Одиниця").grid(row=4, column=0, padx=6, pady=4, sticky="e")
        unit_var = tk.StringVar(value=default_unit)
        ttk.Entry(dlg_product, textvariable=unit_var, width=30).grid(row=4, column=1, padx=6, pady=4, sticky="w")

        result_new: tuple[int, str] | None = None

        def on_save():
            nonlocal result_new
            sku = sku_var.get().strip()
            name = name_var.get().strip()
            unit = unit_var.get().strip() or "pcs"
            if not sku or not name:
                messagebox.showerror("Товари", "Введіть артикул і назву товару")
                return
            brand_idx = brand_combo.current()
            cat_idx = category_combo.current()
            try:
                brand_id = brands[brand_idx]["id"]
                category_id = categories[cat_idx]["id"]
            except Exception:
                messagebox.showerror("Товари", "Оберіть бренд та категорію")
                return
            try:
                new_id = db.add_product(sku, name, brand_id, category_id, unit, True)
            except Exception as exc:
                messagebox.showerror("Товари", f"Не вдалося створити товар: {exc}")
                return
            product_full_name = f"{name} ({sku})"
            product_lookup[product_full_name] = new_id
            products_by_id[new_id] = product_full_name
            product_names.append(product_full_name)
            product_names.sort(key=str.lower)
            product_combo["values"] = product_names
            product_var.set(product_full_name)
            result_new = (new_id, product_full_name)
            dlg_product.destroy()

        def on_cancel():
            dlg_product.destroy()

        btns_new = ttk.Frame(dlg_product)
        btns_new.grid(row=5, column=0, columnspan=2, pady=8)
        ttk.Button(btns_new, text="Зберегти", command=on_save).pack(side=tk.LEFT, padx=4)
        ttk.Button(btns_new, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
        dlg_product.bind("<Return>", lambda e: on_save())
        dlg_product.bind("<Escape>", lambda e: on_cancel())
        dlg_product.wait_window()
        return result_new

    btn_line = ttk.Frame(entry_frame)
    btn_line.grid(row=0, column=8, padx=6)
    ttk.Button(btn_line, text="Новий товар", command=lambda: add_new_product(product_var.get()), state="normal" if editable else "disabled").pack(side=tk.LEFT, padx=4)
    ttk.Button(btn_line, text="Додати/Оновити", command=add_or_update_line, state="normal" if editable else "disabled").pack(side=tk.LEFT)
    ttk.Button(btn_line, text="Видалити", command=delete_line, state="normal" if editable else "disabled").pack(side=tk.LEFT, padx=4)

    refresh_lines()

    row_idx += 1

    result = None

    def on_ok():
        nonlocal result
        try:
            datetime.fromisoformat(date_var.get())
        except ValueError:
            messagebox.showerror("Валідація", "Невірний формат дати (YYYY-MM-DD)")
            return
        if editable and not line_data:
            messagebox.showerror("Валідація", "Додайте хоча б один рядок")
            return
        try:
            rate = float(rate_var.get())
        except ValueError:
            messagebox.showerror("Валідація", "Невірний курс")
            return
        if rate <= 0:
            messagebox.showerror("Валідація", "Курс має бути більшим за 0")
            return
        currency_code = curr_var.get().strip().upper()
        order_expense = 0.0
        if doc_type == "sale":
            try:
                order_expense = float(order_expense_var.get() or 0)
            except ValueError:
                messagebox.showerror("Валідація", "Невірна сума витрат замовлення")
                return
            if order_expense < 0:
                messagebox.showerror("Валідація", "Витрати замовлення не можуть бути від'ємними")
                return
        cp_name = cp_var.get()
        cp_id = None
        if cp_name and cp_name != "-":
            cp_id = next((c["id"] for c in filtered_counterparties if c["name"] == cp_name), None)
        if currency_code != get_base_currency_code() and not db.rate_on_date(currency_code, date_var.get()):
            db.add_currency_rate(currency_code, date_var.get(), rate)
        try:
            warehouse_id = warehouses[wh_combo.current()]["id"]
        except Exception:
            messagebox.showerror("Валідація", "Оберіть склад")
            return
        channel_name = ""
        if doc_type == "sale":
            channel_name = ch_var.get() if ch_var.get() else (channels[0]["name"] if channels else "")
        info = {
            "doc_type": doc_type,
            "doc_date": date_var.get().strip(),
            "counterparty_id": cp_id,
            "warehouse_id": warehouse_id,
            "channel": channel_name,
            "comment": comment_var.get().strip(),
            "currency": curr_var.get(),
            "rate": rate,
            "order_expense_doc": order_expense if doc_type == "sale" else 0.0,
        }
        if doc_type == "sale":
            lines_to_save = [
                (ln["product_id"], ln["quantity"], ln["price"], ln.get("expense", 0.0)) for ln in line_data
            ]
        else:
            lines_to_save = [(ln["product_id"], ln["quantity"], ln["price"]) for ln in line_data]
        result = (info, lines_to_save)
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    def _focus_scan(event=None):
        if editable:
            scan_entry.focus_set()
            scan_entry.selection_range(0, tk.END)
        return "break"

    def _toggle_plus_one(event=None):
        if editable:
            scan_plus_one_var.set(not scan_plus_one_var.get())
            _apply_scan_mode()
            scan_entry.focus_set()
        return "break"

    btns = ttk.Frame(content)
    btns.grid(row=row_idx, column=0, columnspan=2, pady=8, sticky="e")
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Control-Return>", lambda e: on_ok())
    dlg.bind("<Control-KP_Enter>", lambda e: on_ok())
    dlg.bind("<F8>", _focus_scan)
    dlg.bind("<F9>", _toggle_plus_one)
    dlg.bind("<Escape>", lambda e: on_cancel())
    if editable:
        dlg.after_idle(
            lambda: (
                scan_entry.focus_set(),
                scan_entry.selection_range(0, tk.END),
                scan_entry.icursor(tk.END),
            )
        )
    dlg.wait_window()
    return result


def cash_prompt(counterparties, channels):
    dlg = tk.Toplevel()
    dlg.title("Рух коштів")
    dlg.grab_set()

    ttk.Label(dlg, text="Дата (YYYY-MM-DD)").grid(row=0, column=0, padx=6, pady=4, sticky="w")
    date_var = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
    ttk.Entry(dlg, textvariable=date_var, width=15).grid(row=0, column=1, padx=6, pady=4, sticky="w")

    ttk.Label(dlg, text="Тип").grid(row=1, column=0, padx=6, pady=4, sticky="w")
    type_var = tk.StringVar()
    types = [
        ("Оплата від клієнта", "sale_payment"),
        ("Оплата постачальнику", "purchase_payment"),
        ("Інший дохід", "other_income"),
        ("Змінна витрата", "other_variable_expense"),
        ("Постійна витрата", "fixed_expense"),
    ]
    type_combo = ttk.Combobox(dlg, textvariable=type_var, values=[t[0] for t in types], state="readonly")
    type_combo.grid(row=1, column=1, padx=6, pady=4, sticky="w")
    type_combo.current(0)

    ttk.Label(dlg, text="Сума (+ вхід, - вихід)").grid(row=2, column=0, padx=6, pady=4, sticky="w")
    amount_var = tk.StringVar(value="0")
    ttk.Entry(dlg, textvariable=amount_var, width=15).grid(row=2, column=1, padx=6, pady=4, sticky="w")

    ttk.Label(dlg, text="Контрагент").grid(row=3, column=0, padx=6, pady=4, sticky="w")
    cp_var = tk.StringVar()
    cp_names = ["-"] + [c["name"] for c in counterparties]
    cp_combo = ttk.Combobox(dlg, textvariable=cp_var, values=cp_names, state="readonly", width=30)
    cp_combo.grid(row=3, column=1, padx=6, pady=4, sticky="w")
    cp_combo.current(0)

    ttk.Label(dlg, text="Канал").grid(row=4, column=0, padx=6, pady=4, sticky="w")
    ch_var = tk.StringVar()
    ch_names = [c["name"] for c in channels]
    ch_combo = ttk.Combobox(dlg, textvariable=ch_var, values=ch_names, state="readonly", width=20)
    ch_combo.grid(row=4, column=1, padx=6, pady=4, sticky="w")
    if channels:
        ch_combo.current(0)

    ttk.Label(dlg, text="Коментар").grid(row=5, column=0, padx=6, pady=4, sticky="w")
    comment_var = tk.StringVar()
    ttk.Entry(dlg, textvariable=comment_var, width=40).grid(row=5, column=1, padx=6, pady=4, sticky="w")

    result: list[dict] | None = None

    def on_ok():
        nonlocal result
        try:
            datetime.fromisoformat(date_var.get())
            amount = float(amount_var.get())
        except ValueError:
            messagebox.showerror("Валідація", "Невірні значення дати або суми")
            return
        ctype = next((t[1] for t in types if t[0] == type_var.get()), types[0][1])
        cp_name = cp_var.get()
        cp_id = None
        if cp_name and cp_name != "-":
            cp_id = next((c["id"] for c in counterparties if c["name"] == cp_name), None)
        channel = ch_var.get() if ch_var.get() else ""
        result = [
            {
                "date": date_var.get().strip(),
                "amount": amount,
                "ctype": ctype,
                "counterparty_id": cp_id,
                "related_doc_type": None,
                "related_doc_id": None,
                "channel": channel,
                "comment": comment_var.get().strip(),
            }
        ]
        dlg.destroy()

    def on_cancel():
        dlg.destroy()

    btns = ttk.Frame(dlg)
    btns.grid(row=6, column=0, columnspan=2, pady=8)
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=4)
    ttk.Button(btns, text="Скасувати", command=on_cancel).pack(side=tk.LEFT, padx=4)
    dlg.bind("<Return>", lambda e: on_ok())
    dlg.bind("<Escape>", lambda e: on_cancel())
    dlg.wait_window()
    return result


class SettingsDialog(tk.Toplevel):
    def __init__(self, app: InventoryApp, section: str):
        super().__init__(app)
        self.app = app
        self.section = section
        self.title("Налаштування")
        self.resizable(False, False)
        self.grab_set()
        self.vars: dict[str, tk.Variable] = {}
        self.recent_cleared = False

        notebook = ttk.Notebook(self)
        notebook.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.general_tab = ttk.Frame(notebook)
        self.defaults_tab = ttk.Frame(notebook)
        self.files_tab = ttk.Frame(notebook)
        self.ui_tab = ttk.Frame(notebook)
        self.editor_tab = ttk.Frame(notebook)
        self.hotkeys_tab = ttk.Frame(notebook)
        self.support_tab = ttk.Frame(notebook)

        notebook.add(self.general_tab, text="Загальні")
        notebook.add(self.defaults_tab, text="Типові значення")
        notebook.add(self.files_tab, text="Файли й шляхи")
        notebook.add(self.ui_tab, text="Інтерфейс і вікна")
        notebook.add(self.editor_tab, text="Редактор")
        notebook.add(self.hotkeys_tab, text="Гарячі клавіші")
        notebook.add(self.support_tab, text="Допомога й підтримка")

        self.build_general_tab()
        self.build_defaults_tab()
        self.build_files_tab()
        self.build_ui_tab()
        self.build_editor_tab()
        self.build_hotkeys_tab()
        self.build_support_tab()

        tab_index = {
            "general": 0,
            "defaults": 1,
            "files": 2,
            "ui": 3,
            "editor": 4,
            "hotkeys": 5,
            "support": 6,
        }.get(section, 0)
        notebook.select(tab_index)

        buttons = ttk.Frame(self)
        buttons.pack(fill=tk.X, pady=(0, 8))
        ttk.Button(buttons, text="Скасувати", command=self.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(buttons, text="Зберегти", command=self.save).pack(side=tk.RIGHT, padx=4)

    def build_general_tab(self) -> None:
        language_var = self._add_var("general.language", tk.StringVar(value=self.app.settings.get("general", "language")))
        ttk.Label(self.general_tab, text="Мова інтерфейсу:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(self.general_tab, textvariable=language_var, values=["uk", "en", "pl", "de"], width=10).grid(
            row=0, column=1, sticky="w", padx=6, pady=4
        )

        theme_var = self._add_var("general.theme", tk.StringVar(value=self.app.settings.get("general", "theme")))
        ttk.Label(self.general_tab, text="Тема:").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(
            self.general_tab, textvariable=theme_var, values=["system", "light", "dark"], width=10
        ).grid(row=1, column=1, sticky="w", padx=6, pady=4)

    def build_defaults_tab(self) -> None:
        product_frame = ttk.LabelFrame(self.defaults_tab, text="Товари")
        product_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        default_unit = self.app.settings.get("defaults", "product", "unit") or "pcs"
        unit_var = self._add_var("defaults.product.unit", tk.StringVar(value=default_unit))
        ttk.Label(product_frame, text="Одиниця за замовчуванням:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(product_frame, textvariable=unit_var, width=14).grid(row=0, column=1, sticky="w", padx=6, pady=4)

        brands = db.list_brands()
        brand_names = [b["name"] for b in brands]
        brand_default = self.app.settings.get("defaults", "product", "brand") or ""
        brand_var = self._add_var("defaults.product.brand", tk.StringVar(value=brand_default))
        ttk.Label(product_frame, text="Бренд за замовчуванням:").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(product_frame, textvariable=brand_var, values=brand_names, width=30).grid(
            row=1, column=1, sticky="w", padx=6, pady=4
        )
        ttk.Label(product_frame, text="Залиште поле порожнім, щоб вибирати бренд вручну.").grid(
            row=2, column=0, columnspan=2, sticky="w", padx=6, pady=(0, 8)
        )

        categories = self.app.flatten_categories()
        category_labels = [c["label"] for c in categories]
        category_default = self.app.settings.get("defaults", "product", "category") or ""
        category_var = self._add_var("defaults.product.category", tk.StringVar(value=category_default))
        ttk.Label(product_frame, text="Категорія за замовчуванням:").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(product_frame, textvariable=category_var, values=category_labels, width=30).grid(
            row=3, column=1, sticky="w", padx=6, pady=4
        )
        ttk.Label(product_frame, text="Залиште поле порожнім, щоб обирати категорію під час створення.").grid(
            row=4, column=0, columnspan=2, sticky="w", padx=6, pady=(0, 4)
        )

        barcode_prefix = _sanitize_barcode_prefix(self.app.settings.get("defaults", "product", "barcode_prefix") or "")
        barcode_prefix_var = self._add_var("defaults.product.barcode_prefix", tk.StringVar(value=barcode_prefix))
        ttk.Label(product_frame, text="Префікс штрихкоду (додається перед SKU):").grid(
            row=5, column=0, sticky="w", padx=6, pady=4
        )
        ttk.Entry(product_frame, textvariable=barcode_prefix_var, width=30).grid(
            row=5, column=1, sticky="w", padx=6, pady=4
        )
        ttk.Label(product_frame, text="Залиште порожнім — без префікса (немає).").grid(
            row=6, column=0, columnspan=2, sticky="w", padx=6, pady=(0, 8)
        )

        currency_frame = ttk.LabelFrame(self.defaults_tab, text="Валюти")
        currency_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        currencies = db.list_currencies(active_only=False)
        currency_codes = [c["code"] for c in currencies]
        currency_default = self.app.settings.get("defaults", "currency", "base_code") or get_base_currency_code()
        currency_var = self._add_var("defaults.currency.base_code", tk.StringVar(value=currency_default))
        ttk.Label(currency_frame, text="Базова валюта:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(currency_frame, textvariable=currency_var, values=currency_codes, width=12, state="readonly").grid(
            row=0, column=1, sticky="w", padx=6, pady=4
        )

    def build_files_tab(self) -> None:
        workdir_var = self._add_var("files.working_dir", tk.StringVar(value=str(self.app.default_workdir())))
        ttk.Label(self.files_tab, text="Робоча директорія:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(self.files_tab, textvariable=workdir_var, width=40).grid(row=0, column=1, sticky="w", padx=6, pady=4)
        ttk.Button(self.files_tab, text="Огляд", command=lambda: self.choose_workdir(workdir_var)).grid(
            row=0, column=2, padx=4, pady=4
        )

        recent_limit_var = self._add_var(
            "files.recent_limit", tk.IntVar(value=int(self.app.settings.get("files", "recent_limit") or 10))
        )
        ttk.Label(self.files_tab, text="Кількість недавніх:").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Spinbox(self.files_tab, from_=0, to=50, textvariable=recent_limit_var, width=8).grid(
            row=1, column=1, sticky="w", padx=6, pady=4
        )

        self.recent_label = ttk.Label(
            self.files_tab,
            text=f"Збережено недавніх: {len(self.app.settings.get('files', 'recent_items', default=[]))}",
        )
        self.recent_label.grid(row=2, column=0, columnspan=2, sticky="w", padx=6, pady=4)
        ttk.Button(self.files_tab, text="Очистити", command=self.clear_recent).grid(row=2, column=2, padx=4, pady=4)

        encoding_var = self._add_var("files.encoding", tk.StringVar(value=self.app.settings.get("files", "encoding")))
        ttk.Label(self.files_tab, text="Кодування тексту:").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(
            self.files_tab, textvariable=encoding_var, values=["utf-8", "cp1251", "latin-1", "utf-16"], width=12
        ).grid(row=3, column=1, sticky="w", padx=6, pady=4)

    def build_ui_tab(self) -> None:
        layout_var = self._add_var("ui.panel_layout", tk.StringVar(value=self.app.settings.get("ui", "panel_layout")))
        ttk.Label(self.ui_tab, text="Розташування панелей:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(
            self.ui_tab,
            textvariable=layout_var,
            values=["Авто", "Вертикально", "Горизонтально"],
            width=16,
        ).grid(row=0, column=1, sticky="w", padx=6, pady=4)

        status_var = self._add_var("ui.status_bar", tk.BooleanVar(value=bool(self.app.settings.get("ui", "status_bar"))))
        ttk.Checkbutton(self.ui_tab, text="Показувати рядок стану", variable=status_var).grid(
            row=1, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        compact_var = self._add_var("ui.compact_mode", tk.BooleanVar(value=bool(self.app.settings.get("ui", "compact_mode"))))
        ttk.Checkbutton(self.ui_tab, text="Компактний режим", variable=compact_var).grid(
            row=2, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        fullscreen_var = self._add_var(
            "ui.fullscreen", tk.BooleanVar(value=bool(self.app.settings.get("ui", "fullscreen")))
        )
        ttk.Checkbutton(self.ui_tab, text="Повноекранний режим", variable=fullscreen_var).grid(
            row=3, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        volume_var = self._add_var(
            "ui.notifications_volume",
            tk.IntVar(value=int(self.app.settings.get("ui", "notifications_volume") or 70)),
        )
        ttk.Label(self.ui_tab, text="Гучність сповіщень:").grid(row=4, column=0, sticky="w", padx=6, pady=4)
        ttk.Scale(self.ui_tab, from_=0, to=100, variable=volume_var, orient=tk.HORIZONTAL, length=160).grid(
            row=4, column=1, sticky="w", padx=6, pady=4
        )

        duration_var = self._add_var(
            "ui.notifications_duration",
            tk.IntVar(value=int(self.app.settings.get("ui", "notifications_duration") or 3)),
        )
        ttk.Label(self.ui_tab, text="Тривалість сповіщень (сек):").grid(row=5, column=0, sticky="w", padx=6, pady=4)
        ttk.Spinbox(self.ui_tab, from_=1, to=30, textvariable=duration_var, width=8).grid(
            row=5, column=1, sticky="w", padx=6, pady=4
        )

    def build_editor_tab(self) -> None:
        family_var = self._add_var("editor.font_family", tk.StringVar(value=self.app.settings.get("editor", "font_family")))
        ttk.Label(self.editor_tab, text="Шрифт редактора:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        font_box = ttk.Combobox(self.editor_tab, textvariable=family_var, values=sorted(tkfont.families()), width=24)
        font_box.grid(row=0, column=1, sticky="w", padx=6, pady=4)

        size_var = self._add_var("editor.font_size", tk.IntVar(value=int(self.app.settings.get("editor", "font_size") or 10)))
        ttk.Label(self.editor_tab, text="Розмір шрифту:").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Spinbox(self.editor_tab, from_=8, to=28, textvariable=size_var, width=6).grid(
            row=1, column=1, sticky="w", padx=6, pady=4
        )

        syntax_var = self._add_var(
            "editor.syntax_highlighting", tk.BooleanVar(value=bool(self.app.settings.get("editor", "syntax_highlighting")))
        )
        ttk.Checkbutton(self.editor_tab, text="Підсвічування синтаксису", variable=syntax_var).grid(
            row=2, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        indent_tabs_var = self._add_var(
            "editor.indent_with_tabs", tk.BooleanVar(value=bool(self.app.settings.get("editor", "indent_with_tabs")))
        )
        ttk.Checkbutton(self.editor_tab, text="Використовувати табуляцію для відступів", variable=indent_tabs_var).grid(
            row=3, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        tab_width_var = self._add_var("editor.tab_width", tk.IntVar(value=int(self.app.settings.get("editor", "tab_width") or 4)))
        ttk.Label(self.editor_tab, text="Ширина табуляції:").grid(row=4, column=0, sticky="w", padx=6, pady=4)
        ttk.Spinbox(self.editor_tab, from_=2, to=12, textvariable=tab_width_var, width=6).grid(
            row=4, column=1, sticky="w", padx=6, pady=4
        )

        line_numbers_var = self._add_var(
            "editor.line_numbers", tk.BooleanVar(value=bool(self.app.settings.get("editor", "line_numbers")))
        )
        ttk.Checkbutton(self.editor_tab, text="Показувати номери рядків", variable=line_numbers_var).grid(
            row=5, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        minimap_var = self._add_var("editor.minimap", tk.BooleanVar(value=bool(self.app.settings.get("editor", "minimap"))))
        ttk.Checkbutton(self.editor_tab, text="Вмикати мінімеп", variable=minimap_var).grid(
            row=6, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        auto_format_var = self._add_var(
            "editor.auto_format", tk.BooleanVar(value=bool(self.app.settings.get("editor", "auto_format")))
        )
        ttk.Checkbutton(self.editor_tab, text="Автоматичне форматування", variable=auto_format_var).grid(
            row=7, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        autocomplete_var = self._add_var(
            "editor.autocomplete", tk.BooleanVar(value=bool(self.app.settings.get("editor", "autocomplete")))
        )
        ttk.Checkbutton(self.editor_tab, text="Автодоповнення", variable=autocomplete_var).grid(
            row=8, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        length_limit_var = self._add_var(
            "editor.line_length_limit",
            tk.IntVar(value=int(self.app.settings.get("editor", "line_length_limit") or 120)),
        )
        ttk.Label(self.editor_tab, text="Фільтр по довжині рядка:").grid(row=9, column=0, sticky="w", padx=6, pady=4)
        ttk.Spinbox(self.editor_tab, from_=40, to=240, textvariable=length_limit_var, width=6).grid(
            row=9, column=1, sticky="w", padx=6, pady=4
        )

    def build_hotkeys_tab(self) -> None:
        profile_var = self._add_var("hotkeys.profile", tk.StringVar(value=self.app.settings.get("hotkeys", "profile")))
        ttk.Label(self.hotkeys_tab, text="Профіль клавіш:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(
            self.hotkeys_tab, textvariable=profile_var, values=["Типовий", "Emacs", "Vim"], width=12
        ).grid(row=0, column=1, sticky="w", padx=6, pady=4)

        allow_custom_var = self._add_var(
            "hotkeys.allow_custom", tk.BooleanVar(value=bool(self.app.settings.get("hotkeys", "allow_custom")))
        )
        ttk.Checkbutton(self.hotkeys_tab, text="Дозволити переназначення", variable=allow_custom_var).grid(
            row=1, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

    def build_support_tab(self) -> None:
        log_var = self._add_var("support.log_level", tk.StringVar(value=self.app.settings.get("support", "log_level")))
        ttk.Label(self.support_tab, text="Рівень логування:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(self.support_tab, textvariable=log_var, values=["DEBUG", "INFO", "WARNING", "ERROR"], width=12).grid(
            row=0, column=1, sticky="w", padx=6, pady=4
        )

        sysinfo_var = self._add_var(
            "support.collect_system_info", tk.BooleanVar(value=bool(self.app.settings.get("support", "collect_system_info")))
        )
        ttk.Checkbutton(self.support_tab, text="Збирати системну інформацію", variable=sysinfo_var).grid(
            row=1, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        autoreport_var = self._add_var(
            "support.auto_error_reports", tk.BooleanVar(value=bool(self.app.settings.get("support", "auto_error_reports")))
        )
        ttk.Checkbutton(self.support_tab, text="Надсилати звіти про помилки автоматично", variable=autoreport_var).grid(
            row=2, column=0, columnspan=2, sticky="w", padx=6, pady=4
        )

        docs_var = self._add_var("support.docs_url", tk.StringVar(value=self.app.settings.get("support", "docs_url")))
        ttk.Label(self.support_tab, text="Документація/FAQ:").grid(row=3, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(self.support_tab, textvariable=docs_var, width=40).grid(row=3, column=1, sticky="w", padx=6, pady=4)

    def choose_workdir(self, var: tk.StringVar) -> None:
        path = filedialog.askdirectory(initialdir=var.get() or str(self.app.default_workdir()))
        if path:
            var.set(path)

    def clear_recent(self) -> None:
        self.recent_cleared = True
        self.recent_label.configure(text="Збережено недавніх: 0")

    def _add_var(self, path: str, var: tk.Variable) -> tk.Variable:
        self.vars[path] = var
        return var

    def save(self) -> None:
        for path, var in self.vars.items():
            value = var.get()
            if path == "defaults.product.barcode_prefix":
                value = _sanitize_barcode_prefix(str(value))
            keys = path.split(".")
            self.app.settings.set(value, *keys)
        base_code = str(self.vars.get("defaults.currency.base_code", tk.StringVar()).get()).strip().upper()
        if base_code:
            currencies = {c["code"]: c for c in db.list_currencies(active_only=False)}
            selected = currencies.get(base_code)
            if selected:
                self.app.settings.set(selected["name"], "defaults", "currency", "base_name")
                self.app.settings.set(int(selected["decimals"]), "defaults", "currency", "base_decimals")
            else:
                self.app.settings.set(base_code, "defaults", "currency", "base_name")
                self.app.settings.set(get_base_currency_decimals(), "defaults", "currency", "base_decimals")
        if self.recent_cleared:
            self.app.settings.set([], "files", "recent_items")
        self.app.settings.save()
        apply_base_currency_settings(self.app.settings)
        self.app.refresh_currencies()
        self.app.apply_settings()
        self.destroy()


def main() -> None:
    configure_logging()
    logging.info("Starting %s", APP_NAME)
    try:
        with SingleInstance(get_lock_path()):
            settings = Settings()
            apply_base_currency_settings(settings)
            db.init_db()
            app = InventoryApp(settings)
            app.mainloop()
    except RuntimeError:
        messagebox.showwarning(APP_NAME, "Програма вже запущена.")
    except ValueError as exc:
        messagebox.showerror(APP_NAME, str(exc))
    except Exception:
        logging.exception("Fatal error")
        try:
            messagebox.showerror(APP_NAME, f"Критична помилка. Деталі у логах: {get_log_path()}")
        except tk.TclError:
            print("Критична помилка. Деталі у логах:", get_log_path(), file=sys.stderr)
        traceback.print_exc()


if __name__ == "__main__":
    main()

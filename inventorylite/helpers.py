from __future__ import annotations

import csv
import datetime as dt
import math
import re
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook


def _parse_num(text: str) -> float | None:
    if text is None:
        return None
    s = str(text).strip().replace(" ", "").replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _read_rate_two_way(cur: str, base: str, v_direct: str, v_inverse: str) -> float:
    """
    Returns canonical rate: 1 cur = rate base
    Accepts either:
      direct  (1 cur = X base)
      inverse (1 base = X cur) -> rate = 1/X
    """
    a = _parse_num(v_direct)  # 1 cur = a base
    b = _parse_num(v_inverse)  # 1 base = b cur
    if a is None and b is None:
        raise ValueError("Курс не вказано")
    if a is not None and a <= 0:
        raise ValueError("Курс має бути більшим за 0")
    if b is not None and b <= 0:
        raise ValueError("Курс має бути більшим за 0")

    if a is None:
        return 1.0 / b
    if b is None:
        return a

    # Якщо введено обидва — перевір узгодженість (інакше помилка)
    inv = 1.0 / b
    # допустимо 0.5% різниці через округлення
    if a == 0 or abs(a - inv) / a > 0.005:
        raise ValueError("Курси не узгоджуються (перевір обидва поля)")
    return a


def _parse_date_value(raw: str) -> str:
    raw = (raw or "").strip()
    if not raw:
        return dt.datetime.now().strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return dt.datetime.now().strftime("%Y-%m-%d")


def _parse_float_value(raw: str) -> float:
    if raw is None:
        return 0.0
    text = str(raw).replace(" ", "").replace(",", ".").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def _format_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.datetime.min.time()).strftime("%Y-%m-%d")
    return str(value)


def _normalize_header(value: str) -> str:
    """Strip whitespace/BOM from column headers to avoid mapping typos."""

    return (value or "").strip().lstrip("\ufeff")


SalesField = tuple[str, str, tuple[str, ...]]


SALES_FIELDS: list[SalesField] = [
    ("order_no", "Замовлення", ("номер", "замовлення", "order", "order_id", "order_no", "id")),
    ("doc_date", "Дата", ("дата", "date", "order_date", "дата оформлення")),
    ("customer", "Клієнт", ("клієнт", "покупець", "customer", "контрагент")),
    ("phone", "Телефон", ("телефон", "phone")),
    ("email", "Email", ("email", "e-mail")),
    ("sku", "SKU", ("sku", "артикул", "код")),
    ("supplier_sku", "Артикул постачальника", ("артикул постачальника", "supplier_sku", "vendor_sku", "vendor code")),
    ("product_name", "Товар", ("товар", "product", "назва", "item")),
    ("quantity", "Кількість", ("кількість", "к-сть", "qty", "quantity", "шт")),
    ("price", "Ціна", ("ціна", "price", "amount")),
    ("amount", "Сума", ("сума", "amount", "total")),
    ("discount", "Знижка", ("знижка", "discount")),
    ("comment", "Коментар", ("коментар", "примітка", "comment", "note")),
    ("channel", "Канал", ("канал", "channel", "майданчик", "площадка", "platform")),
]


PURCHASE_FIELDS: list[SalesField] = [
    ("sku", "SKU", ("sku", "артикул", "код")),
    ("supplier_sku", "Артикул постачальника", ("артикул постачальника", "supplier_sku", "vendor_sku", "vendor code")),
    ("product_name", "Товар", ("товар", "product", "назва", "item")),
    ("quantity", "Кількість", ("кількість", "к-сть", "qty", "quantity", "шт")),
    ("price", "Ціна", ("ціна", "price", "amount")),
]


PRODUCT_FIELDS: list[SalesField] = [
    ("sku", "SKU", ("sku", "артикул", "код")),
    ("name", "Назва", ("назва", "name", "product", "товар", "item")),
    (
        "supplier_sku",
        "Артикул постачальника",
        ("артикул постачальника", "supplier_sku", "vendor_sku", "vendor code"),
    ),
    ("brand", "Бренд", ("бренд", "brand")),
    ("brand_id", "ID бренду", ("brand id", "id бренду")),
    ("category", "Категорія", ("категорія", "category")),
    ("category_id", "ID категорії", ("category id", "id категорії")),
    ("unit", "Одиниця", ("одиниця", "unit", "шт")),
    ("is_active", "Активний", ("активний", "is_active", "active")),
    (
        "extra_categories",
        "Додаткові категорії",
        ("додаткові категорії", "extra categories", "tags", "категорії"),
    ),
]


def _suggest_sales_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {h.lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for key, _label, aliases in SALES_FIELDS:
        for alias in aliases:
            if alias.lower() in normalized_headers:
                mapping[key] = normalized_headers[alias.lower()]
                break
        else:
            mapping[key] = ""
    return mapping


def _suggest_purchase_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {h.lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for key, _label, aliases in PURCHASE_FIELDS:
        for alias in aliases:
            if alias.lower() in normalized_headers:
                mapping[key] = normalized_headers[alias.lower()]
                break
        else:
            mapping[key] = ""
    return mapping


def _suggest_product_mapping(headers: list[str]) -> dict[str, str]:
    normalized_headers = {h.lower(): h for h in headers}
    mapping: dict[str, str] = {}
    for key, _label, aliases in PRODUCT_FIELDS:
        for alias in aliases:
            if alias.lower() in normalized_headers:
                mapping[key] = normalized_headers[alias.lower()]
                break
        else:
            mapping[key] = ""
    return mapping


def _normalize_sales_records(rows: list[dict[str, object]], mapping: dict[str, str]) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        normalized = {_normalize_header(k): _format_cell_value(v).strip() for k, v in row.items()}

        def pick(field: str, parser=None):
            header = mapping.get(field, "")
            value = normalized.get(header, "") if header else ""
            return parser(value) if parser else value

        records.append(
            {
                "order_no": pick("order_no"),
                "doc_date": pick("doc_date", _parse_date_value),
                "customer": pick("customer"),
                "phone": pick("phone"),
                "email": pick("email"),
                "sku": pick("sku"),
                "supplier_sku": pick("supplier_sku"),
                "product_name": pick("product_name"),
                "quantity": pick("quantity", _parse_float_value),
                "price": pick("price", _parse_float_value),
                "amount": pick("amount", _parse_float_value),
                "discount": pick("discount", _parse_float_value),
                "comment": pick("comment"),
                "channel": pick("channel"),
            }
        )
    return records


def _parse_bool_value(raw: str) -> Optional[int]:
    text = (raw or "").strip().lower()
    if not text:
        return None
    if text in {"1", "true", "yes", "так", "y", "t", "on"}:
        return 1
    if text in {"0", "false", "no", "ні", "n", "off", "f"}:
        return 0
    return None


def _normalize_product_records(rows: list[dict[str, object]], mapping: dict[str, str]) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        normalized = {_normalize_header(k): _format_cell_value(v) for k, v in row.items()}

        def pick(field: str, parser=None):
            header = mapping.get(field, "")
            value = normalized.get(header, "") if header else ""
            return parser(value) if parser else (value.strip() if isinstance(value, str) else value)

        extra_raw = pick("extra_categories") or ""
        extra_items = []
        for part in re.split(r"[;,|]", extra_raw):
            clean = part.strip()
            if clean and clean not in extra_items:
                extra_items.append(clean)

        def parse_int(value: object) -> Optional[int]:
            try:
                num = int(str(value).strip())
                return num if num > 0 else None
            except Exception:
                return None

        records.append(
            {
                "sku": pick("sku"),
                "name": pick("name"),
                "supplier_sku": pick("supplier_sku"),
                "brand": pick("brand"),
                "brand_id": parse_int(pick("brand_id")),
                "category": pick("category"),
                "category_id": parse_int(pick("category_id")),
                "unit": pick("unit"),
                "is_active": _parse_bool_value(pick("is_active")),
                "extra_categories": extra_items,
            }
        )
    return records


def _normalize_purchase_records(
    rows: list[dict[str, object]],
    mapping: dict[str, str],
    default_supplier: str | None = None,
    default_doc_date: str | None = None,
    default_order_no: str | None = None,
    default_comment: str | None = None,
) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        normalized = {_normalize_header(k): _format_cell_value(v).strip() for k, v in row.items()}

        def pick(field: str, parser=None):
            header = mapping.get(field, "")
            value = normalized.get(header, "") if header else ""
            return parser(value) if parser else value

        records.append(
            {
                "order_no": (default_order_no or "").strip(),
                "doc_date": default_doc_date or pick("doc_date", _parse_date_value),
                "supplier": default_supplier if default_supplier else pick("supplier"),
                "sku": pick("sku"),
                "supplier_sku": pick("supplier_sku"),
                "product_name": pick("product_name"),
                "quantity": pick("quantity", _parse_float_value),
                "price": pick("price", _parse_float_value),
                "amount": pick("amount", _parse_float_value),
                "comment": (default_comment or "").strip(),
            }
        )
    return records


def _read_import_csv(path: Path, encoding: str) -> tuple[list[dict[str, object]], list[str]]:
    with path.open("r", encoding=encoding, newline="") as f:
        sample = f.read(2048)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample) if sample else csv.excel
        except Exception:
            dialect = csv.excel
        reader = csv.DictReader(f, dialect=dialect)
        raw_headers = list(reader.fieldnames or [])
        normalized_headers = [_normalize_header(h or "") for h in raw_headers]
        header_map = {raw or "": normalized for raw, normalized in zip(raw_headers, normalized_headers)}

        rows = []
        for row in reader:
            cleaned_row: dict[str, object] = {}
            for raw_key, value in row.items():
                normalized_key = header_map.get(raw_key or "", _normalize_header(raw_key or ""))
                cleaned_row[normalized_key] = value
            rows.append(cleaned_row)

        return rows, normalized_headers


def _read_import_xlsx(path: Path) -> tuple[list[dict[str, object]], list[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [_normalize_header(_format_cell_value(cell)) for cell in rows[0]]
    records: list[dict[str, object]] = []
    for row in rows[1:]:
        record: dict[str, object] = {}
        for idx, value in enumerate(row):
            header = headers[idx] if idx < len(headers) else ""
            record[header] = value
        records.append(record)
    return records, headers


def _read_import_xls(path: Path) -> tuple[list[dict[str, object]], list[str]]:
    try:
        # Деякі сервіси експортують XLSX-файли з розширенням .xls, тому
        # спершу пробуємо прочитати їх через openpyxl.
        return _read_import_xlsx(path)
    except Exception:
        # Якщо це справді старий XLS, повертаємося до xlrd.
        try:
            import xlrd
        except ImportError as exc:
            raise ImportError(
                "Для імпорту XLS-файлів потрібно встановити залежність 'xlrd'."
            ) from exc

    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return [], []

    headers = [_normalize_header(_format_cell_value(sheet.cell_value(0, col))) for col in range(sheet.ncols)]
    records: list[dict[str, object]] = []
    for row_idx in range(1, sheet.nrows):
        record: dict[str, object] = {}
        for col_idx in range(sheet.ncols):
            header = headers[col_idx] if col_idx < len(headers) else ""
            cell = sheet.cell(row_idx, col_idx)
            value: object = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                except Exception:
                    pass
            record[header] = value
        records.append(record)
    return records, headers


def parse_import_file(path: Path, encoding: str = "utf-8") -> tuple[list[dict[str, object]], list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        raw_rows, headers = _read_import_xlsx(path)
    elif suffix == ".xls":
        raw_rows, headers = _read_import_xls(path)
    else:
        raw_rows, headers = _read_import_csv(path, encoding)

    return raw_rows, headers


def parse_sales_file(path: Path, encoding: str = "utf-8") -> tuple[list[dict[str, object]], list[str]]:
    return parse_import_file(path, encoding)


def _find_index_by_name(items: list[str], target: str | None) -> int | None:
    if not target:
        return None
    target_lower = target.lower()
    return next((i for i, name in enumerate(items) if str(name).lower() == target_lower), None)


def _sanitize_barcode_prefix(prefix: str) -> str:
    cleaned = (prefix or "").strip()
    if not cleaned:
        return ""
    return re.sub(r"\s+", "-", cleaned)
